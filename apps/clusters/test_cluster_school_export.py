"""The Cluster Profile exports the cluster's schools.

Owner, 2026-10-17: "I want to export the schools from the cluster list. the
export should have the school ID, school name, Phone, Primary Contact, and
Shipping Address" — and "the export buttons should be inside the cluster
profile".

The schools are the ones the Cluster Page lists: its operating members,
without the Champion schools that have their own table.
"""

from __future__ import annotations

import io

from django.contrib.auth import get_user_model

from apps.accounts.models import StaffSupervisorAssignment
from apps.clusters.test_eligibility import EligibilityFixture
from apps.core.rbac import EdifyRole

User = get_user_model()

HEADERS = ["School ID", "School Name", "Phone", "Primary Contact", "Shipping Address"]


def _rows(response):
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


class ClusterSchoolExportTest(EligibilityFixture):
    def setUp(self):
        super().setUp()
        self._join(self.school, self.chegere_north)
        self.school.school_phone = "0772 000 111"
        self.school.primary_contact_name = "Grace Akello"
        self.school.shipping_address = "Plot 4, Chegere Road, Apac"
        self.school.save()

        self.second = self._join(
            self._school(
                "ST-PAUL",
                "St. Paul Primary School",
                self.james_profile,
                district=self.apac,
                sub_county=self.chegere,
            ),
            self.chegere_north,
        )
        self.elsewhere = self._join(
            self._school(
                "AKO-1",
                "Akokoro Primary School",
                self.james_profile,
                district=self.apac,
                sub_county=self.akokoro,
            ),
            self.akokoro_cluster,
        )
        self.admin = User.objects.create(
            email="export-admin@edify.test",
            name="Export Admin",
            roles=[EdifyRole.ADMIN.value],
            active_role=EdifyRole.ADMIN.value,
            is_active=True,
        )
        self.url = f"/clusters/{self.chegere_north.id}/schools/export"

    def _join(self, school, cluster):
        school.cluster_id = cluster.id
        school.save()
        return school

    def _export(self, user=None):
        self.client.force_login(user or self.admin)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        return _rows(response)

    def test_the_export_has_the_five_columns_with_each_schools_details(self):
        rows = self._export()

        self.assertEqual(rows[0], HEADERS)
        self.assertIn(
            [
                "ST-MARY",
                "St. Mary Primary School",
                "0772 000 111",
                "Grace Akello",
                "Plot 4, Chegere Road, Apac",
            ],
            rows[1:],
        )

    def test_it_lists_this_clusters_schools_and_no_others(self):
        ids = [row[0] for row in self._export()[1:]]

        self.assertEqual(ids, ["ST-MARY", "ST-PAUL"])

    def test_champion_schools_are_left_out_as_they_are_on_the_cluster_page(self):
        self.second.school_type = "champion"
        self.second.save()

        ids = [row[0] for row in self._export()[1:]]

        self.assertEqual(ids, ["ST-MARY"])

    def test_a_missing_detail_is_a_blank_cell(self):
        rows = self._export()

        self.assertEqual(
            rows[2], ["ST-PAUL", "St. Paul Primary School", None, None, None]
        )

    def test_a_typed_formula_is_text_but_an_international_phone_is_kept(self):
        self.second.primary_contact_name = '=HYPERLINK("http://x","y")'
        self.second.school_phone = "+256 772 000 222"
        self.second.save()

        row = self._export()[2]

        self.assertEqual(row[3], '\'=HYPERLINK("http://x","y")')
        self.assertEqual(row[2], "+256 772 000 222")

    def test_the_file_is_named_after_the_cluster(self):
        self.client.force_login(self.admin)

        response = self.client.get(self.url)

        self.assertIn(
            'filename="Chegere North Cluster schools.xlsx"',
            response["Content-Disposition"],
        )

    def test_a_cluster_outside_the_readers_scope_is_not_exported(self):
        lead = User.objects.create(
            email="export-lead@edify.test",
            name="Export Lead",
            roles=[EdifyRole.COUNTRY_PROGRAM_LEAD.value],
            active_role=EdifyRole.COUNTRY_PROGRAM_LEAD.value,
            is_active=True,
        )
        from apps.accounts.models import StaffProfile

        profile = StaffProfile.objects.create(user=lead, title="Lead")
        StaffSupervisorAssignment.objects.create(
            supervisee=self.mary_profile, supervisor=profile
        )
        self.client.force_login(lead)

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 404)

    def test_the_cluster_profile_offers_the_export(self):
        self.client.force_login(self.admin)

        response = self.client.get(f"/clusters/{self.chegere_north.id}/profile")

        self.assertContains(response, f'href="{self.url}"')
        self.assertContains(response, "Export schools")

    def test_whoever_opens_the_profile_is_offered_and_served_it(self):
        """Owner, 2026-10-31: everyone who can open a cluster's profile exports
        its schools, CCEOs included, though the role holds no `data.export`.
        It is a copy of a list they already read on the cluster's own page,
        so the page's scope is the control (apps/core/tests/
        test_export_authorization.py: extraction of what you already see)."""
        self.client.force_login(self.james)

        profile = self.client.get(f"/clusters/{self.chegere_north.id}/profile")

        self.assertContains(profile, f'href="{self.url}"')
        self.assertContains(profile, "Export schools")
        ids = [row[0] for row in self._export(self.james)[1:]]
        self.assertEqual(ids, ["ST-MARY", "ST-PAUL"])

    def test_a_cceo_who_does_not_hold_the_cluster_is_not_served_it(self):
        self.client.force_login(self.mary)

        response = self.client.get(self.url)

        self.assertNotIn("spreadsheetml", response.get("Content-Type", ""))
