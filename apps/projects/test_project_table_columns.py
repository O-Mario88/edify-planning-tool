"""The project table's columns, and its export.

Owner, 2026-10-05:

  "on project table here are the columns (Project name, School ID, School
  Name, District, Training, Purpose of assignment, SSA intervention ("General"
  for Alumni), Previous SSA Score, Current SSA Score, SSA Improvement (Current
  SSA score - Previous SSA Score), Status, Activity date, Enrolled on,
  Planning Stage (Awaiting {Partner Name}), Planned By (Partner Name),
  Execution, Activity Status (Scheduled, Completed, Canceled, rescheduled),
  Activity Cost (Only fetch if scheduled). The IA, project coordinator or CD
  needs to export either all projects in 1 file or select a project to
  export"
"""

from __future__ import annotations

import io
import re
from datetime import timedelta

# Imported here, before any test freezes the clock: openpyxl notes the date
# types it formats when it is first imported, and under a frozen clock those
# are freezegun's.
from openpyxl import load_workbook

from apps.activities.models import Activity, ActivityScheduleCostLine
from apps.projects import monitoring, school_table
from apps.projects.models import (
    GENERAL_INTERVENTION,
    Project,
    ProjectSchoolAssignment,
)
from apps.projects.test_project_monitoring_schools import FOCUS, _Fixture, _user
from apps.schools.models import School

#: On the page a cost is written with its currency; in the workbook the cell
#: is a number and the heading carries it.
PAGE_COLUMNS = [
    "Activity Cost" if heading == "Activity Cost (UGX)" else heading
    for heading in school_table.COLUMNS
]


def _headings(body: str) -> list[str]:
    start = body.index("data-project-schools")
    table = body[start : body.index("</thead>", start)]
    return [
        re.sub(r"\s*<br>\s*", " ", heading)
        for heading in re.findall(r'<th scope="col"[^>]*>(.+?)</th>', table)
    ]


def _cost(activity, amount: int, key: str = "test_cost") -> None:
    ActivityScheduleCostLine.objects.create(
        activity=activity,
        cost_setting_key=key,
        label="Test cost",
        unit_cost=amount,
        quantity=1,
        amount=amount,
    )


class _ColumnsFixture(_Fixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cd_user, _ = _user("pms-cd", "CountryDirector", "PMS Director")
        cls.cceo_user, _ = _user("pms-cceo", "CCEO", "PMS Officer")
        cls.visit = Activity.objects.get(school=cls.coordinated.school)
        _cost(cls.visit, 75_000)

    def workbook(self, user, path, **query):
        self.client.force_login(user)
        response = self.client.get(path, {"fy": self.fy, **query})
        self.assertEqual(response.status_code, 200, path)
        return response, load_workbook(io.BytesIO(response.content))


class TheOwnersColumnsTest(_ColumnsFixture):
    def test_the_table_has_the_eighteen_columns_in_the_owners_order(self):
        self.assertEqual(
            list(school_table.COLUMNS),
            [
                "Project Name",
                "School ID",
                "School Name",
                "District",
                "Training",
                "Purpose of Assignment",
                "SSA Intervention",
                "Previous SSA Score",
                "Current SSA Score",
                "SSA Improvement",
                "Status",
                "Activity Date",
                "Enrolled On",
                "Planning Stage",
                "Planned By",
                "Execution",
                "Activity Status",
                "Activity Cost (UGX)",
            ],
        )

    def test_project_monitoring_draws_them_then_its_actions(self):
        for user in (self.lead_user, self.ia_user, self.cd_user, self.coord_user):
            with self.subTest(role=user.active_role):
                self.client.force_login(user)
                body = self.client.get(
                    "/projects/monitoring", {"fy": self.fy}
                ).content.decode()
                self.assertEqual(_headings(body), [*PAGE_COLUMNS, "Actions"])

    def test_planning_oversight_s_special_projects_tab_draws_the_same(self):
        self.client.force_login(self.ia_user)
        body = self.client.get(
            "/team-planning-oversight/", {"view": "projects", "fy": self.fy}
        ).content.decode()
        self.assertEqual(_headings(body), [*PAGE_COLUMNS, "Actions"])

    def test_the_workbook_s_first_sheet_is_exactly_the_columns(self):
        _response, book = self.workbook(self.ia_user, "/projects/export")
        sheet = book["Project Schools"]
        self.assertEqual([cell.value for cell in sheet[1]], list(school_table.COLUMNS))


class WhatEachCellSaysTest(_ColumnsFixture):
    def test_every_row_names_its_project(self):
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(
            {row.project_name for row in rows.values()}, {"PMS Learning Project"}
        )

    def test_the_scores_are_the_one_at_enrolment_and_the_latest_since(self):
        _result, rows = self.rows_for(self.ia_user)
        verified = rows["Verified Primary"]
        self.assertEqual(
            (verified.previous_score, verified.current_score, verified.ssa_improvement),
            (4.0, 6.0, 2.0),
        )

    def test_no_follow_up_ssa_is_no_current_score_and_no_improvement(self):
        _result, rows = self.rows_for(self.ia_user)
        delivered = rows["Delivered Primary"]
        self.assertEqual(delivered.previous_score, 3.0)
        self.assertIsNone(delivered.current_score)
        self.assertIsNone(delivered.ssa_improvement)
        unplanned = rows["Unplanned Primary"]
        self.assertIsNone(unplanned.previous_score)
        self.assertIsNone(unplanned.ssa_improvement)

    def test_a_score_that_fell_is_a_negative_improvement(self):
        self._ssa(self.delivered.school, self.today - timedelta(days=1), 2.5, 4.0)
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(rows["Delivered Primary"].ssa_improvement, -0.5)

    def test_improvement_is_green_up_red_down_and_neutral_for_no_change(self):
        """Owner, 2026-10-05: "SSA improvement should be color coded green for
        +(improved), Red for -(declined) and neutral for no change"."""
        self._ssa(self.delivered.school, self.today - timedelta(days=1), 2.5, 4.0)
        self._ssa(self.coordinated.school, self.today - timedelta(days=200), 5.0, 5.0)
        self._ssa(self.coordinated.school, self.today - timedelta(days=1), 5.0, 5.0)
        ProjectSchoolAssignment.objects.filter(id=self.coordinated.id).update(
            start_date=self.today - timedelta(days=100)
        )
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(
            {
                name: (rows[name].ssa_improvement, rows[name].ssa_improvement_tone)
                for name in (
                    "Verified Primary",
                    "Delivered Primary",
                    "Coordinated Primary",
                    "Unplanned Primary",
                )
            },
            {
                "Verified Primary": (2.0, "success"),
                "Delivered Primary": (-0.5, "danger"),
                "Coordinated Primary": (0.0, "neutral"),
                "Unplanned Primary": (None, "neutral"),
            },
        )

        self.client.force_login(self.ia_user)
        body = self.client.get("/projects/monitoring", {"fy": self.fy}).content.decode()
        cells = re.findall(r"data-ssa-improvement>(.*?)</td>", body, flags=re.S)
        self.assertIn(
            '<span class="partner-status-chip" data-tone="success">+2.0</span>', cells
        )
        self.assertIn(
            '<span class="partner-status-chip" data-tone="danger">-0.5</span>', cells
        )
        self.assertIn(
            '<span class="partner-status-chip" data-tone="neutral">0.0</span>', cells
        )
        # No score to compare is no colour at all.
        self.assertIn("—", cells)

    def test_the_workbook_colours_the_improvement_the_same_way(self):
        _response, book = self.workbook(
            self.ia_user, f"/projects/{self.project.id}/export"
        )
        sheet = book["Project Schools"]
        column = school_table.COLUMNS.index("SSA Improvement") + 1
        formats = {
            sheet.cell(row, column).number_format for row in range(2, sheet.max_row + 1)
        }
        # One format for the column: green above nought, red below, plain at.
        self.assertEqual(formats, {"[Color10]+0.0#;[Red]-0.0#;0.0#"})

    def test_work_a_partner_holds_undated_awaits_that_partner_by_name(self):
        _result, rows = self.rows_for(self.ia_user)
        handed = rows["Handed Primary"]
        self.assertEqual(handed.planning_stage, "Awaiting PMS Partner")
        # Nobody has planned it yet: the Partner is named once it has.
        self.assertEqual(handed.planned_by_name, "")
        self.assertEqual(
            rows["Returned Primary"].planning_stage, "Returned by PMS Partner"
        )

    def test_planned_by_is_the_partner_once_the_partner_has_scheduled(self):
        _result, rows = self.rows_for(self.ia_user)
        partnered = rows["Partnered Primary"]
        self.assertEqual(partnered.planning_stage, "Partner scheduled")
        self.assertEqual(partnered.planned_by_name, "PMS Partner")

    def test_planned_by_is_the_staff_member_who_planned_it(self):
        _result, rows = self.rows_for(self.ia_user)
        coordinated = rows["Coordinated Primary"]
        self.assertEqual(coordinated.planning_stage, "Coordinator planned")
        self.assertEqual(coordinated.planned_by_name, "PMS Coordinator")
        self.assertEqual(rows["Unplanned Primary"].planned_by_name, "")

    def test_execution_says_how_much_of_the_work_is_done(self):
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(rows["Unplanned Primary"].execution_summary, "No activity")
        self.assertEqual(
            rows["Verified Primary"].execution_summary, "Verified · 1/1 done"
        )


class ActivityStatusTest(_ColumnsFixture):
    """Scheduled, Rescheduled, Completed or Cancelled — and nothing for work
    that has no day on it."""

    def test_scheduled_and_completed(self):
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(
            {
                name: rows[name].activity_state_label
                for name in (
                    "Coordinated Primary",
                    "Partnered Primary",
                    "Delivered Primary",
                    "Verified Primary",
                )
            },
            {
                "Coordinated Primary": "Scheduled",
                "Partnered Primary": "Scheduled",
                "Delivered Primary": "Completed",
                "Verified Primary": "Completed",
            },
        )

    def test_work_with_no_day_has_no_activity_status(self):
        _result, rows = self.rows_for(self.ia_user)
        for name in ("Unplanned Primary", "Handed Primary", "Returned Primary"):
            with self.subTest(school=name):
                self.assertEqual(rows[name].activity_state_label, "")

    def test_a_moved_date_reads_rescheduled(self):
        Activity.objects.filter(id=self.visit.id).update(
            status="rescheduled", reschedule_count=1
        )
        _result, rows = self.rows_for(self.ia_user)
        row = rows["Coordinated Primary"]
        self.assertEqual(row.activity_state_label, "Rescheduled")
        self.assertEqual(row.status_label, "Scheduled")

    def test_cancelled_work_not_planned_again_reads_cancelled(self):
        Activity.objects.filter(id=self.visit.id).update(status="cancelled")
        _result, rows = self.rows_for(self.ia_user)
        row = rows["Coordinated Primary"]
        self.assertEqual(row.activity_state_label, "Cancelled")
        # The school is the coordinator's to plan again: no date stands, and
        # cancelled work is nobody's plan and nobody's cost.
        self.assertEqual(row.status_label, "Awaiting Project Coordinator Action")
        self.assertEqual(row.purpose_label, "School Visit")
        self.assertIsNone(row.activity_date)
        self.assertIsNone(row.activity_cost)
        self.assertEqual((row.planned, row.is_planned), (0, False))

    def test_work_planned_after_a_cancellation_speaks_for_the_school(self):
        Activity.objects.filter(id=self.visit.id).update(status="cancelled")
        again = self._work(
            self.coordinated.school,
            "school_visit",
            self.visit.planned_date + timedelta(days=1),
            self.coord.id,
        )
        _cost(again, 40_000)
        _result, rows = self.rows_for(self.ia_user)
        row = rows["Coordinated Primary"]
        self.assertEqual(row.activity_state_label, "Scheduled")
        self.assertEqual(row.activity_cost, 40_000)


class ActivityCostTest(_ColumnsFixture):
    """A cost is read once the activity is scheduled, from its cost lines."""

    def test_a_scheduled_activity_reads_its_cost_lines(self):
        _cost(self.visit, 25_000, "test_second_cost")
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(rows["Coordinated Primary"].activity_cost, 100_000)

    def test_a_school_with_nothing_scheduled_has_no_cost(self):
        _result, rows = self.rows_for(self.ia_user)
        for name in ("Unplanned Primary", "Handed Primary", "Returned Primary"):
            with self.subTest(school=name):
                self.assertIsNone(rows[name].activity_cost)

    def test_work_handed_to_a_partner_undated_has_no_cost_yet(self):
        held = self._work(
            self.unplanned.school,
            "school_visit",
            self.visit.planned_date,
            None,
            status="assigned_to_partner",
            partner=self.partner,
        )
        _cost(held, 60_000)
        _result, rows = self.rows_for(self.ia_user)
        row = rows["Unplanned Primary"]
        self.assertEqual(row.status_label, "Awaiting Partner Schedule")
        self.assertEqual(row.activity_state_label, "")
        self.assertIsNone(row.activity_cost)

    def test_a_scheduled_activity_with_no_cost_line_has_no_recorded_cost(self):
        _result, rows = self.rows_for(self.ia_user)
        partnered = rows["Partnered Primary"]
        self.assertEqual(partnered.activity_state_label, "Scheduled")
        self.assertIsNone(partnered.activity_cost)

    def test_an_in_school_training_reads_the_cost_its_visit_carries(self):
        school = self.unplanned.school
        visit = self._work(
            school, "school_visit", self.visit.planned_date, self.coord.id
        )
        training = self._work(
            school, "in_school_training", self.visit.planned_date, self.coord.id
        )
        Activity.objects.filter(id=training.id).update(paired_school_visit=visit)
        _cost(visit, 90_000)
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(rows["Unplanned Primary"].activity_cost, 90_000)

    def test_completed_work_keeps_the_cost_it_was_scheduled_at(self):
        done = Activity.objects.get(school=self.verified.school)
        _cost(done, 55_000)
        _result, rows = self.rows_for(self.ia_user)
        self.assertEqual(rows["Verified Primary"].activity_cost, 55_000)


class GeneralForAlumniTest(_ColumnsFixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.alumni = Project.objects.create(
            name="Alumni",
            code="SP-ALUMNI",
            category="pilot",
            status="active",
            target_interventions=[GENERAL_INTERVENTION],
            manager_staff_id=cls.coord.id,
        )
        cls.alumni_school = School.objects.create(
            school_id="PMS-9",
            name="Alumni Primary",
            region=cls.region,
            district=cls.district,
            school_type="client",
        )
        ProjectSchoolAssignment.objects.create(
            project=cls.alumni,
            school=cls.alumni_school,
            assigned_by=cls.lead_user.id,
            matched_intervention=FOCUS,
        )
        cls._ssa(cls.alumni_school, cls.today - timedelta(days=200), 4.0, 5.0)

    def alumni_row(self):
        result = monitoring.project_monitoring(self.ia_user, fy=self.fy)
        project = next(row for row in result.rows if row.id == self.alumni.id)
        return project.school_rows[0]

    def test_an_alumni_school_reads_general_and_carries_no_ssa_scores(self):
        row = self.alumni_row()
        self.assertEqual(row.intervention_label, "General")
        self.assertIsNone(row.previous_score)
        self.assertIsNone(row.current_score)

    def test_alumni_work_reads_general_whatever_its_record_names(self):
        when = self.visit.planned_date
        Activity.objects.create(
            activity_type="school_visit",
            school=self.alumni_school,
            project_id=self.alumni.id,
            fy=self.fy,
            planned_date=when,
            planned_month=when.month,
            status="scheduled",
            responsible_staff_id=self.coord.id,
            delivery_type="staff",
            focus_intervention=FOCUS,
        )
        row = self.alumni_row()
        self.assertEqual(row.intervention_label, "General")
        self.assertEqual([line.intervention_label for line in row.work], ["General"])
        self.assertIsNone(row.previous_score)


class _ExportFixture(_ColumnsFixture):
    """A second project, so "all projects" and "one project" differ."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.second = Project.objects.create(
            name="PMS Second Project",
            code="SP-PMS2",
            category="pilot",
            status="active",
            intervention=FOCUS,
            target_interventions=[FOCUS],
            manager_staff_id=cls.coord.id,
        )
        ProjectSchoolAssignment.objects.create(
            project=cls.second,
            school=cls.unplanned.school,
            assigned_by=cls.lead_user.id,
        )

    def rows(self, book) -> list[dict]:
        sheet = book["Project Schools"]
        headings = [cell.value for cell in sheet[1]]
        return [
            dict(zip(headings, [cell.value for cell in line], strict=True))
            for line in sheet.iter_rows(min_row=2)
        ]


class TheExportTest(_ExportFixture):
    def test_all_projects_are_one_file_and_one_sheet(self):
        for user in (self.ia_user, self.cd_user, self.coord_user):
            with self.subTest(role=user.active_role):
                response, book = self.workbook(user, "/projects/export")
                self.assertIn(
                    f"special-projects-schools-FY{self.fy}.xlsx",
                    response["Content-Disposition"],
                )
                rows = self.rows(book)
                self.assertEqual(
                    {row["Project Name"] for row in rows},
                    {"PMS Learning Project", "PMS Second Project"},
                )
                self.assertEqual(len(rows), 9)

    def test_one_project_is_a_file_of_its_own_schools(self):
        for user in (self.ia_user, self.cd_user, self.coord_user):
            with self.subTest(role=user.active_role):
                response, book = self.workbook(
                    user, f"/projects/{self.second.id}/export"
                )
                self.assertIn(
                    "project-sp-pms2-schools", response["Content-Disposition"]
                )
                rows = self.rows(book)
                self.assertEqual(
                    [(row["Project Name"], row["School Name"]) for row in rows],
                    [("PMS Second Project", "Unplanned Primary")],
                )

    def test_the_file_says_what_the_table_says(self):
        _response, book = self.workbook(
            self.ia_user, f"/projects/{self.project.id}/export"
        )
        by_school = {row["School Name"]: row for row in self.rows(book)}

        verified = by_school["Verified Primary"]
        self.assertEqual(verified["School ID"], "PMS-6")
        self.assertEqual(verified["District"], "PMS District")
        self.assertEqual(verified["SSA Intervention"], "Learning Environment")
        self.assertEqual(
            (
                verified["Previous SSA Score"],
                verified["Current SSA Score"],
                verified["SSA Improvement"],
            ),
            (4, 6, 2),
        )
        self.assertEqual(verified["Status"], "Completed")
        self.assertEqual(verified["Activity Status"], "Completed")
        self.assertEqual(verified["Execution"], "Verified · 1/1 done")
        self.assertEqual(
            verified["Activity Date"].date(), self.today - timedelta(days=20)
        )
        self.assertEqual(
            verified["Enrolled On"].date(), self.today - timedelta(days=100)
        )

        coordinated = by_school["Coordinated Primary"]
        self.assertEqual(coordinated["Purpose of Assignment"], "School Visit")
        self.assertEqual(coordinated["Planning Stage"], "Coordinator planned")
        self.assertEqual(coordinated["Planned By"], "PMS Coordinator")
        self.assertEqual(coordinated["Activity Status"], "Scheduled")
        self.assertEqual(coordinated["Activity Cost (UGX)"], 75_000)

        handed = by_school["Handed Primary"]
        self.assertEqual(handed["Planning Stage"], "Awaiting PMS Partner")
        self.assertEqual(handed["Activity Date"], "Awaiting scheduling")
        self.assertIsNone(handed["Planned By"])
        self.assertIsNone(handed["Activity Status"])
        self.assertIsNone(handed["Activity Cost (UGX)"])

        unplanned = by_school["Unplanned Primary"]
        self.assertIsNone(unplanned["Previous SSA Score"])
        self.assertIsNone(unplanned["SSA Improvement"])
        self.assertIsNone(unplanned["Activity Cost (UGX)"])

    def test_the_file_is_the_readers_own_lens(self):
        # A Programme Lead reads the schools they added, on the page and in
        # the file alike.
        _response, book = self.workbook(self.other_lead_user, "/projects/export")
        self.assertEqual(
            [row["School Name"] for row in self.rows(book)], ["Theirs Primary"]
        )

    def test_the_stage_filter_narrows_the_file_as_it_does_the_page(self):
        _response, book = self.workbook(
            self.ia_user, "/projects/export", stage=monitoring.PLAN_PARTNER_AWAITING
        )
        self.assertEqual(
            [row["School Name"] for row in self.rows(book)], ["Handed Primary"]
        )

    def test_a_role_without_export_gets_no_file(self):
        self.client.force_login(self.cceo_user)
        for path in ("/projects/export", f"/projects/{self.project.id}/export"):
            with self.subTest(path=path):
                response = self.client.get(path, {"fy": self.fy})
                self.assertNotEqual(response.status_code, 200)

    def test_a_coordinator_cannot_export_a_project_they_do_not_run(self):
        self.client.force_login(self.stranger_user)
        response = self.client.get(f"/projects/{self.project.id}/export")
        self.assertEqual(response.status_code, 404)


class WhereTheExportsAreTest(_ExportFixture):
    """Two exports in the page's heading: every project in one file, and the
    project whose tab is open."""

    def test_project_monitoring_offers_all_projects_and_the_open_project(self):
        for user in (self.ia_user, self.cd_user, self.coord_user):
            with self.subTest(role=user.active_role):
                self.client.force_login(user)
                body = self.client.get(
                    "/projects/monitoring", {"fy": self.fy}
                ).content.decode()
                self.assertIn(f'href="/projects/export?fy={self.fy}"', body)
                self.assertEqual(body.count("data-project-export-all"), 1)
                self.assertEqual(body.count("data-project-export="), 1)
                self.assertIn(
                    f'href="/projects/{self.project.id}/export?fy={self.fy}"', body
                )

    def test_the_exports_keep_the_page_s_stage_filter(self):
        self.client.force_login(self.ia_user)
        stage = monitoring.PLAN_PARTNER_AWAITING
        body = self.client.get(
            "/projects/monitoring", {"fy": self.fy, "stage": stage}
        ).content.decode()
        self.assertIn(f"/projects/export?fy={self.fy}&amp;stage={stage}", body)
        self.assertIn(
            f"/projects/{self.project.id}/export?fy={self.fy}&amp;stage={stage}", body
        )

    def test_a_reader_without_export_is_offered_none(self):
        self.client.force_login(self.cceo_user)
        ProjectSchoolAssignment.objects.filter(id=self.unplanned.id).update(
            assigned_by=self.cceo_user.id
        )
        body = self.client.get("/projects/monitoring", {"fy": self.fy}).content.decode()
        self.assertIn("data-project-schools", body)
        self.assertNotIn("data-project-export", body)

    def test_planning_oversight_offers_all_projects_and_the_open_project(self):
        self.client.force_login(self.ia_user)
        every = f'href="/team-planning-oversight/projects/export?fy={self.fy}'
        one = f"/team-planning-oversight/projects/{self.second.id}/export?fy={self.fy}"

        # All Projects: one file of every project is the only export.
        body = self.client.get(
            "/team-planning-oversight/", {"view": "projects", "fy": self.fy}
        ).content.decode()
        self.assertIn(every, body)
        self.assertEqual(body.count("data-project-export-all"), 1)
        self.assertNotIn("data-project-export=", body)

        # A project's tab: that project can be exported on its own as well.
        body = self.client.get(
            "/team-planning-oversight/",
            {"view": "projects", "fy": self.fy, "project": self.second.id},
        ).content.decode()
        self.assertIn(every, body)
        self.assertIn(one, body)
        self.assertEqual(body.count("data-project-export="), 1)

    def test_the_coordinators_projects_page_offers_all_projects(self):
        self.client.force_login(self.coord_user)
        body = self.client.get("/projects").content.decode()
        self.assertIn('href="/projects/export"', body)
        self.assertIn(f'data-project-export="{self.project.id}"', body)


class TheColumnsCostNoQueryPerRowTest(_ColumnsFixture):
    def test_more_scheduled_schools_cost_no_more_queries(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        def count() -> int:
            with CaptureQueriesContext(connection) as queries:
                monitoring.project_monitoring(self.ia_user, fy=self.fy)
            return len(queries)

        before = count()
        when = self.visit.planned_date
        for index in range(6):
            enrolment = self._enrol(f"PMS-X{index}", f"Extra Primary {index}")
            visit = self._work(enrolment.school, "school_visit", when, self.coord.id)
            _cost(visit, 10_000)
            cancelled = self._work(
                enrolment.school,
                "school_visit",
                when,
                self.coord.id,
                status="cancelled",
            )
            _cost(cancelled, 10_000)
        self.assertEqual(count(), before)
