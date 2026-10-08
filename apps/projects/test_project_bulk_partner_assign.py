"""Any project's schools can be handed to a partner, several at a time.

Owner, 2026-10-05, on "Could not assign the selection: Select a Catalogue
Activity eligible for every selected Project School": "lift the restriction on
projects that lacks any partner approved activities so that all projects
schools added to a project can be assigned to any partner", and "use
checkboxes on the left to check all the project schools to be assigned to the
partner and the button for assign pops up just like the way bulk assign to
partner by cceo and pl is designed".

The drawer used to offer only a Catalogue Activity the project approved, that
a partner may deliver to a single school and that every selected school's SSA
called for. A project approving none — or approving group trainings such as
Leadership — offered nothing, and the save was refused. What is handed over is
now the planner's choice, as on every other partner drawer: the purpose and,
for an In-school Training, the training.

Asked next whether a project hand-over should go past the earlier rules that
still left schools out (Champion schools, the Core partner half, the
two-partner limit, a partner's recorded trainings, the reason for a partner
outside the project's list), the owner said "yes for now and apply to the
staff(project coordinator)": those rules sit behind one switch
(apps.partners.handover_policy), and the coordinator's own bulk Schedule asks
for a purpose and a training too.
"""

from __future__ import annotations

import io
import json
from unittest import mock

from django.test import TestCase
from openpyxl import load_workbook

from apps.accounts.models import StaffProfile, User
from apps.activity_catalogue.models import ActivityCatalogueItem, ActivityProjectMapping
from apps.core.enums import SsaIntervention
from apps.core.fy import get_operational_fy
from apps.core.rbac import EdifyRole
from apps.geography.models import District, Region
from apps.partners.models import Partner, PartnerAssignment
from apps.projects.models import Project, ProjectSchoolAssignment
from apps.schools.models import School

DRAWER = "/projects/planning/bulk-partner"
SCHEDULE = "/projects/planning/bulk-schedule"

#: The school rules a project hand-over goes past for now, switched back on
#: for the tests that hold them for the day they return.
KEEP_SCHOOL_RULES = mock.patch(
    "apps.partners.handover_policy.PROJECT_HANDOVERS_KEEP_SCHOOL_RULES", True
)


class _Fixture(TestCase):
    """A coordinated project that approves no Catalogue Activity, with three
    client schools in two districts and one Champion school."""

    @classmethod
    def setUpTestData(cls):
        from apps.activity_catalogue.seeding import seed_activity_catalogue

        seed_activity_catalogue(actor_id="test")
        cls.fy = get_operational_fy()
        cls.coord_user = User.objects.create_user(
            email="pba-coord@edify.org",
            name="PBA Coordinator",
            roles=[EdifyRole.PROJECT_COORDINATOR.value],
            active_role=EdifyRole.PROJECT_COORDINATOR.value,
            password="x",
            is_active=True,
        )
        cls.coord = StaffProfile.objects.create(
            user=cls.coord_user, title="Coordinator", country="Uganda"
        )
        cls.ia_user = User.objects.create_user(
            email="pba-ia@edify.org",
            name="PBA Impact",
            roles=[EdifyRole.IMPACT_ASSESSMENT.value],
            active_role=EdifyRole.IMPACT_ASSESSMENT.value,
            password="x",
            is_active=True,
        )
        StaffProfile.objects.create(user=cls.ia_user, title="IA", country="Uganda")
        region = Region.objects.create(name="PBA Region")
        cls.north = District.objects.create(name="PBA North", region=region)
        cls.south = District.objects.create(name="PBA South", region=region)

        def school(code, name, district, school_type="client"):
            return School.objects.create(
                school_id=code,
                name=name,
                region=region,
                district=district,
                school_type=school_type,
                planning_readiness="ready_for_support_planning",
            )

        cls.schools = [
            school("PBA-1", "Acacia Primary", cls.north),
            school("PBA-2", "Baobab Primary", cls.north),
            school("PBA-3", "Cedar Primary", cls.south),
        ]
        cls.champion = school("PBA-4", "Date Palm Primary", cls.south, "champion")
        cls.project = Project.objects.create(
            name="PBA Leadership Project",
            category="pilot",
            status="active",
            intervention=SsaIntervention.LEADERSHIP,
            target_interventions=[SsaIntervention.LEADERSHIP],
            manager_staff_id=cls.coord.id,
        )
        cls.enrolments = [
            ProjectSchoolAssignment.objects.create(project=cls.project, school=s)
            for s in cls.schools
        ]
        cls.champion_enrolment = ProjectSchoolAssignment.objects.create(
            project=cls.project, school=cls.champion
        )
        cls.partner = Partner.objects.create(name="PBA Partner", active_status=True)

    def setUp(self):
        self.client.force_login(self.coord_user)

    def ids(self, *enrolments):
        return ",".join(e.id for e in (enrolments or self.enrolments))

    def drawer(self, *enrolments):
        response = self.client.get(DRAWER, {"assignments": self.ids(*enrolments)})
        self.assertEqual(response.status_code, 200)
        return response

    def hand_over(self, *enrolments, **fields):
        return self.client.post(
            DRAWER,
            {
                "assignments": self.ids(*enrolments),
                "partner_id": self.partner.id,
                **fields,
            },
        )

    def courses(self, *enrolments):
        return json.loads(self.drawer(*enrolments).context["training_courses_json"])


class AProjectWithNoApprovedActivityTest(_Fixture):
    def test_the_project_approves_nothing(self):
        self.assertFalse(ActivityProjectMapping.objects.filter(project=self.project))

    def test_the_drawer_asks_for_a_purpose_and_refuses_nothing(self):
        response = self.drawer()
        self.assertContains(response, "Assign 3 project schools to a partner")
        self.assertContains(response, "Purpose of assignment")
        self.assertNotContains(response, "eligible for every selected")
        self.assertNotContains(response, 'name="catalogue_item_id"')
        self.assertEqual(len(response.context["assignments"]), 3)

    def test_every_ticked_school_is_handed_over_for_the_purpose(self):
        response = self.hand_over(purpose_of_visit="ssa_support")
        self.assertEqual(response.status_code, 200, response.content[:400])
        self.assertContains(response, "Assigned 3 project schools to PBA Partner.")
        handovers = PartnerAssignment.objects.filter(project=self.project)
        self.assertEqual(
            sorted(h.school.school_id for h in handovers), ["PBA-1", "PBA-2", "PBA-3"]
        )
        for handover in handovers:
            with self.subTest(school=handover.school.school_id):
                self.assertEqual(handover.partner_id, self.partner.id)
                self.assertEqual(handover.purpose_of_visit, "ssa_support")
                self.assertTrue(handover.catalogue_item.standard_support)
                self.assertEqual(
                    handover.status, PartnerAssignment.STATUS_PENDING_SCHEDULING
                )
                # SSA Support moves no single intervention.
                self.assertFalse(handover.focus_intervention)

    def test_the_date_is_the_partner_s_to_choose(self):
        """Owner: the project is assigned "to the partner to schedule"."""
        # The hand-over shows the day it is made and asks for no date
        # (owner, 2026-10-05): a date typed here read as the partner's plan.
        self.assertContains(self.drawer(), "Assigned date")
        self.assertNotContains(self.drawer(), "Target date")
        self.assertNotContains(self.drawer(), 'name="scheduled_date"')
        response = self.hand_over(purpose_of_visit="ssa_support")
        self.assertEqual(response.status_code, 200, response.content[:400])
        self.assertFalse(
            PartnerAssignment.objects.filter(
                project=self.project, scheduled_date__isnull=False
            )
        )

    def test_a_training_is_any_of_the_governed_courses(self):
        courses = self.courses()
        self.assertGreater(len(courses), 10)
        # No project course to lead the list: the project approves none.
        self.assertFalse([c for c in courses if c["projectCourse"]])
        course = courses[0]
        response = self.hand_over(
            purpose_of_visit="in_school_training", training_course_id=course["id"]
        )
        self.assertEqual(response.status_code, 200, response.content[:400])
        handover = PartnerAssignment.objects.filter(project=self.project).first()
        # The course is named; the standard In-school Training workflow
        # delivers and prices it, as on every other partner hand-over.
        self.assertEqual(handover.training_course_id, course["id"])
        self.assertEqual(handover.catalogue_item.workflow_kind, "in_school_training")
        self.assertTrue(handover.catalogue_item.standard_support)
        self.assertEqual(
            handover.focus_intervention,
            course["ssaIntervention"] or SsaIntervention.LEADERSHIP,
        )

    def test_a_training_names_its_course(self):
        response = self.hand_over(purpose_of_visit="in_school_training")
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Select the Training", status_code=400)
        self.assertFalse(PartnerAssignment.objects.filter(project=self.project))

    def test_a_purpose_is_asked_for(self):
        response = self.hand_over()
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Select the reason", status_code=400)

    def test_for_now_a_champion_school_is_handed_over_too(self):
        """Owner, 2026-10-05: "all projects schools added to a project can be
        assigned to any partner" — "yes for now"."""
        everyone = (*self.enrolments, self.champion_enrolment)
        response = self.drawer(*everyone)
        self.assertContains(response, "Assign 4 project schools to a partner")
        self.assertNotContains(response, "data-handover-not-included")

        saved = self.hand_over(*everyone, purpose_of_visit="ssa_support")
        self.assertEqual(saved.status_code, 200, saved.content[:400])
        self.assertContains(saved, "Assigned 4 project schools")
        self.assertTrue(
            PartnerAssignment.objects.filter(school=self.champion, project=self.project)
        )

    @KEEP_SCHOOL_RULES
    def test_a_champion_school_is_named_and_left_out(self):
        """Staff deliver a Champion school's work themselves (owner,
        2026-09-21); with that rule on, the drawer says so before the save
        rather than after."""
        everyone = (*self.enrolments, self.champion_enrolment)
        response = self.drawer(*everyone)
        self.assertContains(response, "Assign 3 project schools to a partner")
        self.assertContains(response, "data-handover-not-included")
        self.assertContains(response, "Date Palm Primary is a Champion school")

        saved = self.hand_over(*everyone, purpose_of_visit="ssa_support")
        self.assertEqual(saved.status_code, 200, saved.content[:400])
        self.assertContains(saved, "Assigned 3 project schools")
        self.assertFalse(PartnerAssignment.objects.filter(school=self.champion))

    @KEEP_SCHOOL_RULES
    def test_a_champion_school_alone_has_nothing_to_assign(self):
        response = self.drawer(self.champion_enrolment)
        self.assertContains(response, "Date Palm Primary is a Champion school")
        self.assertNotContains(response, 'type="submit"')
        saved = self.hand_over(self.champion_enrolment, purpose_of_visit="ssa_support")
        self.assertEqual(saved.status_code, 400)
        self.assertFalse(PartnerAssignment.objects.filter(school=self.champion))


class TheProjectsOwnTrainingsLeadTest(_Fixture):
    """What a project approves no longer limits the choice; it leads it."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        # A group training (never offered before: not delivered to a single
        # school) and a project activity that is not a governed course.
        cls.leadership = ActivityCatalogueItem.objects.get(
            stable_code="SCHOOL_LEADERSHIP"
        )
        cls.cc_sel = ActivityCatalogueItem.objects.get(stable_code="CC_SEL")
        for item in (cls.leadership, cls.cc_sel):
            ActivityProjectMapping.objects.create(
                project=cls.project, catalogue_item=item
            )

    def test_the_project_s_trainings_come_first(self):
        courses = self.courses()
        own = [c["stableCode"] for c in courses if c["projectCourse"]]
        self.assertEqual(sorted(own), ["CC_SEL", "SCHOOL_LEADERSHIP"])
        self.assertEqual(
            sorted(c["stableCode"] for c in courses[: len(own)]), sorted(own)
        )
        self.assertGreater(len(courses), 10, "every other course is still offered")
        self.assertContains(self.drawer(), "This project's trainings")

    def test_a_group_training_the_project_approves_can_be_handed_over(self):
        self.assertFalse(self.leadership.individual_school_allowed)
        response = self.hand_over(
            purpose_of_visit="in_school_training",
            training_course_id=self.leadership.id,
        )
        self.assertEqual(response.status_code, 200, response.content[:400])
        handover = PartnerAssignment.objects.filter(project=self.project).first()
        self.assertEqual(handover.training_course_id, self.leadership.id)
        self.assertEqual(handover.catalogue_item.workflow_kind, "in_school_training")

    def test_a_project_activity_outside_the_courses_is_handed_over_as_itself(self):
        self.assertFalse(self.cc_sel.is_training_course)
        response = self.hand_over(
            purpose_of_visit="in_school_training", training_course_id=self.cc_sel.id
        )
        self.assertEqual(response.status_code, 200, response.content[:400])
        handovers = PartnerAssignment.objects.filter(project=self.project)
        self.assertEqual(handovers.count(), 3)
        for handover in handovers:
            self.assertEqual(handover.catalogue_item_id, self.cc_sel.id)
            self.assertIsNone(handover.training_course_id)
            self.assertEqual(
                handover.focus_intervention, SsaIntervention.CHRISTLIKE_BEHAVIOUR
            )

    def test_another_project_s_activity_is_not_handed_over_as_itself(self):
        """Outside the governed courses, only what this project approves."""
        other = ActivityCatalogueItem.objects.get(
            stable_code="LEARNER_CENTERED_APPROACHES"
        )
        response = self.hand_over(
            purpose_of_visit="in_school_training", training_course_id=other.id
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(PartnerAssignment.objects.filter(project=self.project))


class AProjectHandOverGoesPastTheSchoolRulesForNowTest(_Fixture):
    """Each rule is asked at the one creation door (and the two-partner limit
    at the model); a project's hand-over goes past them while the switch is
    off, and nothing changes for a hand-over that names no project."""

    RULES = (
        "apps.partners.services._assert_school_takes_partner_work",
        "apps.partners.services._assert_partner_half_open",
        "apps.partners.services._assert_project_partner",
        "apps.partners.services._assert_partner_delivers",
    )

    def _create(self, **fields):
        from apps.partners import services as partner_services

        return partner_services.create_assignment(
            school=self.schools[0],
            partner=self.partner,
            assigning_staff_id=self.coord.id,
            expected_activity_type="school_visit",
            **fields,
        )

    def test_a_project_hand_over_is_asked_none_of_them(self):
        from apps.core.exceptions import BadRequest

        for rule in self.RULES:
            with (
                self.subTest(rule=rule),
                mock.patch(rule, side_effect=BadRequest("no")),
            ):
                handover = self._create(project=self.project)
                self.assertEqual(handover.project_id, self.project.id)
                handover.delete()

    def test_a_hand_over_outside_a_project_is_asked_every_one(self):
        from apps.core.exceptions import BadRequest

        for rule in self.RULES:
            with (
                self.subTest(rule=rule),
                mock.patch(rule, side_effect=BadRequest("no")),
            ):
                with self.assertRaisesMessage(BadRequest, "no"):
                    self._create()

    @KEEP_SCHOOL_RULES
    def test_the_switch_puts_them_back_on_project_hand_overs(self):
        from apps.core.exceptions import BadRequest

        for rule in self.RULES:
            with (
                self.subTest(rule=rule),
                mock.patch(rule, side_effect=BadRequest("no")),
            ):
                with self.assertRaisesMessage(BadRequest, "no"):
                    self._create(project=self.project)

    def test_a_third_partner_is_held_outside_a_project_only(self):
        from apps.core.exceptions import ConflictError

        limit = (
            "apps.partners.support_responsibility.assert_school_accepts_another_partner"
        )
        # At a client school one hand-over a year answers before this limit
        # does (owner, 2026-10-08; `_assert_school_has_room`). It is held
        # aside here so the limit itself, which a Core school still reaches,
        # is what is asked.
        one_visit = mock.patch("apps.partners.services._assert_school_has_room")
        with one_visit, mock.patch(limit, side_effect=ConflictError("two already")):
            self.assertTrue(self._create(project=self.project).id)
            with self.assertRaisesMessage(ConflictError, "two already"):
                self._create()
        with KEEP_SCHOOL_RULES, mock.patch(limit, side_effect=ConflictError("two")):
            PartnerAssignment.objects.all().delete()
            with self.assertRaisesMessage(ConflictError, "two"):
                self._create(project=self.project)

    def test_a_closed_school_takes_no_hand_over_either_way(self):
        from apps.core.exceptions import BadRequest

        with mock.patch(
            "apps.schools.lifecycle_service.assert_operating",
            side_effect=BadRequest("closed"),
        ):
            with self.assertRaisesMessage(BadRequest, "closed"):
                self._create(project=self.project)

    def test_a_partner_with_recorded_trainings_is_offered_every_training(self):
        """A partner's recorded trainings narrow the list only while the rule
        is on; for now a project's work goes to any partner."""
        self.partner.activity_codes = ["SCHOOL_LEADERSHIP"]
        self.partner.save(update_fields=["activity_codes"])
        self.assertEqual(self.drawer().context["partner_courses_json"], "{}")
        other = next(
            c for c in self.courses() if c["stableCode"] != "SCHOOL_LEADERSHIP"
        )
        response = self.hand_over(
            purpose_of_visit="in_school_training", training_course_id=other["id"]
        )
        self.assertEqual(response.status_code, 200, response.content[:400])
        self.assertEqual(
            PartnerAssignment.objects.filter(project=self.project).count(), 3
        )
        with KEEP_SCHOOL_RULES:
            self.assertIn(
                self.partner.id,
                json.loads(self.drawer().context["partner_courses_json"]),
            )


class ThePartnerDatesWhatItWasHandedTest(_Fixture):
    """A lift at the hand-over alone would be a trap: the partner's own
    scheduling door asked the same Champion rule, so the school was handed
    over and could never be dated."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.partner_user = User.objects.create(
            id="pba-partner-user",
            email="pba-partner@edify.org",
            name="PBA Partner User",
            roles=["PartnerFieldOfficer"],
            active_role="PartnerFieldOfficer",
            is_active=True,
        )
        cls.partner.user_id = cls.partner_user.id
        cls.partner.save(update_fields=["user_id"])

    def setUp(self):
        super().setUp()
        self.handover = PartnerAssignment.objects.create(
            school=self.champion,
            partner=self.partner,
            project=self.project,
            assigning_staff_id=self.coord.id,
            monitoring_staff_id=self.coord.id,
            expected_activity_type="school_visit",
            purpose_of_visit="training_follow_up",
            status=PartnerAssignment.STATUS_PENDING_SCHEDULING,
        )

    def date_it(self):
        from datetime import date, timedelta

        from apps.activities.services import _partner_schedule_from_assignment
        from apps.planning.test_visit_gate import _schedulable

        when = _schedulable(date.today() + timedelta(days=10))
        return _partner_schedule_from_assignment(
            self.handover.id,
            {"scheduledDate": when.isoformat(), "deliveryContactName": "PBA Visitor"},
            self.partner_user,
        )

    def test_the_partner_dates_a_project_hand_over_at_a_champion_school(self):
        created = self.date_it()
        self.assertEqual(created["deliveryType"], "partner")
        self.handover.refresh_from_db()
        self.assertEqual(self.handover.scheduled_activity_id, created["id"])

    @KEEP_SCHOOL_RULES
    def test_with_the_rules_on_the_champion_school_is_refused_there_too(self):
        from apps.core.exceptions import BadRequest

        with self.assertRaisesMessage(BadRequest, "Champion school"):
            self.date_it()

    def test_the_partner_s_queue_offers_schedule_for_it(self):
        self.client.force_login(self.partner_user)
        response = self.client.get("/partner/assigned-schools")
        self.assertEqual(response.status_code, 200, response.get("Location"))
        self.assertContains(response, "Date Palm Primary")
        self.assertNotContains(response, 'data-visit-locked="true"')
        with KEEP_SCHOOL_RULES:
            locked = self.client.get("/partner/assigned-schools")
        self.assertContains(locked, 'data-visit-locked="true"')


class TheCoordinatorSchedulesInBulkTest(_Fixture):
    """Owner, 2026-10-05: "apply to the staff(project coordinator)". The
    coordinator's bulk Schedule asked for the same project-approved Catalogue
    Activity and offered this project nothing."""

    def schedule(self, *enrolments, **fields):
        from datetime import timedelta

        from django.utils import timezone

        return self.client.post(
            SCHEDULE,
            {
                "assignments": self.ids(*enrolments),
                "scheduled_date": (
                    timezone.localdate() + timedelta(days=7)
                ).isoformat(),
                **fields,
            },
        )

    def test_the_drawer_asks_for_a_purpose_and_refuses_nothing(self):
        response = self.client.get(SCHEDULE, {"assignments": self.ids()})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Schedule 3 project schools")
        self.assertContains(response, "Purpose of visit")
        self.assertNotContains(response, "eligible for every selected")
        self.assertNotContains(response, 'name="catalogue_item_id"')
        courses = json.loads(response.context["training_courses_json"])
        self.assertGreater(len(courses), 10)

    def test_a_purpose_is_costed_by_its_own_catalogue_activity(self):
        with mock.patch(
            "apps.frontend.views.planning_views.schedule_school_visit",
            return_value={"id": "x"},
        ) as scheduled:
            response = self.schedule(purpose_of_visit="ssa_support")
        self.assertEqual(response.status_code, 200, response.content[:400])
        self.assertContains(response, "Scheduled 3 project schools.")
        payloads = [call.args[0] for call in scheduled.call_args_list]
        self.assertEqual(
            sorted(p["schoolId"] for p in payloads), ["PBA-1", "PBA-2", "PBA-3"]
        )
        item = ActivityCatalogueItem.objects.get(id=payloads[0]["catalogueItemId"])
        self.assertEqual(item.workflow_kind, "school_visit_ssa_collection")
        for payload in payloads:
            self.assertEqual(payload["projectId"], self.project.id)
            self.assertEqual(payload["purposeType"], "ssa_support")
            self.assertEqual(payload["deliveryType"], "staff")
            # SSA Support moves no single intervention.
            self.assertNotIn("focusIntervention", payload)

    def test_a_training_is_the_course_and_its_school_visit_together(self):
        response = self.client.get(SCHEDULE, {"assignments": self.ids()})
        course = json.loads(response.context["training_courses_json"])[0]
        with mock.patch(
            "apps.frontend.views.planning_views.schedule_in_school_training_pair",
            return_value={"id": "x"},
        ) as paired:
            saved = self.schedule(
                purpose_of_visit="in_school_training", training_course_id=course["id"]
            )
        self.assertEqual(saved.status_code, 200, saved.content[:400])
        self.assertEqual(paired.call_count, 3)
        payload = paired.call_args_list[0].args[0]
        self.assertEqual(payload["catalogueItemId"], course["id"])
        self.assertEqual(payload["purposeType"], "in_school_training")
        self.assertEqual(payload["projectId"], self.project.id)

    def test_a_project_activity_outside_the_courses_is_scheduled_as_itself(self):
        cc_sel = ActivityCatalogueItem.objects.get(stable_code="CC_SEL")
        ActivityProjectMapping.objects.create(
            project=self.project, catalogue_item=cc_sel
        )
        with mock.patch(
            "apps.frontend.views.planning_views.schedule_school_visit",
            return_value={"id": "x"},
        ) as scheduled:
            saved = self.schedule(
                purpose_of_visit="in_school_training", training_course_id=cc_sel.id
            )
        self.assertEqual(saved.status_code, 200, saved.content[:400])
        payload = scheduled.call_args_list[0].args[0]
        self.assertEqual(payload["catalogueItemId"], cc_sel.id)
        self.assertEqual(
            payload["focusIntervention"], SsaIntervention.CHRISTLIKE_BEHAVIOUR
        )
        # Scheduled as the Catalogue Activity it is, as before.
        self.assertNotIn("purposeType", payload)

    def test_a_purpose_and_a_training_are_asked_for(self):
        no_purpose = self.schedule()
        self.assertEqual(no_purpose.status_code, 400)
        self.assertContains(no_purpose, "Select the purpose", status_code=400)
        no_course = self.schedule(purpose_of_visit="in_school_training")
        self.assertEqual(no_course.status_code, 400)
        self.assertContains(no_course, "Select the Training", status_code=400)

    def test_a_school_the_rules_refuse_is_named_and_the_rest_scheduled(self):
        """The coordinator's own scheduling still follows the staff-side
        rules; a refused school does not cost the selection."""
        from apps.core.exceptions import BadRequest

        def schedule(payload, principal):
            if payload["schoolId"] == "PBA-2":
                raise BadRequest("Baobab Primary has had its visit.")
            return {"id": "x"}

        with mock.patch(
            "apps.frontend.views.planning_views.schedule_school_visit",
            side_effect=schedule,
        ):
            saved = self.schedule(purpose_of_visit="donor_visit")
        self.assertEqual(saved.status_code, 200, saved.content[:400])
        self.assertContains(saved, "Scheduled 2 project schools.")
        self.assertContains(saved, "Baobab Primary has had its visit.")


class TickTheSchoolsOnProjectMonitoringTest(_Fixture):
    def page(self, user=None, **query):
        self.client.force_login(user or self.coord_user)
        response = self.client.get("/projects/monitoring", {"fy": self.fy, **query})
        self.assertEqual(response.status_code, 200)
        return response

    def test_the_coordinator_ticks_schools_and_the_assign_bar_opens_the_drawer(self):
        html = self.page().content.decode()
        self.assertIn("data-project-pick-all", html)
        for enrolment in self.enrolments:
            self.assertIn(f'value="{enrolment.id}" x-model="selected"', html)
        # The bar opens the project hand-over for what is ticked.
        self.assertIn("data-project-bulk-bar", html)
        self.assertIn(
            'hx-get="/projects/planning/bulk-partner" hx-target="#drawer-container"',
            html,
        )
        self.assertIn('name="assignments" :value="selected.join', html)
        self.assertIn("data-project-assign-selected", html)
        # ...or the coordinator's own bulk Schedule for the same schools.
        self.assertIn(
            'hx-get="/projects/planning/bulk-schedule" hx-target="#drawer-container"',
            html,
        )
        self.assertIn("data-project-schedule-selected", html)
        # No header button: the bar comes up once a school is ticked.
        self.assertNotIn("Assign project to a partner", html)

    def test_for_now_every_open_school_has_a_box(self):
        html = self.page().content.decode()
        self.assertIn(f'value="{self.champion_enrolment.id}" x-model="selected"', html)

    @KEEP_SCHOOL_RULES
    def test_a_school_no_partner_may_take_has_no_box(self):
        html = self.page().content.decode()
        self.assertNotIn(f'value="{self.champion_enrolment.id}" x-model', html)
        # Its row still opens the drawer, which says why.
        self.assertIn(
            f"/projects/planning/bulk-partner?assignments={self.champion_enrolment.id}",
            html,
        )

    def test_a_reader_gets_no_boxes_and_no_bar(self):
        html = self.page(self.ia_user).content.decode()
        self.assertIn("Acacia Primary", html)
        for mark in ("data-project-pick", "data-project-bulk-bar", DRAWER):
            self.assertNotIn(mark, html)

    def test_the_page_reads_itself_again_after_a_hand_over(self):
        """The drawer saves without leaving the page; a school just handed
        over must stop reading "Not planned"."""
        html = self.page().content.decode()
        self.assertIn("data-project-monitoring-refresh", html)
        self.assertIn(
            f'hx-get="/projects/monitoring?project={self.project.id}&amp;stage='
            f'&amp;fy={self.fy}" hx-trigger="planning-saved from:body"',
            html,
        )
        self.assertIn('hx-select="#project-monitoring-root"', html)


class TheFilterRowTest(_Fixture):
    def page(self, **query):
        response = self.client.get("/projects/monitoring", {"fy": self.fy, **query})
        self.assertEqual(response.status_code, 200)
        return response

    def schools_on(self, response):
        (row,) = response.context["rows"]
        return [school.school_name for school in row.school_rows]

    def test_the_filters_sit_above_the_figures_before_the_tab_pill(self):
        """Owner, 2026-10-05: "move the filters above the KPI strip next to
        Monitor and analytic tabs. Start with the filters then the tab pill"."""
        html = self.page().content.decode()
        controls = html.index("data-project-monitoring-controls")
        form = html.index('class="edify-filter-bar', controls)
        pill = html.index('class="edify-section-nav"', controls)
        figures = html.index("Live projects in this lens")
        self.assertLess(controls, form)
        self.assertLess(form, pill)
        self.assertLess(pill, figures)

    def test_the_open_project_s_activity_figures_are_in_the_strip(self):
        """Owner, 2026-10-05, of the four figures that stood inside the
        project's card: "can you move that to the KPI Strip. it looks ugly
        there"."""
        html = self.page().content.decode()
        card = html.index('data-project-monitoring="')
        for label in ("Trainings", "School visits", "Other project work"):
            with self.subTest(figure=label):
                self.assertIn(label, html[:card])
                self.assertNotIn(label, html[card:])
        self.assertIn("Who delivered", html[:card])
        self.assertNotIn("Who delivered", html[card:])
        # Each says which project it counts; the figures before it count all.
        self.assertIn("Delivered of scheduled · PBA Leadership Project", html[:card])

    def test_district_leads_the_filters_the_table_s_columns_follow(self):
        self.assertEqual(
            [f["key"] for f in self.page().context["monitor_fields"]],
            ["district", "stage", "fy", "partner", "training", "purpose", "activity"],
        )

    def test_a_filter_offers_what_the_schools_hold(self):
        filters = {f["key"]: f for f in self.page().context["monitor_fields"]}
        self.assertEqual(
            [value for value, _label in filters["district"]["options"]],
            ["", "PBA North", "PBA South"],
        )
        # Nothing is with a partner yet, so there is no partner to pick.
        self.assertEqual(filters["partner"]["options"], [("", "All partners")])

    def test_district_narrows_the_schools_and_not_the_project(self):
        response = self.page(district="PBA North")
        self.assertEqual(
            self.schools_on(response), ["Acacia Primary", "Baobab Primary"]
        )
        (row,) = response.context["rows"]
        self.assertEqual(
            len(row.all_school_rows), 4, "the project's figures stay whole"
        )
        (tab,) = response.context["project_tabs"]
        self.assertEqual(tab["count"], 2)

    def test_partner_and_purpose_filter_what_was_handed_over(self):
        self.hand_over(self.enrolments[0], purpose_of_visit="ssa_support")
        filters = {f["key"]: f for f in self.page().context["monitor_fields"]}
        self.assertIn(("PBA Partner", "PBA Partner"), filters["partner"]["options"])
        self.assertIn(("SSA Support", "SSA Support"), filters["purpose"]["options"])
        self.assertEqual(
            self.schools_on(self.page(partner="PBA Partner")), ["Acacia Primary"]
        )
        self.assertEqual(
            self.schools_on(self.page(purpose="SSA Support", district="PBA South")), []
        )
        self.assertContains(
            self.page(purpose="SSA Support", district="PBA South"),
            "No school in this project matches these filters.",
        )

    def test_the_filters_travel_with_the_tabs_the_exports_and_clear(self):
        response = self.page(district="PBA North")
        html = response.content.decode()
        self.assertIn(
            f"?project={self.project.id}&stage=&fy={self.fy}&district=PBA+North", html
        )
        self.assertIn(f"/projects/export?fy={self.fy}&amp;district=PBA+North", html)
        self.assertEqual(
            response.context["filter_reset_url"],
            f"/projects/monitoring?project={self.project.id}&fy={self.fy}",
        )
        self.assertEqual(self.page().context["filter_reset_url"], "")

    def test_the_export_keeps_the_filter(self):
        response = self.client.get(
            "/projects/export", {"fy": self.fy, "district": "PBA North"}
        )
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        names = [row[2].value for row in sheet.iter_rows(min_row=2) if row[2].value]
        self.assertEqual(names, ["Acacia Primary", "Baobab Primary"])
