"""My Plan's export: its tables, a sheet each (owner, 2026-10-26: "tables
are exported exactly the way they are in My Plan tables"), and the budget for
those activities in the last sheet (owner, 2026-10-12: "the budget for those
selected activities in the last sheet and the budget should be fetched direct
from those planned activities").
"""

from __future__ import annotations

import csv
import io

from django.test import Client

from apps.activities.models import Activity, ActivityScheduleCostLine
from apps.core import plan_export
from apps.frontend.test_oversight_filters_export import FilterExportFixture
from apps.my_plan import table_export


def _workbook(response):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(response.content), read_only=True)


class PlanExportSheetsTest(FilterExportFixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.meeting = cls._activity("cluster_meeting", cls.james, cost=15_000)

    def my_plan(self, export: str):
        client = Client()
        client.force_login(self.james_user)
        response = client.get("/my-plan/", {"fy": self.fy, "export": export})
        self.assertEqual(response.status_code, 200)
        return response

    def my_plan_lines(self) -> list[list[str]]:
        body = self.my_plan("csv").content.decode()
        return list(csv.reader(io.StringIO(body)))

    def test_the_sheets_are_my_plans_tables_and_the_budget_last(self):
        book = _workbook(self.my_plan("xlsx"))
        self.assertEqual(book.sheetnames, [*table_export.FIXED_TITLES, "Budget"])

    def test_each_kind_of_work_is_on_its_own_sheet(self):
        tables = table_export.tables_from_csv(self.my_plan_lines())
        self.assertEqual(
            [(row["School ID"], row["Activity"]) for row in tables["School Visits"]],
            [("s1", "School Visit")],
        )
        self.assertEqual([row["School ID"] for row in tables["Trainings"]], ["s1"])
        self.assertEqual(tables["Group Trainings"], [])
        self.assertEqual(len(tables["Cluster Meetings"]), 1)
        # The partner's work is not the officer's plan.
        names = {row.get("Executor") for rows in tables.values() for row in rows}
        self.assertNotIn("Partner X", names)

    def test_the_budget_sheet_is_those_activities_own_cost_lines(self):
        book = _workbook(self.my_plan("xlsx"))
        budget = [list(r) for r in book["Budget"].iter_rows(values_only=True)]
        exported = Activity.objects.filter(
            id__in=[self.visit.id, self.training.id, self.meeting.id]
        )
        expected = sum(
            line.amount
            for line in ActivityScheduleCostLine.objects.filter(activity__in=exported)
        )
        self.assertEqual(expected, 115_000)
        # The Budget page's ledger: Grand total on the first row and the last.
        self.assertEqual(budget[0][4], "Grand total")
        self.assertEqual(budget[0][5], expected)
        self.assertEqual(budget[-1][5], expected)
        self.assertNotIn(99_000, [cell for row in budget for cell in row])
        # Exactly the Budget export's Activity cost plan (owner, 2026-10-12:
        # "use the actual budget format and use that exactly"): the same
        # sheet `cost_plan_sheet` gives the Budget page's export for these
        # lines, row for row, styles included.
        from apps.accounts.models import User
        from apps.budget.services import (
            _EXPORT_ACTIVITY_FIELDS,
            _EXPORT_LINE_FIELDS,
            COST_PLAN_CAPTION,
            budget_groups,
            cost_line_rows,
            cost_plan_sheet,
        )

        lines = cost_line_rows(
            ActivityScheduleCostLine.objects.filter(activity__in=exported).order_by(
                "planned_date", "activity_id", "id"
            ),
            _EXPORT_LINE_FIELDS,
            _EXPORT_ACTIVITY_FIELDS,
        )
        names = dict(User.objects.values_list("id", "name"))
        day = self.visit.planned_date
        page = cost_plan_sheet(
            budget_groups(lines, names),
            period_label=f"Period: {day:%-d %b %Y} – {day:%-d %b %Y}",
        )
        blank = lambda rows: [[c for c in r] + [None] * (6 - len(r)) for r in rows]  # noqa: E731
        self.assertEqual(blank(budget), blank(page["rows"]))
        self.assertEqual(budget[0][0], "Activity cost plan")
        self.assertEqual(budget[1][0], COST_PLAN_CAPTION)
        made = plan_export.budget_sheet(
            [{"activity_id": a.id, "planned_date": day} for a in exported]
        )
        for key in (
            "row_styles",
            "widths",
            "number_formats",
            "align",
            "filter",
            "freeze",
        ):
            self.assertEqual(made[key], page[key], key)

    def test_the_csv_carries_the_same_sheets_and_no_budget_table(self):
        lines = self.my_plan_lines()
        tables = table_export.tables_from_csv(lines)
        self.assertEqual(list(tables), list(table_export.FIXED_TITLES))
        self.assertNotIn("Budget", tables)
        for rows in tables.values():
            for row in rows:
                self.assertIn("Status", row)
