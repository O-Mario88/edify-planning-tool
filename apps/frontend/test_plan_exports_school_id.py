"""Every exported plan carries the School ID (owner, 2026-09-27).

"All exported plan should have school ID." The School ID is the school's
business code — what the School Directory, Salesforce and the field call the
school — never the platform's own primary key. Work that is at no single
school (a cluster session read as a whole, a programme event) leaves it blank.
"""

from __future__ import annotations

import csv
import io
from datetime import date, timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone

from apps.accounts.models import StaffProfile
from apps.activities.models import Activity
from apps.core.fy import get_operational_fy
from apps.core.rbac import EdifyRole
from apps.core_schools.models import CorePlan, cplan_id
from apps.core_schools.services import create_package_slots
from apps.geography.models import District, Region
from apps.partners.models import Partner, PartnerAssignment
from apps.schools.models import School

User = get_user_model()


def _workbook_rows(response):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(response.content))
    sheet = book[book.sheetnames[0]]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _weekday_ahead(days: int) -> date:
    day = timezone.localdate() + timedelta(days=days)
    while day.weekday() == 6:
        day += timedelta(days=1)
    return day


class _ExportFixture(TestCase):
    def setUp(self):
        self.fy = get_operational_fy()
        self.region = Region.objects.create(name="Export Region")
        self.district = District.objects.create(
            name="Export District", region=self.region
        )
        self.user = User.objects.create(
            email="export-cceo@edify.org",
            name="Export Officer",
            roles=[EdifyRole.CCEO.value],
            active_role=EdifyRole.CCEO.value,
            status="active",
            is_active=True,
        )
        self.staff = StaffProfile.objects.create(user=self.user, title="CCEO")
        self.client_school = self._school("EXP-CLIENT-7", "Export Client Primary")
        self.core_school = self._school("EXP-CORE-3", "Export Core Primary", "core")
        plan = CorePlan.objects.create(
            id=cplan_id("EXP-CORE-3", fy=self.fy),
            school_id="EXP-CORE-3",
            fy=self.fy,
            status="Active",
        )
        create_package_slots(plan, "EXP-CORE-3", ["leadership"])

    def _school(self, code, name, school_type="client"):
        return School.objects.create(
            school_id=code,
            name=name,
            school_type=school_type,
            region=self.region,
            district=self.district,
            account_owner_id=self.staff.id,
        )

    def _visit(self, school, **fields):
        day = fields.pop("on", None) or _weekday_ahead(3)
        with self.captureOnCommitCallbacks(execute=True):
            return Activity.objects.create(
                activity_type=fields.pop("activity_type", "school_visit"),
                school=school,
                fy=get_operational_fy(day),
                quarter="Q4",
                planned_date=day,
                planned_month=day.month,
                status="scheduled",
                delivery_type="staff",
                responsible_staff_id=self.staff.id,
                **fields,
            )


class MyPlanExportTest(_ExportFixture):
    def test_every_row_carries_its_school_id_core_work_included(self):
        day = _weekday_ahead(1)
        self._visit(self.client_school, on=day)
        self._visit(self.core_school, on=day)
        self.client.force_login(self.user)

        rows = _workbook_rows(
            self.client.get(
                f"/my-plan?export=xlsx&period=fy&fy={get_operational_fy(day)}"
            )
        )

        self.assertEqual(rows[0][:3], ["Activity ID", "School ID", "Type"])
        by_school = {row[1]: row for row in rows[1:]}
        self.assertIn("EXP-CLIENT-7", by_school)
        # The Core School card used to be left out of the file altogether.
        self.assertIn("EXP-CORE-3", by_school)
        self.assertIn("(V1)", by_school["EXP-CORE-3"][2])

    def test_the_csv_carries_it_too(self):
        self._visit(self.client_school)
        self.client.force_login(self.user)

        body = self.client.get("/my-plan?export=csv&period=fy").content.decode()
        rows = list(csv.reader(io.StringIO(body)))

        self.assertEqual(rows[0][1], "School ID")
        self.assertIn("EXP-CLIENT-7", [row[1] for row in rows[1:]])


class OversightExportColumnsTest(TestCase):
    def test_the_planning_oversight_export_writes_the_business_code(self):
        from apps.planning.oversight_service import PlanningOversightItem, export_rows

        at_school = PlanningOversightItem(
            stage="staff_scheduled",
            activity_id="a1",
            school_id="pk-of-the-school",
            school_code="OVS-12",
        )
        no_code = PlanningOversightItem(
            stage="staff_scheduled",
            activity_id="a2",
            school_id="pk-2",
            school_code="pk-2",
        )
        cluster = PlanningOversightItem(
            stage="staff_scheduled", activity_id="a3", school_code="—"
        )

        rows = list(export_rows([at_school, no_code, cluster]))
        column = rows[0].index("School ID")

        self.assertEqual(
            [row[column] for row in rows[1:]],
            ["OVS-12", "", ""],
            "the business code, never the pk fallback, blank without a school",
        )

    def test_the_partner_oversight_export_writes_the_business_code(self):
        from apps.planning.partner_oversight_service import (
            PartnerOversightItem,
            export_rows,
        )

        school_item = PartnerOversightItem(
            stage="awaiting_schedule",
            partner_assignment_id="pa1",
            school_id="pk-1",
            school_code="PART-4",
            school_name="Partner School",
        )
        cluster_item = PartnerOversightItem(
            stage="awaiting_schedule",
            partner_assignment_id="pa2",
            cluster_id="cl-1",
            cluster_name="North Cluster",
        )

        rows = list(export_rows([school_item, cluster_item]))

        self.assertEqual(rows[0][1], "School ID")
        self.assertEqual([row[1] for row in rows[1:]], ["PART-4", ""])


class PartnerWorkspaceExportTest(_ExportFixture):
    def test_the_partner_organisation_s_export_names_each_school(self):
        # A Project Coordinator holds the partner workspace and its export but
        # not Partner Oversight, so it reads this page rather than being sent
        # to oversight.
        reader = User.objects.create(
            email="export-pc@edify.org",
            name="Export Coordinator",
            roles=[EdifyRole.PROJECT_COORDINATOR.value],
            active_role=EdifyRole.PROJECT_COORDINATOR.value,
            status="active",
            is_active=True,
        )
        partner = Partner.objects.create(name="Export Partner Org", active_status=True)
        PartnerAssignment.objects.create(
            school=self.client_school,
            partner=partner,
            assigning_staff_id=self.staff.id,
            expected_activity_type="school_visit",
            status=PartnerAssignment.STATUS_PENDING_SCHEDULING,
        )
        self.client.force_login(reader)

        response = self.client.get(f"/partners?export=csv&fy={self.fy}")

        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(io.StringIO(response.content.decode())))
        self.assertEqual(rows[0][:3], ["Partner", "School ID", "School / target"])
        self.assertIn("EXP-CLIENT-7", [row[1] for row in rows[1:]])
