"""A plan export is the plan, and a group training's per-school cost is the
catalogue's.

Owner, 2026-10-12: "The export for plan is not working ... it export blank";
the default export "is meant to export only planned activities (visits and
trainings and cluster meetings) with dates of activity and their budget";
and "the schools group training are still showing UGX 5000 even when it is
changed from cost catalogue ... delete the hardcoded figure and fetch
everything from the database".

Two causes of the blank: the Export options drawer's kind of work
(``visits_trainings``, ``group_trainings``...) reached the page's own
activity-type filter, which matched no activity; and the drawer opened on
"Completed" so a plan with nothing completed yet exported nothing.
"""

from __future__ import annotations

import io
from datetime import date, timedelta

from django.test import Client, RequestFactory, SimpleTestCase, TestCase
from django.utils import timezone

from apps.accounts.models import (
    StaffProfile,
    StaffSchoolAssignment,
    StaffSupervisorAssignment,
    User,
)
from apps.activities.models import Activity, ActivityScheduleCostLine
from apps.budget.costing import GROUP_TRAINING_MEALS_RATE_KEY
from apps.budget.costing_service import catalogue_rate, planned_meal_rates
from apps.budget.models import ActivityCostSnapshot, CostSetting
from apps.budget.reference import ensure_active_catalogue
from apps.clusters.models import Cluster
from apps.core import export_kinds
from apps.core.fy import get_operational_fy
from apps.core.rbac import EdifyRole
from apps.geography.models import District, Region
from apps.planning import oversight_service as oversight
from apps.schools.models import School


def _sheet(response) -> list[list]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(response.content), read_only=True)
    return [list(row) for row in book.worksheets[0].iter_rows(values_only=True)]


class ExportVocabularyTest(SimpleTestCase):
    def test_the_plan_is_visits_trainings_and_cluster_meetings(self):
        self.assertEqual(export_kinds.KIND_CHOICES[0][0], export_kinds.PLANNED)
        keeps = export_kinds.keeps
        self.assertTrue(keeps("planned", "school_visit"))
        self.assertTrue(keeps("planned", "in_school_training"))
        self.assertTrue(keeps("planned", "cluster_training"))
        self.assertTrue(keeps("planned", "cluster_meeting"))
        self.assertFalse(keeps("planned", "programme_event"))

    def test_a_kind_of_work_is_never_read_as_an_activity_type(self):
        factory = RequestFactory()
        for kind in export_kinds.KINDS:
            request = factory.get("/", {"activity_type": kind, "district_id": "d1"})
            self.assertEqual(
                oversight.read_filters(request), {"district_id": "d1"}, kind
            )
        request = factory.get("/", {"activity_type": "school_visit"})
        self.assertEqual(
            oversight.read_filters(request), {"activity_type": "school_visit"}
        )

    def test_the_drawers_plan_status_is_the_exports_not_an_exact_status(self):
        request = RequestFactory().get("/", {"status": "completed", "risk": "any"})
        self.assertEqual(oversight.read_filters(request)["status"], "completed")
        self.assertEqual(oversight.export_filters(request), {"risk": "any"})


class PlanExportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.fy = get_operational_fy()
        cls.region = Region.objects.create(id="r1", name="Central")
        cls.district = District.objects.create(
            id="d1", name="Kampala", region=cls.region
        )
        cls.pl_user, cls.pl = cls._staff(
            "pl@p.test", "Team Lead", EdifyRole.COUNTRY_PROGRAM_LEAD
        )
        cls.cceo_user, cls.cceo = cls._staff("c@p.test", "Officer", EdifyRole.CCEO)
        StaffSupervisorAssignment.objects.create(supervisee=cls.cceo, supervisor=cls.pl)
        cls.school = School.objects.create(
            school_id="s1",
            name="Alpha Primary",
            district=cls.district,
            region=cls.region,
        )
        StaffSchoolAssignment.objects.create(staff=cls.cceo, school_id=cls.school.id)
        cls.cluster = Cluster.objects.create(
            name="Kampala North", district=cls.district, region=cls.region
        )
        cls.day = date.today() + timedelta(days=6)
        cls.visit = cls._activity("school_visit", cost=40_000)
        cls.meeting = cls._activity("cluster_meeting", cost=10_000, cluster=cls.cluster)
        cls.training = cls._activity(
            "cluster_training", cost=90_000, cluster=cls.cluster
        )
        cls.event = cls._activity("programme_event", cost=5_000)

    @classmethod
    def _staff(cls, email, name, role):
        user = User.objects.create(
            email=email,
            name=name,
            roles=[role.value],
            active_role=role.value,
            is_active=True,
        )
        return user, StaffProfile.objects.create(user=user, title=name)

    @classmethod
    def _activity(cls, kind, *, cost, cluster=None):
        activity = Activity.objects.create(
            activity_type=kind,
            school=None if cluster else cls.school,
            cluster=cluster,
            fy=get_operational_fy(cls.day),
            quarter="Q1",
            planned_date=cls.day,
            planned_month=cls.day.month,
            status="scheduled",
            responsible_staff_id=cls.cceo.id,
        )
        ActivityScheduleCostLine.objects.create(
            activity=activity,
            cost_setting_key="transport",
            label="Transport",
            unit_cost=cost,
            quantity=1,
            amount=cost,
            planned_date=cls.day,
            fiscal_year=activity.fy,
            responsible_user=cls.cceo_user.id,
        )
        return activity

    def _drawer_query(self, **extra) -> dict:
        """What the Export options drawer submits, as it opens."""
        return {
            "format": "xlsx",
            "period_type": "month_range",
            "from_month": "10",
            "to_month": "9",
            "fy": get_operational_fy(self.day),
            "status": "all",
            "activity_type": export_kinds.PLANNED,
            **extra,
        }

    #: What each fixture activity reads as in the export's sheets
    #: (apps.core.plan_export): the sheet it is on and how it is named there.
    VISIT = ("School Visits", "School Visit")
    MEETING = ("Cluster Meetings", "Cluster Meeting")
    TRAINING = ("Group Trainings", "Cluster Training")

    def _work(self, response) -> list:
        """The exported work, a (sheet, name) per row, sorted."""
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(response.content), read_only=True)
        found = []
        for sheet in book.worksheets:
            if sheet.title == "Budget":
                continue
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            header = rows[0]
            named = header.index("Activity" if "Activity" in header else "Training")
            found += [(sheet.title, row[named]) for row in rows[1:]]
        return sorted(found)

    def _statuses(self, response) -> list:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(response.content), read_only=True)
        out = []
        for sheet in book.worksheets:
            if sheet.title == "Budget":
                continue
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            out += [row[rows[0].index("Status")] for row in rows[1:]]
        return out

    def test_my_plan_exports_the_plan_as_the_drawer_opens(self):
        client = Client()
        client.force_login(self.cceo_user)
        response = client.get("/my-plan/", {**self._drawer_query(), "export": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            self._work(response), sorted([self.VISIT, self.MEETING, self.TRAINING])
        )

    def test_completed_only_is_completed_work_not_a_label_with_the_word(self):
        """ "Not Complete" has "complet" in it; the old filter kept every row."""
        done = self._activity("school_visit", cost=1_000)
        Activity.objects.filter(pk=done.pk).update(status="completed")
        started = self._activity("school_visit", cost=1_000)
        Activity.objects.filter(pk=started.pk).update(status="completion_started")
        client = Client()
        client.force_login(self.cceo_user)
        response = client.get(
            "/my-plan/",
            {**self._drawer_query(status="completed"), "export": "xlsx"},
        )
        # One row of the three visits: the completed one.
        self.assertEqual(self._work(response), [self.VISIT])
        self.assertEqual(len(self._statuses(response)), 1)
        client.force_login(self.pl_user)
        response = client.get(
            "/team-planning-oversight/export", self._drawer_query(status="completed")
        )
        self.assertEqual(self._work(response), [self.VISIT])
        self.assertEqual(self._statuses(response), ["Completed"])

    def test_my_plan_exports_one_kind_when_asked(self):
        client = Client()
        client.force_login(self.cceo_user)
        response = client.get(
            "/my-plan/",
            {**self._drawer_query(activity_type="group_trainings"), "export": "xlsx"},
        )
        self.assertEqual(self._work(response), [self.TRAINING])

    def test_my_plan_drawer_offers_the_kinds_and_opens_on_the_plan(self):
        client = Client()
        client.force_login(self.cceo_user)
        html = client.get("/my-plan/export-drawer").content.decode()
        self.assertIn("Group trainings only", html)
        self.assertIn('<option value="all" selected>', html)
        self.assertNotIn('value="completed" selected', html)

    def test_team_plan_exports_every_kind_the_drawer_names(self):
        client = Client()
        client.force_login(self.pl_user)
        expected = {
            export_kinds.PLANNED: [self.VISIT, self.MEETING, self.TRAINING],
            export_kinds.GROUP_TRAININGS: [self.TRAINING],
            export_kinds.MEETINGS: [self.MEETING],
            export_kinds.VISITS: [self.VISIT],
        }
        for kind, ids in expected.items():
            response = client.get(
                "/team-planning-oversight/export",
                self._drawer_query(activity_type=kind),
            )
            self.assertEqual(response.status_code, 200, kind)
            self.assertEqual(self._work(response), sorted(ids), kind)


class GroupTrainingMealRateTest(TestCase):
    """The per-school cost of a group training is heads × the meal rate the
    database holds, never a literal 5,000."""

    @classmethod
    def setUpTestData(cls):
        cls.catalogue = ensure_active_catalogue()
        cls.region = Region.objects.create(id="r2", name="West")
        cls.district = District.objects.create(
            id="d2", name="Mbarara", region=cls.region
        )
        cls.user = User.objects.create(
            email="m@p.test",
            name="Officer",
            roles=[EdifyRole.CCEO.value],
            active_role=EdifyRole.CCEO.value,
            is_active=True,
        )
        cls.staff = StaffProfile.objects.create(user=cls.user, title="Officer")
        cls.cluster = Cluster.objects.create(
            name="Mbarara East", district=cls.district, region=cls.region
        )
        cls.day = date.today() + timedelta(days=6)

    def _training(self, *, priced_meals=None):
        activity = Activity.objects.create(
            activity_type="cluster_training",
            cluster=self.cluster,
            fy=get_operational_fy(self.day),
            quarter="Q1",
            planned_date=self.day,
            planned_month=self.day.month,
            status="scheduled",
            responsible_staff_id=self.staff.id,
            expected_participants=10,
        )
        if priced_meals is not None:
            ActivityCostSnapshot.objects.create(
                activity=activity,
                operational_rate_card=self.catalogue,
                operational_cost=priced_meals * 10,
                operational_breakdown=[
                    {
                        "key": GROUP_TRAINING_MEALS_RATE_KEY,
                        "label": "Group Training - Participant Meals",
                        "unit": priced_meals,
                        "qty": 10,
                        "amount": priced_meals * 10,
                        "missing": False,
                    }
                ],
                calculated_at=timezone.now(),
            )
        return activity

    def _set_rate(self, amount):
        CostSetting.objects.filter(
            catalogue=self.catalogue, key=GROUP_TRAINING_MEALS_RATE_KEY
        ).update(unit_cost=amount, approved_minimum=amount)

    def test_the_rate_is_the_catalogues_and_follows_its_edits(self):
        costed = self._training(priced_meals=5000)
        uncosted = self._training()
        self._set_rate(7000)
        self.assertEqual(catalogue_rate(GROUP_TRAINING_MEALS_RATE_KEY), 7000)
        rates = planned_meal_rates([costed, uncosted], GROUP_TRAINING_MEALS_RATE_KEY)
        # The costed session re-prices its own meals line at the new rate, the
        # way the rest of My Plan prices a row; the uncosted one reads the card.
        self.assertEqual(rates, {costed.id: 7000, uncosted.id: 7000})

    def test_a_rate_the_card_does_not_carry_is_nothing_not_five_thousand(self):
        training = self._training()
        CostSetting.objects.filter(
            catalogue=self.catalogue, key=GROUP_TRAINING_MEALS_RATE_KEY
        ).delete()
        self.assertIsNone(catalogue_rate(GROUP_TRAINING_MEALS_RATE_KEY))
        self.assertEqual(
            planned_meal_rates([training], GROUP_TRAINING_MEALS_RATE_KEY),
            {training.id: None},
        )

    def test_my_plan_shows_the_schools_share_at_the_catalogue_rate(self):
        for code in ("MR-1", "MR-2"):
            School.objects.create(
                school_id=code,
                name=f"Member {code}",
                region=self.region,
                district=self.district,
                cluster_id=self.cluster.id,
                cluster_status="clustered",
            )
        self._training()
        self._set_rate(7000)
        client = Client()
        client.force_login(self.user)
        html = client.get(
            "/my-plan/", {"fy": get_operational_fy(self.day)}
        ).content.decode()
        # The session's ten heads over its two member schools: 5 × 7,000. It
        # read two heads a school when the session named no split, a figure
        # nobody planned (owner, 2026-10-15: "no misleading costs").
        self.assertIn("UGX 35,000", html)
        self.assertNotIn("UGX 25,000", html)
