"""Planning Oversight exports one table (owner, 2026-10-20).

"the planning oversight table export columns should have (Financial year,
Table, School ID, School name, District, Cluster, Activity type, Training
name, SSA intervention, Purpose of visit, Planned date, Delivery type,
Participants, Planning stage, Planned by, Operational owner, Executor type,
Executor, Managing staff, Supervising PL, Partner, Planned cost (UGX),
Activity status, Assignment status, Salesforce ID, IA, Finance, Risks, Next
action owner)". Asked, the owner chose: one table and no Budget sheet, a
cluster training once, and My Plan and Cluster Oversight keep their sheets.
"""

from __future__ import annotations

import csv
import io

from django.test import Client

from apps.core.rbac import EdifyRole
from apps.frontend.test_export_period_and_sessions import (
    FACILITATION,
    MEALS,
    VENUE,
    SessionFixture,
)
from apps.frontend.test_oversight_filters_export import PL_EXPORT

COUNTRY_EXPORT = "/country-planning-oversight/export"

OWNER_COLUMNS = [
    "Financial year",
    "Table",
    "School ID",
    "School name",
    "District",
    "Cluster",
    "Activity type",
    "Training name",
    "SSA intervention",
    "Purpose of visit",
    "Planned date",
    "Delivery type",
    "Participants",
    "Planning stage",
    "Planned by",
    "Operational owner",
    "Executor type",
    "Executor",
    "Managing staff",
    "Supervising PL",
    "Partner",
    "Planned cost (UGX)",
    "Activity status",
    "Assignment status",
    "Salesforce ID",
    "IA",
    "Finance",
    "Risks",
    "Next action owner",
]


class PlanningOversightExportsOneTable(SessionFixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.cd_user, _ = cls._staff("cd@f.test", "Director", EdifyRole.COUNTRY_DIRECTOR)
        cls.visit.salesforce_activity_id = "a0V5g00000VISIT1"
        cls.visit.save(update_fields=["salesforce_activity_id"])

    def lines(self, client=None, path=PL_EXPORT, **params):
        client = client or self.as_pl()
        response = client.get(
            path, {"fy": self.fy, "owner": self.james.id, "format": "csv", **params}
        )
        self.assertEqual(response.status_code, 200)
        body = b"".join(response.streaming_content).decode()
        return list(csv.reader(io.StringIO(body)))

    def records(self, **params) -> list[dict]:
        header, *rows = self.lines(**params)
        return [dict(zip(header, row)) for row in rows]

    def test_the_header_is_the_owners_columns_in_order(self):
        self.assertEqual(self.lines()[0], OWNER_COLUMNS)

    def test_it_is_one_table_with_no_budget_sheet(self):
        lines = self.lines()
        self.assertTrue(all(len(line) == len(OWNER_COLUMNS) for line in lines))
        self.assertEqual(lines.count(OWNER_COLUMNS), 1)
        self.assertNotIn(["Budget"], lines)

    def test_each_row_names_the_table_the_page_shows_it_on(self):
        # The officer's visit and the partner's, the training at the school
        # and the cluster training.
        self.assertEqual(
            sorted(row["Table"] for row in self.records()),
            ["Client School Visit", "Client School Visit", "Training", "Training"],
        )

    def test_a_cluster_training_is_one_row_with_its_people_and_its_cost(self):
        rows = [r for r in self.records() if r["Cluster"] == "Kasangati Cluster"]
        self.assertEqual(len(rows), 1, rows)
        self.assertEqual(rows[0]["School ID"], "")
        self.assertEqual(rows[0]["Participants"], "12")
        self.assertEqual(
            rows[0]["Planned cost (UGX)"], str(MEALS + VENUE + FACILITATION)
        )

    def test_the_salesforce_id_is_written(self):
        ids = {r["Salesforce ID"] for r in self.records(activity_type="school_visit")}
        self.assertIn("a0V5g00000VISIT1", ids)

    def test_the_pages_filters_still_narrow_it(self):
        rows = self.records(activity_type="training")
        self.assertEqual([r["Activity type"] for r in rows], ["training"])

    def test_the_country_export_is_the_same_table(self):
        client = Client()
        client.force_login(self.cd_user)
        self.assertEqual(self.lines(client, COUNTRY_EXPORT)[0], OWNER_COLUMNS)

    def test_the_workbook_is_the_same_one_table(self):
        from openpyxl import load_workbook

        response = self.as_pl().get(
            PL_EXPORT, {"fy": self.fy, "owner": self.james.id, "format": "xlsx"}
        )
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(book.sheetnames), 1, book.sheetnames)
        sheet = book.worksheets[0]
        self.assertEqual([cell.value for cell in sheet[1]], OWNER_COLUMNS)
        self.assertEqual(sheet.max_row, 1 + len(self.records()))


class MyPlanKeepsItsSheets(SessionFixture):
    def test_my_plan_still_exports_its_sheets(self):
        tables = self.my_plan_tables(fy=self.fy)
        self.assertIn("School Visits", tables)
        self.assertIn("Group Trainings", tables)
