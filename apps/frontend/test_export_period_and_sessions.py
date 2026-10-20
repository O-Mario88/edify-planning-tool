"""Exports pick the chosen period, export their own page, and a group training
reads its planned people and its own cost (owner, 2026-10-15).

"The Exported group training should export the actual number of
participants (Number of schools x number of Participants) as planned. and it
should reflect the actual (cost of participant * 12000)+venue Fee+ Facilitation
Fee. Also it is not selecting the period it is still exporting the entire
plan. make sure it only picks the plan for the selected period not
everything." Then: "if it is planning oversight page, it should fetch data
from planning oversight, if it is any other page, it should fetch data for
that specific page".

Asked, the owner chose: a group training's Cost is the participants' meals,
the venue fee and the facilitation fee — the officer's own transport and lunch
that day stay on the Budget sheet as their own lines; Cluster Oversight
exports the tables it shows.

Found:

  1. The drawer always sends both the month pair and the custom dates. The
     export read the months first, so "Custom Dates" exported the whole
     October-to-September year.
  2. The drawer opened From the page's month To September: a month on the
     page became the rest of the year in the file.
  3. The oversight exports wrote a group training's per-school figure (3)
     where the page shows the people invited (12).
  4. Planning Oversight's drawer chose the country or team export by the
     reader's role, not by the page it was opened on.
  5. Cluster Oversight exported Planning Oversight's activity sheets.
"""

from __future__ import annotations

import io
import re
from datetime import timedelta

from django.test import Client

from apps.activities.models import (
    Activity,
    ActivityScheduleCostLine,
    ClusterActivityAttendance,
)
from apps.clusters.models import Cluster
from apps.core.fy import get_fy_date_range
from apps.core.rbac import EdifyRole
from apps.frontend.test_oversight_filters_export import PL_EXPORT, FilterExportFixture
from apps.schools.models import School

MEALS, VENUE, FACILITATION = 12 * 12_000, 30_000, 50_000
TRANSPORT, LUNCH = 50_000, 12_000


def _chosen(html: str, name: str) -> str:
    block = re.search(rf'name="{name}".*?</select>', html, re.S).group(0)
    return re.search(r'<option value="(\d+)"\s+selected', block).group(1)


class SessionFixture(FilterExportFixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        fy_start, _ = get_fy_date_range(cls.fy)
        # Mid-year: in this year whatever day the suite runs.
        cls.day = fy_start.date() + timedelta(days=200)
        cls.cluster = Cluster.objects.create(
            name="Kasangati Cluster",
            region=cls.region,
            district=cls.district,
            responsible_staff_id=cls.james.id,
        )
        cls.members = [
            School.objects.create(
                school_id=f"gt{n}",
                name=f"Member {n}",
                district=cls.district,
                region=cls.region,
                cluster_id=cls.cluster.id,
                cluster_status="clustered",
            )
            for n in range(4)
        ]
        cls.session = Activity.objects.create(
            activity_type="cluster_training",
            cluster=cls.cluster,
            fy=cls.fy,
            quarter="Q3",
            planned_date=cls.day,
            planned_month=cls.day.month,
            status="scheduled",
            delivery_type="staff",
            responsible_staff_id=cls.james.id,
            participants_per_school=3,
            schools_invited=4,
            expected_participants=12,
        )
        for school in cls.members:
            ClusterActivityAttendance.objects.create(
                activity=cls.session, school=school, invited=True
            )
        for key, kind, label, rate, qty in (
            ("group_training_meals", "participant_meals", "Meals", 12_000, 12),
            ("group_training_venue_cost", "venue", "Venue Fee", VENUE, 1),
            (
                "group_training_facilitation_fee",
                "facilitation",
                "Facilitation Fee",
                FACILITATION,
                1,
            ),
            ("primary_transport_per_day", "transport", "Transport", TRANSPORT, 1),
            ("primary_lunch_per_day", "lunch", "Lunch", LUNCH, 1),
        ):
            ActivityScheduleCostLine.objects.create(
                activity=cls.session,
                cost_setting_key=key,
                line_item_type=kind,
                label=label,
                unit_cost=rate,
                quantity=qty,
                amount=rate * qty,
            )

    def my_plan_tables(self, **params):
        import csv

        from apps.core import plan_export

        client = Client()
        client.force_login(self.james_user)
        response = client.get("/my-plan/", {"export": "csv", **params})
        self.assertEqual(response.status_code, 200)
        lines = list(csv.reader(io.StringIO(response.content.decode())))
        return plan_export.tables_from_csv(lines)

    def oversight_dates(self, **params):
        return {
            row["Planned date"]
            for row in self.export_records(**params)
            if row["Planned date"]
        }

    @staticmethod
    def dates(tables):
        return {
            row["Planned Date"]
            for rows in tables.values()
            for row in rows
            if row.get("Planned Date")
        }


class AGroupTrainingReadsItsPeopleAndItsOwnCost(SessionFixture):
    def _row(self, tables):
        rows = [
            row
            for row in tables["Group Trainings"]
            if row["Cluster Name"] == "Kasangati Cluster"
        ]
        self.assertEqual(len(rows), 1, tables["Group Trainings"])
        return rows[0]

    def test_planning_oversight(self):
        # One table (owner, 2026-10-20): the cluster training is one row.
        rows = [
            row
            for row in self.export_records(fy=self.fy)
            if row["Cluster"] == "Kasangati Cluster"
        ]
        self.assertEqual(len(rows), 1, rows)
        self.assertEqual(rows[0]["Participants"], "12")
        self.assertEqual(
            rows[0]["Planned cost (UGX)"], str(MEALS + VENUE + FACILITATION)
        )

    def test_my_plan(self):
        row = self._row(self.my_plan_tables(fy=self.fy))
        self.assertEqual(row["Participants"], "12")
        self.assertEqual(row["Cost (UGX)"], str(MEALS + VENUE + FACILITATION))

    def test_the_officers_day_stays_on_the_budget_sheet(self):
        from apps.core import plan_export

        made = plan_export.budget_sheet(
            [{"activity_id": self.session.id, "planned_date": self.day}]
        )
        self.assertEqual(
            made["rows"][0][5], MEALS + VENUE + FACILITATION + TRANSPORT + LUNCH
        )


class TheChosenPeriodIsTheOneExported(SessionFixture):
    """The drawer sends every field it holds; the period type chosen decides
    which of them is read."""

    def _custom(self):
        return {
            "fy": self.fy,
            "period_type": "custom_dates",
            "start_date": self.day.isoformat(),
            "end_date": self.day.isoformat(),
            "format": "csv",
            # Hidden in the drawer, still sent.
            "from_month": "10",
            "to_month": "9",
        }

    def _months(self):
        fy_start, fy_end = get_fy_date_range(self.fy)
        return {
            "fy": self.fy,
            "period_type": "month_range",
            "from_month": str(self.day.month),
            "to_month": str(self.day.month),
            "start_date": fy_start.date().isoformat(),
            "end_date": (fy_end.date() - timedelta(days=1)).isoformat(),
            "format": "csv",
        }

    def test_custom_dates_on_planning_oversight(self):
        self.assertEqual(self.oversight_dates(**self._custom()), {self.day.isoformat()})

    def test_custom_dates_on_my_plan(self):
        self.assertEqual(
            self.dates(self.my_plan_tables(**self._custom())), {self.day.isoformat()}
        )

    def test_a_month_on_planning_oversight(self):
        dates = self.oversight_dates(**self._months())
        self.assertIn(self.day.isoformat(), dates)
        self.assertEqual({d[:7] for d in dates}, {self.day.isoformat()[:7]})

    def test_a_month_on_my_plan(self):
        dates = self.dates(self.my_plan_tables(**self._months()))
        self.assertIn(self.day.isoformat(), dates)
        self.assertEqual({d[:7] for d in dates}, {self.day.isoformat()[:7]})

    def test_the_workbook_holds_the_period_alone(self):
        from openpyxl import load_workbook

        response = self.as_pl().get(
            PL_EXPORT, {**self._custom(), "format": "xlsx", "owner": self.james.id}
        )
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        header, *rows = book.worksheets[0].iter_rows(values_only=True)
        at = header.index("Planned date")
        self.assertEqual({row[at].date() for row in rows}, {self.day})


class TheDrawerOpensOnThePagesPeriod(SessionFixture):
    def _selected(self, **params):
        html = (
            self.as_pl()
            .get("/team-planning-oversight/export-drawer", {"fy": self.fy, **params})
            .content.decode()
        )
        return _chosen(html, "from_month"), _chosen(html, "to_month")

    def test_a_month(self):
        self.assertEqual(self._selected(period="month", month="11"), ("11", "11"))

    def test_a_quarter(self):
        self.assertEqual(self._selected(period="quarter", quarter="Q2"), ("1", "3"))

    def test_the_year_opens_on_the_current_month(self):
        # Owner, 2026-10-15: "the planning drawer should default to the
        # current month".
        from django.utils import timezone

        month = str(timezone.localdate().month)
        self.assertEqual(self._selected(period="fy"), (month, month))
        self.assertEqual(self._selected(), (month, month))

    def test_my_plan_with_a_month(self):
        client = Client()
        client.force_login(self.james_user)
        html = client.get(
            "/my-plan/export-drawer", {"fy": self.fy, "month": "11"}
        ).content.decode()
        self.assertEqual(
            (_chosen(html, "from_month"), _chosen(html, "to_month")), ("11", "11")
        )


class EachPageExportsItself(SessionFixture):
    def test_team_planning_oversight_exports_the_team_page_for_a_director(self):
        director, _ = self._staff("cd@f.test", "Director", EdifyRole.COUNTRY_DIRECTOR)
        client = Client()
        client.force_login(director)
        html = client.get(
            "/team-planning-oversight/export-drawer", {"fy": self.fy}
        ).content.decode()
        self.assertIn('action="/team-planning-oversight/export"', html)
        html = client.get(
            "/country-planning-oversight/export-drawer", {"fy": self.fy}
        ).content.decode()
        self.assertIn('action="/country-planning-oversight/export"', html)
