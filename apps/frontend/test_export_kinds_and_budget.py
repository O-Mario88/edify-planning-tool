"""The Export options' kinds, the budget column, and the Budget and Weekly
Advance Request exports.

Owner, 2026-10-12: "modify the export drawer for users to decide what they
want to export (Only group training, Only visit, only cluster meeting, only
online meeting or All). The budget column is one of the most important column
it should not be left out"; "extend the same export functionality to budget
and fund requests as well"; "make sure the budget is an accurate extract from
the actual planned activities for that period".
"""

from __future__ import annotations

import io
from datetime import date, timedelta
from pathlib import Path

from django.conf import settings
from django.test import Client, SimpleTestCase, TestCase

from apps.accounts.models import (
    StaffProfile,
    StaffSchoolAssignment,
    StaffSupervisorAssignment,
    User,
)
from apps.activities.models import Activity, ActivityScheduleCostLine
from apps.clusters.models import Cluster
from apps.core import export_kinds
from apps.core.fy import get_operational_fy
from apps.core.rbac import EdifyRole
from apps.geography.models import District, Region
from apps.schools.models import School

DRAWER = Path(settings.BASE_DIR) / (
    "templates/partials/oversight/export_customization_drawer.html"
)


def _sheets(response) -> dict[str, list[list]]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(response.content), read_only=True)
    return {
        sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
        for sheet in book.worksheets
    }


class KindsTest(SimpleTestCase):
    def test_the_drawer_offers_every_kind(self):
        """The select is the vocabulary, looped; nothing hand-typed beside it."""
        source = DRAWER.read_text()
        self.assertIn("{% for value, text in kind_choices %}", source)
        self.assertNotIn('value="visits_trainings" selected', source)
        keys = [key for key, _ in export_kinds.KIND_CHOICES]
        for wanted in ("visits", "group_trainings", "meetings", "online", "all"):
            self.assertIn(wanted, keys)

    def test_each_kind_keeps_its_own_work(self):
        keeps = export_kinds.keeps
        self.assertTrue(keeps("visits", "school_visit"))
        self.assertFalse(keeps("visits", "cluster_training"))
        self.assertTrue(keeps("group_trainings", "cluster_training"))
        self.assertFalse(keeps("group_trainings", "in_school_training"))
        self.assertTrue(keeps("meetings", "cluster_meeting"))
        self.assertFalse(keeps("meetings", "school_visit"))
        # Online is a group training delivered online, nothing else.
        self.assertTrue(keeps("online", "cluster_training", online=True))
        self.assertFalse(keeps("online", "cluster_training", online=False))
        self.assertFalse(keeps("online", "school_visit", online=True))
        self.assertTrue(keeps("visits_trainings", "in_school_training"))
        self.assertFalse(keeps("visits_trainings", "cluster_meeting"))
        self.assertTrue(keeps("all", "programme_event"))

    def test_an_unknown_kind_is_the_whole_page(self):
        self.assertEqual(export_kinds.read_kind({"activity_type": "bananas"}), "all")
        self.assertEqual(export_kinds.read_kind({}), "all")
        self.assertEqual(export_kinds.read_kind({"activity_type": "visits"}), "visits")

    def test_every_plan_export_names_its_budget_column(self):
        from apps.core import plan_export
        from apps.planning.partner_oversight_service import EXPORT_HEADER
        from apps.projects.school_table import COLUMNS
        from apps.budget.services import (
            ACTIVITY_SHEET_HEADERS,
            LINE_SHEET_HEADERS,
            SUMMARY_SHEET_HEADERS,
        )
        from apps.frontend.views.budget_views import (
            REQUEST_LINE_SHEET_HEADERS,
            REQUEST_SHEET_HEADERS,
        )

        for headers in (
            *plan_export.ALL_HEADERS,
            EXPORT_HEADER,
            COLUMNS,
            ACTIVITY_SHEET_HEADERS,
            LINE_SHEET_HEADERS,
            SUMMARY_SHEET_HEADERS,
            REQUEST_SHEET_HEADERS,
            REQUEST_LINE_SHEET_HEADERS,
        ):
            self.assertTrue(
                any("(UGX)" in h for h in headers), f"no UGX column in {headers}"
            )


class BudgetExportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.fy = get_operational_fy()
        cls.region = Region.objects.create(id="r1", name="Central")
        cls.district = District.objects.create(
            id="d1", name="Kampala", region=cls.region
        )
        cls.pl_user, cls.pl = cls._staff(
            "pl@b.test", "Team Lead", EdifyRole.COUNTRY_PROGRAM_LEAD
        )
        cls.cceo_user, cls.cceo = cls._staff("c@b.test", "Officer", EdifyRole.CCEO)
        StaffSupervisorAssignment.objects.create(supervisee=cls.cceo, supervisor=cls.pl)
        cls.school = School.objects.create(
            school_id="s1",
            name="Alpha Primary",
            district=cls.district,
            region=cls.region,
        )
        StaffSchoolAssignment.objects.create(staff=cls.cceo, school_id=cls.school.id)
        cls.cluster = Cluster.objects.create(
            name="Kampala North", district=cls.district, region=cls.region
        )
        cls.today = date.today()
        cls.start = cls.today
        cls.end = cls.today + timedelta(days=10)
        inside = cls.today + timedelta(days=6)
        outside = cls.today + timedelta(days=40)
        cls.visit = cls._activity("school_visit", inside, cost=40_000)
        cls.meeting = cls._activity(
            "cluster_meeting", inside, cost=10_000, cluster=cls.cluster
        )
        cls.online = cls._activity(
            "cluster_training", inside, cost=0, cluster=cls.cluster, mode="online"
        )
        # Called off: no longer anybody's plan, so no longer anybody's budget.
        cls.cancelled = cls._activity(
            "school_visit", inside, cost=99_000, status="cancelled"
        )
        cls.later = cls._activity("school_visit", outside, cost=70_000)

    @classmethod
    def _staff(cls, email, name, role):
        user = User.objects.create(
            email=email,
            name=name,
            roles=[role.value],
            active_role=role.value,
            is_active=True,
        )
        return user, StaffProfile.objects.create(user=user, title=name)

    @classmethod
    def _activity(cls, kind, planned, *, cost, cluster=None, mode=None, status=None):
        activity = Activity.objects.create(
            activity_type=kind,
            school=None if cluster else cls.school,
            cluster=cluster,
            fy=get_operational_fy(planned),
            quarter="Q1",
            planned_date=planned,
            scheduled_date=None,
            planned_month=planned.month,
            status=status or "scheduled",
            responsible_staff_id=cls.cceo.id,
            programme_delivery_mode=mode,
        )
        if cost:
            ActivityScheduleCostLine.objects.create(
                activity=activity,
                cost_setting_key="transport",
                label="Transport",
                unit_cost=cost,
                quantity=1,
                amount=cost,
                planned_date=planned,
                fiscal_year=activity.fy,
                responsible_user=cls.cceo_user.id,
            )
        return activity

    def export(self, **params):
        client = Client()
        client.force_login(self.pl_user)
        query = {
            "period_type": "custom_dates",
            "start_date": self.start.isoformat(),
            "end_date": self.end.isoformat(),
            "budget_scope": "team",
            "activity_type": "all",
            "status": "all",
            **params,
        }
        response = client.get("/budget/export", query)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        return _sheets(response)

    def _budgets(self, sheets) -> dict[str, int]:
        rows = sheets["Planned Activities"]
        header = rows[0]
        self.assertIn("Budget (UGX)", header)
        at_id, at_budget = header.index("Activity ID"), header.index("Budget (UGX)")
        return {row[at_id]: row[at_budget] for row in rows[1:]}

    def test_the_budget_is_the_periods_planned_activities_and_their_lines(self):
        budgets = self._budgets(self.export())
        self.assertEqual(budgets, {self.visit.id: 40_000, self.meeting.id: 10_000})
        summary = self.export()["Summary"]
        total = [row for row in summary[1:] if row[0] == "Total"][0]
        self.assertEqual(total[1:], [2, 50_000])

    def test_a_kind_narrows_the_budget_to_that_work(self):
        self.assertEqual(
            self._budgets(self.export(activity_type="visits")),
            {self.visit.id: 40_000},
        )
        self.assertEqual(
            self._budgets(self.export(activity_type="meetings")),
            {self.meeting.id: 10_000},
        )
        # An online training has no cost line, so no budget row either.
        self.assertEqual(self._budgets(self.export(activity_type="online")), {})

    def test_completed_only_leaves_scheduled_work_out(self):
        self.assertEqual(self._budgets(self.export(status="completed")), {})
        self.assertEqual(
            self._budgets(self.export(status="scheduled")),
            {self.visit.id: 40_000, self.meeting.id: 10_000},
        )

    def test_the_cost_lines_sheet_carries_every_line_of_the_period(self):
        lines = self.export()["Cost Lines"]
        header = lines[0]
        at_amount = header.index("Amount (UGX)")
        self.assertEqual(sorted(row[at_amount] for row in lines[1:]), [10_000, 40_000])

    def test_the_first_sheet_is_the_pages_activity_cost_plan(self):
        """Owner, 2026-10-12: "make sure the budget is exported exactly in the
        budget format" — the page's ledger: title and Grand total, a band per
        activity type, that table's own headings, items, subtotal, Grand
        total."""
        sheets = self.export()
        self.assertEqual(next(iter(sheets)), "Activity cost plan")
        rows = [
            [cell for cell in row] + [None] * (6 - len(row))
            for row in sheets["Activity cost plan"]
        ]
        self.assertEqual(rows[0][0], "Activity cost plan")
        self.assertEqual(rows[0][4:6], ["Grand total", 50_000])
        self.assertIn("grouped by activity and item", rows[1][0])
        first = [row[0] for row in rows]
        visit = first.index("School Visit · 1 activity")
        self.assertEqual(
            rows[visit + 1],
            [
                "Item",
                "No. of Days",
                "No. of Persons",
                "Rate / Unit",
                "Staff cost",
                "Activity",
            ],
        )
        self.assertEqual(rows[visit + 2], ["Transport", 1, 1, 40_000, 40_000, 40_000])
        self.assertEqual(
            rows[visit + 3], ["School Visit subtotal", None, None, None, 40_000, 40_000]
        )
        meeting = first.index("Cluster Meeting · 1 activity")
        self.assertEqual(rows[meeting + 1][1:3], ["Meetings", "Participants"])
        self.assertEqual(rows[-1], ["Grand total", None, None, None, 50_000, 50_000])
        # The larger table first, as on the page.
        self.assertLess(visit, meeting)

    def test_the_cost_plan_is_the_pages_own_ledger(self):
        """Every row of the sheet is a row `budget_groups` gave the page."""
        from apps.budget.services import budget_workspace, cost_plan_sheet

        page = budget_workspace(
            self.pl_user,
            {
                "budget_scope": "team",
                "period": "fy",
                "fy": self.fy,
                "date": self.today.isoformat(),
                "plan_only": True,
            },
        )
        sheet = cost_plan_sheet(page["groups"])
        self.assertEqual(sheet["rows"][0][5], page["total"])
        self.assertEqual(sheet["rows"][-1][4:6], [page["staff_total"], page["total"]])
        bands = [i for i, name in sheet["row_styles"].items() if name == "band"]
        self.assertEqual(len(bands), len(page["groups"]))
        for at, group in zip(bands, page["groups"]):
            body = sheet["rows"][at + 1 : at + 1 + len(group["rows"])]
            self.assertEqual([r[0] for r in body], [r["item"] for r in group["rows"]])
            self.assertEqual([r[5] for r in body], [r["total"] for r in group["rows"]])

    def test_the_cost_plan_shows_ugx_and_each_actual_rate(self):
        from openpyxl import load_workbook

        # A second visit at another rate: two Transport rows, each its rate.
        self._activity("school_visit", self.today + timedelta(days=7), cost=5_000)
        client = Client()
        client.force_login(self.pl_user)
        response = client.get(
            "/budget/export",
            {
                "period_type": "custom_dates",
                "start_date": self.start.isoformat(),
                "end_date": self.end.isoformat(),
                "budget_scope": "team",
                "activity_type": "all",
                "status": "all",
            },
        )
        sheet = load_workbook(io.BytesIO(response.content))["Activity cost plan"]
        rows = [r for r in sheet.iter_rows() if r[0].value == "Transport"]
        # "Pick the actual rate from the cost catalogue and as calculated".
        self.assertEqual(
            # The School Visit table's two, then the Cluster Meeting's one.
            [(r[3].value, r[5].value) for r in rows],
            [(40_000, 40_000), (5_000, 5_000), (10_000, 10_000)],
        )
        row = rows[0]
        self.assertEqual(row[5].number_format, '"UGX" #,##0')
        # The Grand total beside the title is money too.
        top = next(sheet.iter_rows())
        self.assertEqual(top[4].value, "Grand total")
        self.assertEqual(top[5].number_format, '"UGX" #,##0')
        self.assertEqual(row[5].alignment.horizontal, "right")
        self.assertEqual(row[1].alignment.horizontal, "center")
        band = next(
            r
            for r in sheet.iter_rows()
            if str(r[0].value) == "School Visit · 2 activities"
        )
        # The page's light band, dark ink: not the navy of a flat export.
        self.assertEqual(band[0].fill.fgColor.rgb[-6:], "DCE6EC")
        self.assertTrue(band[0].font.i)
        # A ledger, not a filterable flat table.
        self.assertIsNone(sheet.auto_filter.ref)
        self.assertIsNone(sheet.freeze_panes)

    def test_the_drawers_open_for_the_budget_and_the_advance_requests(self):
        client = Client()
        client.force_login(self.pl_user)
        for path in ("/budget/export-drawer", "/fund-requests/weekly/export-drawer"):
            response = client.get(path, {"fy": self.fy})
            self.assertEqual(response.status_code, 200, path)
            body = response.content.decode()
            self.assertIn('value="group_trainings"', body)
            self.assertIn('value="meetings"', body)
            self.assertIn('value="online"', body)
        self.assertIn('name="export" value="xlsx"', body)
        self.assertIn('value="disbursed"', body)

    def test_the_advance_requests_export_is_a_workbook_with_an_amount_column(self):
        client = Client()
        client.force_login(self.pl_user)
        response = client.get(
            "/fund-requests/weekly",
            {
                "export": "xlsx",
                "fy": self.fy,
                "period_type": "custom_dates",
                "start_date": self.start.isoformat(),
                "end_date": self.end.isoformat(),
            },
        )
        self.assertEqual(response.status_code, 200)
        sheets = _sheets(response)
        self.assertIn("Amount (UGX)", sheets["Cost Lines"][0])
        self.assertIn("Total Amount (UGX)", sheets["Requests"][0])

    def test_my_plan_exports_its_budget_column(self):
        client = Client()
        client.force_login(self.cceo_user)
        response = client.get("/my-plan", {"export": "xlsx", "fy": self.fy})
        self.assertEqual(response.status_code, 200)
        # My Plan's School Visits table (owner, 2026-10-26): its Cost column,
        # written in UGX as the page writes it.
        from openpyxl import load_workbook

        from apps.my_plan import table_export

        sheet = load_workbook(io.BytesIO(response.content))["School Visits"]
        header = [cell.value for cell in sheet[1]]
        self.assertIn("Cost", header)
        at = header.index("Cost") + 1
        self.assertEqual(sheet.cell(2, at).number_format, table_export.MONEY)
