"""Planning Oversight's export holds what the page was narrowed to (owner,
2026-10-19: "When i select Visit only it was exporting everything and for the
whole FY ... everything exported should be limited to what i have selected").

Found on the deployed page: the School Visits, Cluster Meetings and Group
Trainings tabs (``?activity=``) narrowed the page, but neither the Export
button nor its drawer heard of them. The drawer opened on "Visits, trainings &
cluster meetings", so the file and its Budget held every kind of work. The
drawer also dropped the page's district, executor, partner, risk and context
filters, which the plain Export link kept.
"""

from __future__ import annotations

import io
import re

from django.test import Client

from apps.frontend.test_oversight_filters_export import (
    PL_EXPORT,
    PL_URL,
    FilterExportFixture,
)

DRAWER = "/team-planning-oversight/export-drawer"
# James's visit and the partner's visit.
VISITS = [("School Visits", "James"), ("School Visits", "Partner X")]
# And James's training.
EVERY = sorted([*VISITS, ("Trainings", "James")])


def _rows(response) -> list[tuple[str, str]]:
    """Every row of the workbook as (its sheet, who delivers it): the export
    is My Plan's tables, a sheet each (owner, 2026-10-26)."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(response.content), read_only=True)
    found = []
    for sheet in book.worksheets:
        header, *rows = sheet.iter_rows(values_only=True)
        if "Executor" not in header:
            continue
        at = header.index("Executor")
        found.extend((sheet.title, row[at]) for row in rows)
    return sorted(found)


def _selected_kind(html: str) -> str:
    block = re.search(r'name="activity_type".*?</select>', html, re.S).group(0)
    return re.search(r'<option value="([a-z_]+)"\s+selected', block).group(1)


class TheExportFollowsThePagesTab(FilterExportFixture):
    def page(self, **params):
        return self.as_pl().get(PL_URL, {"owner": self.james.id, **params})

    def test_the_export_button_carries_the_tab(self):
        body = self.page(activity="visits").content.decode()
        button = re.search(r"<a[^>]*data-oversight-export[^>]*>", body, re.S).group(0)
        link, drawer = button.split("hx-get")
        self.assertIn("activity=visits", link)
        self.assertIn("activity=visits", drawer)

    def test_all_activities_adds_nothing(self):
        body = self.page().content.decode()
        button = re.search(r"<a[^>]*data-oversight-export[^>]*>", body, re.S).group(0)
        self.assertNotIn("activity=", button)

    def test_the_drawer_opens_on_the_tab(self):
        for tab, kind in (
            ("visits", "visits"),
            ("meetings", "meetings"),
            ("trainings", "trainings"),
            ("all", "planned"),
        ):
            with self.subTest(tab=tab):
                html = self.as_pl().get(DRAWER, {"activity": tab}).content.decode()
                self.assertEqual(_selected_kind(html), kind)

    def test_the_export_link_alone_holds_the_tabs_work(self):
        response = self.as_pl().get(
            PL_EXPORT,
            {"owner": self.james.id, "activity": "visits", "format": "xlsx"},
        )
        self.assertEqual(_rows(response), VISITS)

    def test_the_drawers_choice_wins_over_the_tab(self):
        response = self.as_pl().get(
            PL_EXPORT,
            {
                "owner": self.james.id,
                "activity": "visits",
                "activity_type": "all",
                "format": "xlsx",
            },
        )
        self.assertEqual(_rows(response), EVERY)

    def test_visits_only_from_the_drawer_holds_visits_alone(self):
        response = self.as_pl().get(
            PL_EXPORT,
            {
                "owner": self.james.id,
                "period_type": "fy",
                "fy": self.fy,
                "activity_type": "visits",
                "status": "all",
                "format": "xlsx",
            },
        )
        self.assertEqual(_rows(response), VISITS)


class TheDrawerKeepsThePagesFilters(FilterExportFixture):
    def test_the_pages_filters_ride_along_as_hidden_fields(self):
        html = (
            self.as_pl()
            .get(
                DRAWER,
                {
                    "owner": self.james.id,
                    "executor_type": "partner",
                    "district_id": self.district.id,
                },
            )
            .content.decode()
        )
        self.assertIn('name="executor_type" value="partner"', html)
        self.assertIn(f'name="district_id" value="{self.district.id}"', html)

    def test_a_partner_only_page_exports_partner_work_only(self):
        client = Client()
        client.force_login(self.pl_user)
        response = client.get(
            PL_EXPORT,
            {
                "owner": self.james.id,
                "executor_type": "partner",
                "period_type": "fy",
                "fy": self.fy,
                "activity_type": "planned",
                "status": "all",
                "format": "xlsx",
            },
        )
        self.assertEqual(_rows(response), [("School Visits", "Partner X")])
