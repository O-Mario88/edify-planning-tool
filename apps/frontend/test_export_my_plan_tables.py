"""Exports are My Plan's tables (owner, 2026-10-26).

"If the IA select all then it should group those activities in different
sheets in the excel exported sheet not everything in one table ... tables are
exported exactly the way they are in My Plan tables. I REPEAT the export should
be exact no inventing new columns or removing columns."

And the Export options drawer: Plan Status is "All scheduled plans" or
"Scheduled and completed plans"; Activity Type adds Schools with Planned
Training, Schools with Planned Cluster Meeting, Core School Visits, Client
School Visits and Core Schools with Planned Trainings. Asked, the owner chose:
every drawer offers them; Planning Oversight and My Plan export My Plan's
tables, a sheet each; the other pages keep their own layout and give the
matching My Plan table for the five; a cluster meeting's schools are one row a
school in the Trainings table's columns.
"""

from __future__ import annotations

import io
import re
from datetime import date, timedelta
from html.parser import HTMLParser
from pathlib import Path

from django.conf import settings
from django.test import Client

from apps.activities.models import Activity, ClusterActivityAttendance
from apps.core import export_kinds
from apps.core.rbac import EdifyRole
from apps.frontend import test_export_period_and_sessions as sessions
from apps.my_plan import table_export
from apps.schools.models import School

PL_EXPORT = "/team-planning-oversight/export"
COUNTRY_EXPORT = "/country-planning-oversight/export"
TEMPLATES = Path(settings.BASE_DIR) / "templates/partials/my_plan"


def book(response) -> dict[str, list[list]]:
    from openpyxl import load_workbook

    loaded = load_workbook(io.BytesIO(response.content))
    return {
        sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
        for sheet in loaded.worksheets
    }


def records(sheet: list[list]) -> list[dict]:
    header, *rows = sheet
    return [dict(zip(header, row)) for row in rows]


def template_headings(name: str) -> list[str]:
    """The headings a My Plan table draws: its <th> cells, the switched-off
    ones (inside {# #}) and the Actions column left out."""
    source = re.sub(r"\{#.*?#\}", "", (TEMPLATES / f"{name}.html").read_text())
    source = source.split("<thead", 1)[1].split("</thead>", 1)[0]
    headings = [
        re.sub(r"<[^>]+>", "", cell).strip()
        for cell in re.findall(r"<th[^>]*>(.*?)</th>", source, re.S)
    ]
    return [h for h in headings if h and h != "Actions"]


class _Tables(HTMLParser):
    """Every <table> of a page: its headings and each row's cell texts. A
    cell holding a tick box is the row picker, not a column."""

    def __init__(self):
        super().__init__()
        self.tables: list[dict] = []
        self._cell = None
        self._row = None
        self._picker = False

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append({"headings": [], "rows": []})
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            classes = dict(attrs).get("class") or ""
            self._cell = []
            self._picker = "school-plan-table__select" in classes
        elif tag == "input" and self._cell is not None:
            self._picker = True

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            text = " ".join(" ".join(self._cell).split())
            if not self._picker:
                self._row.append((tag, text))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if all(kind == "th" for kind, _ in self._row) and self._row:
                self.tables[-1]["headings"] = [text for _, text in self._row]
            elif self._row:
                self.tables[-1]["rows"].append([text for _, text in self._row])
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def shown(value, number_format: str) -> str:
    """A workbook cell as the page writes it."""
    if isinstance(value, (date,)):
        if number_format == table_export.SHORT_DAY:
            return f"{value:%b} {value.day}, {value:%a}"
        return f"{value.day} {value:%b %Y}"
    if isinstance(value, int) and number_format == table_export.MONEY:
        return f"UGX {value:,}"
    return "" if value is None else str(value)


class MyPlanTablesFixture(sessions.SessionFixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.ia_user, _ = cls._staff("ia@t.test", "Impact", EdifyRole.IMPACT_ASSESSMENT)
        soon = date.today() + timedelta(days=5)
        cls.core = School.objects.create(
            school_id="core1",
            name="Core Hill",
            school_type="core",
            district=cls.district,
            region=cls.region,
        )
        cls.core_visit = Activity.objects.create(
            activity_type="core_visit",
            school=cls.core,
            fy=cls.fy,
            quarter="Q1",
            planned_date=soon,
            planned_month=soon.month,
            status="scheduled",
            delivery_type="staff",
            responsible_staff_id=cls.james.id,
        )
        cls.meeting = Activity.objects.create(
            activity_type="cluster_meeting",
            cluster=cls.cluster,
            fy=cls.fy,
            quarter="Q1",
            planned_date=soon,
            planned_month=soon.month,
            status="scheduled",
            delivery_type="staff",
            responsible_staff_id=cls.james.id,
            participants_per_school=2,
        )
        for school in cls.members[:2]:
            ClusterActivityAttendance.objects.create(
                activity=cls.meeting, school=school, invited=True
            )
        cls.done = Activity.objects.create(
            activity_type="school_visit",
            school=cls.school,
            fy=cls.fy,
            quarter="Q1",
            planned_date=soon,
            planned_month=soon.month,
            status="completed",
            delivery_type="staff",
            responsible_staff_id=cls.james.id,
        )

    def export(self, path=PL_EXPORT, client=None, **params):
        client = client or self.as_pl()
        query = {"fy": self.fy, "format": "xlsx", "status": "all", **params}
        if path == PL_EXPORT:
            query.setdefault("owner", self.james.id)
        response = client.get(path, query)
        self.assertEqual(response.status_code, 200)
        return book(response)

    def as_user(self, user):
        client = Client()
        client.force_login(user)
        return client


class TheDrawerOffersTheOwnersChoices(MyPlanTablesFixture):
    def test_plan_status_is_two_choices(self):
        self.assertEqual(
            list(export_kinds.STATUS_CHOICES),
            [
                ("scheduled", "All scheduled plans"),
                ("all", "Scheduled and completed plans"),
            ],
        )
        body = self.as_pl().get(f"{PL_EXPORT}-drawer", {"fy": self.fy})
        html = body.content.decode()
        status = html.split('name="status"', 1)[1].split("</select>", 1)[0]
        self.assertEqual(
            re.findall(r"<option[^>]*>([^<]+)</option>", status),
            ["All scheduled plans", "Scheduled and completed plans"],
        )

    def test_activity_type_adds_the_five_lists(self):
        labels = dict(export_kinds.KIND_CHOICES)
        for key, label in (
            ("training_schools", "Schools with Planned Training"),
            ("meeting_schools", "Schools with Planned Cluster Meeting"),
            ("core_visits", "Core School Visits"),
            ("client_visits", "Client School Visits"),
            ("core_trainings", "Core Schools with Planned Trainings"),
        ):
            with self.subTest(kind=key):
                self.assertEqual(labels[key], label)
        # Today's choices stay.
        for key in ("planned", "visits", "group_trainings", "meetings", "online"):
            self.assertIn(key, labels)

    def test_every_drawer_offers_them(self):
        client = self.as_pl()
        for path in (
            f"{PL_EXPORT}-drawer",
            "/partner-oversight/export-drawer",
            "/my-plan/export-drawer",
            "/budget/export-drawer",
            "/fund-requests/weekly/export-drawer",
        ):
            response = client.get(path, {"fy": self.fy})
            if response.status_code != 200:
                continue  # a page this role does not open
            with self.subTest(drawer=path):
                self.assertIn('value="core_trainings"', response.content.decode())


class AllIsEveryMyPlanTableOnItsOwnSheet(MyPlanTablesFixture):
    def test_each_table_is_a_sheet_in_my_plans_order(self):
        sheets = self.export(activity_type="all")
        self.assertEqual(
            [name for name in sheets if name in table_export.FIXED_TITLES],
            list(table_export.FIXED_TITLES),
        )

    def test_each_sheet_has_exactly_my_plans_columns(self):
        sheets = self.export(activity_type="all")
        for table in table_export.TABLES:
            with self.subTest(table=table.key):
                self.assertEqual(
                    sheets[table.title][0], template_headings(table.template)
                )
        self.assertEqual(
            sheets[table_export.MEETING_SCHOOLS_TITLE][0],
            template_headings("cluster_trainings"),
        )

    def test_the_rows_are_on_the_sheet_of_the_table_my_plan_lists_them_in(self):
        sheets = self.export(activity_type="all")
        ids = {
            title: {row.get("School ID") for row in records(sheets[title])}
            for title in sheets
        }
        self.assertIn("s1", ids["School Visits"])
        self.assertIn("core1", ids["Core School Visits"])
        self.assertNotIn("core1", ids["School Visits"])
        # A group training is one row a school on Trainings, one on Group
        # Trainings; a meeting one row on Cluster Meetings.
        self.assertTrue({"gt0", "gt1", "gt2", "gt3"} <= ids["Trainings"])
        self.assertEqual(len(records(sheets["Group Trainings"])), 1)
        self.assertEqual(len(records(sheets["Cluster Meetings"])), 1)

    def test_the_country_export_is_the_same_sheets(self):
        sheets = self.export(COUNTRY_EXPORT, self.as_user(self.ia_user))
        self.assertEqual(
            [name for name in sheets if name in table_export.FIXED_TITLES],
            list(table_export.FIXED_TITLES),
        )
        self.assertIn(
            "core1", {r["School ID"] for r in records(sheets["Core School Visits"])}
        )

    def test_the_csv_holds_the_same_tables_one_after_another(self):
        response = self.as_pl().get(
            PL_EXPORT,
            {"fy": self.fy, "owner": self.james.id, "format": "csv", "status": "all"},
        )
        body = b"".join(response.streaming_content).decode()
        for title in table_export.FIXED_TITLES:
            self.assertIn(title, body)


class EachListIsItsMyPlanTable(MyPlanTablesFixture):
    def test_schools_with_planned_training_is_the_trainings_table(self):
        sheets = self.export(activity_type="training_schools")
        self.assertEqual(list(sheets), ["Trainings"])
        # The group training's four schools, and the training at Alpha.
        self.assertEqual(
            sorted(r["School ID"] for r in records(sheets["Trainings"])),
            ["gt0", "gt1", "gt2", "gt3", "s1"],
        )

    def test_a_cluster_meetings_schools_are_one_row_each(self):
        sheets = self.export(activity_type="meeting_schools")
        self.assertEqual(list(sheets), [table_export.MEETING_SCHOOLS_TITLE])
        rows = records(sheets[table_export.MEETING_SCHOOLS_TITLE])
        self.assertEqual(sorted(r["School ID"] for r in rows), ["gt0", "gt1"])
        for row in rows:
            self.assertEqual(row["School / Cluster"][:7], "Member ")
            self.assertEqual(row["Cluster Planned From"], "Kasangati Cluster Invited")
            self.assertEqual(row["Mode of Delivery"], "Cluster Meeting")
            self.assertEqual(row["Participants"], 2)

    def test_core_and_client_school_visits(self):
        core = self.export(activity_type="core_visits")
        self.assertEqual(list(core), ["Core School Visits"])
        self.assertEqual(
            [r["School ID"] for r in records(core["Core School Visits"])], ["core1"]
        )
        client = self.export(activity_type="client_visits")
        self.assertEqual(list(client), ["School Visits"])
        self.assertEqual(
            {r["School ID"] for r in records(client["School Visits"])}, {"s1"}
        )

    def test_core_schools_with_planned_trainings_is_the_core_trainings_table(self):
        self.members[0].school_type = "core"
        self.members[0].save(update_fields=["school_type"])
        sheets = self.export(activity_type="core_trainings")
        self.assertEqual(list(sheets), ["Core School Trainings"])
        self.assertEqual(
            [r["School ID"] for r in records(sheets["Core School Trainings"])],
            ["gt0"],
        )


class PlanStatusDecidesTheRows(MyPlanTablesFixture):
    """Alpha has James's completed visit beside the open ones."""

    def days(self, status):
        sheet = self.export(activity_type="client_visits", status=status)
        return [
            r["Planned Date"].date()
            for r in records(sheet["School Visits"])
            if r["School ID"] == "s1"
        ]

    def test_all_scheduled_plans_leaves_completed_work_out(self):
        self.assertNotIn(self.done.planned_date, self.days("scheduled"))
        self.assertIn(self.visit.planned_date, self.days("scheduled"))

    def test_scheduled_and_completed_keeps_both(self):
        days = self.days("all")
        self.assertIn(self.done.planned_date, days)
        self.assertIn(self.visit.planned_date, days)


class TheCellsAreMyPlansCells(MyPlanTablesFixture):
    """What James's My Plan page draws, row for row and cell for cell, is
    what his export writes."""

    def test_every_row_my_plan_draws_is_exported_as_drawn(self):
        from openpyxl import load_workbook

        client = self.as_user(self.james_user)
        page = client.get("/my-plan/", {"fy": self.fy, "period": "fy"})
        self.assertEqual(page.status_code, 200)
        parser = _Tables()
        parser.feed(page.content.decode())
        response = client.get(
            "/my-plan/",
            {"fy": self.fy, "export": "xlsx", "status": "all", "activity_type": "all"},
        )
        loaded = load_workbook(io.BytesIO(response.content))
        compared = 0
        for table in table_export.TABLES:
            sheet = loaded[table.title]
            headers = [cell.value for cell in sheet[1]]
            exported = [
                [shown(cell.value, cell.number_format) for cell in row[: len(headers)]]
                for row in sheet.iter_rows(min_row=2)
            ]
            drawn = [
                t for t in parser.tables if t["headings"][: len(headers)] == headers
            ]
            # A group training's Cost on the export is its own meals, venue
            # and facilitation (owner, 2026-10-15, for every export); the
            # page writes its estimate. Every other cell is the page's.
            skip = {headers.index("Cost")} if table.key == "group_trainings" else set()

            def kept(cells, skip=skip):
                return [c for i, c in enumerate(cells) if i not in skip]

            exported = [kept(cells) for cells in exported]
            for page_table in drawn:
                for row in page_table["rows"]:
                    if len(row) < len(headers):
                        continue  # an empty table's one message row
                    with self.subTest(table=table.key, row=row[:2]):
                        self.assertIn(kept(row[: len(headers)]), exported)
                        compared += 1
        self.assertGreater(compared, 3)


class MyPlanExportsItsTables(MyPlanTablesFixture):
    def test_my_plan_is_its_tables_and_the_budget(self):
        client = self.as_user(self.james_user)
        response = client.get(
            "/my-plan/", {"fy": self.fy, "export": "xlsx", "status": "all"}
        )
        sheets = book(response)
        self.assertEqual(
            [name for name in sheets if name in table_export.FIXED_TITLES],
            list(table_export.FIXED_TITLES),
        )
        self.assertEqual(list(sheets)[-1], "Budget")


class TheOtherPagesGiveTheMatchingTable(MyPlanTablesFixture):
    def test_partner_oversight_budget_and_advance_requests(self):
        client = self.as_pl()
        for path, extra in (
            ("/partner-oversight/export", {"format": "xlsx"}),
            ("/budget/export", {}),
            ("/fund-requests/weekly", {"export": "xlsx"}),
        ):
            response = client.get(
                path,
                {
                    "fy": self.fy,
                    "period_type": "fy",
                    "activity_type": "client_visits",
                    **extra,
                },
            )
            if response.status_code in (302, 403):
                continue  # a page this role does not open
            with self.subTest(page=path):
                self.assertEqual(response.status_code, 200)
                self.assertEqual(list(book(response)), ["School Visits"])


OWNER_CORE_TRAINING_COLUMNS = [
    "School ID",
    "School Name",
    "Planned From",
    "Executor",
    "Training Name",
    "SSA Intervention",
    "Planned Date",
    "Completed Date",
    "Core Training Package",
    "Participants",
    "Delivery Type",
    "Cost",
    "Salesforce ID",
    "Status",
]


class CoreSchoolOversightHasTheCorePlans(MyPlanTablesFixture):
    """Owner, 2026-10-26: "Core schools oversight should have Core school
    Visits plans, Core schools with trainings planned" with the owner's
    columns for the trainings."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.members[0].school_type = "core"
        cls.members[0].save(update_fields=["school_type"])
        soon = date.today() + timedelta(days=7)
        Activity.objects.create(
            activity_type="in_school_training",
            school=cls.core,
            fy=cls.fy,
            quarter="Q1",
            planned_date=soon,
            planned_month=soon.month,
            status="scheduled",
            delivery_type="staff",
            responsible_staff_id=cls.james.id,
        )

    def tab(self, tab, **params):
        return self.as_user(self.ia_user).get(
            "/core-schools-oversight/", {"fy": self.fy, "tab": tab, **params}
        )

    def test_core_school_visits_is_my_plans_table(self):
        page = self.tab("core_visits")
        self.assertEqual(page.status_code, 200)
        self.assertIn('data-core-plan-row="core1"', page.content.decode())
        sheet = book(self.tab("core_visits", format="xlsx"))["Core School Visits"]
        self.assertEqual(sheet[0], template_headings("core_school_visits"))
        self.assertEqual([r["School ID"] for r in records(sheet)], ["core1"])

    def test_core_school_trainings_has_the_owners_columns(self):
        page = self.tab("core_trainings")
        self.assertEqual(page.status_code, 200)
        sheet = book(self.tab("core_trainings", format="xlsx"))["Core School Trainings"]
        self.assertEqual(sheet[0], OWNER_CORE_TRAINING_COLUMNS)
        rows = {r["School ID"]: r for r in records(sheet)}
        self.assertEqual(set(rows), {"core1", "gt0"})
        self.assertEqual(rows["core1"]["Planned From"], "In-school Visit")
        self.assertEqual(rows["core1"]["Delivery Type"], "In-school")
        self.assertEqual(rows["gt0"]["Planned From"], "Kasangati Cluster")
        self.assertEqual(rows["gt0"]["Delivery Type"], "Group Training")
        self.assertEqual(rows["gt0"]["Executor"], "James")


class EveryPlanStatusExports(MyPlanTablesFixture):
    """Both Plan Statuses, and the old links' two, answer on every export."""

    def test_each_status_on_each_export(self):
        client = self.as_user(self.ia_user)
        for path in (
            PL_EXPORT,
            COUNTRY_EXPORT,
            "/team-planning-oversight/projects/export",
            "/partner-oversight/export",
            "/budget/export",
        ):
            for status in ("scheduled", "all", "completed", "in_progress"):
                response = client.get(
                    path,
                    {
                        "fy": self.fy,
                        "period_type": "fy",
                        "status": status,
                        "format": "xlsx",
                    },
                )
                if response.status_code in (302, 403):
                    continue  # a page this role does not open
                with self.subTest(page=path, status=status):
                    self.assertEqual(response.status_code, 200)
