"""Exported text cannot run as a spreadsheet formula (owner, 2026-10-19).

"Protect against spreadsheet formula injection when exporting untrusted
text ... Do not prefix all IDs with arbitrary text if that would change the
authoritative identifier." A school name, purpose or note typed as a formula
is written as the text it is: in a workbook as a text cell with nothing
added, in a CSV with the leading apostrophe spreadsheets read as text.
Numbers, IDs and placeholders are left exactly as they were.
"""

from __future__ import annotations

import io

import openpyxl
from django.test import SimpleTestCase

from apps.core.excel import (
    STREAM_ROWS,
    safe_csv_writer,
    spreadsheet_text,
    workbook_response,
)

FORMULA = '=HYPERLINK("http://example.test","Open")'


class CsvCells(SimpleTestCase):
    def test_formula_shaped_text_is_neutralised(self):
        for value in (FORMULA, "+cmd", "@SUM(A1)", "-2+3", "\tx"):
            with self.subTest(value=value):
                self.assertEqual(spreadsheet_text(value), "'" + value)

    def test_numbers_ids_and_placeholders_are_left_alone(self):
        for value in ("-5", "-1,200", "35438", "SVE-00123", "-", "—", "", 12, None):
            with self.subTest(value=value):
                self.assertEqual(spreadsheet_text(value), value)

    def test_the_writer_applies_it_to_every_cell(self):
        stream = io.StringIO()
        safe_csv_writer(stream).writerow(["35438", FORMULA, 4])
        self.assertEqual(
            stream.getvalue().strip(), f"35438,\"'{FORMULA.replace('\"', '\"\"')}\",4"
        )


class WorkbookCells(SimpleTestCase):
    def _first_cell(self, rows):
        response = workbook_response(
            "t.xlsx", [{"title": "T", "headers": ["School ID", "Note"], "rows": rows}]
        )
        sheet = openpyxl.load_workbook(io.BytesIO(response.content))["T"]
        return sheet["A2"], sheet["B2"]

    def test_a_formula_is_stored_as_text_with_nothing_added(self):
        for count in (1, STREAM_ROWS + 1):
            with self.subTest(rows=count):
                school_id, note = self._first_cell([["35438", FORMULA]] * count)
                self.assertEqual(note.value, FORMULA)
                self.assertEqual(note.data_type, "s")
                self.assertEqual(school_id.value, "35438")
