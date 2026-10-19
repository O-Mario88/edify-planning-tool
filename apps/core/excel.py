"""One .xlsx writer, so every plan export looks like the same document.

Owner, 2026-09-22: "IA and PL and CD and Regional Programme Leads, and CCEO
should be able to export all of their plans into Excel."

Every one of those roles already had an export; four of the five got a CSV.
A CSV opens in Excel and then argues with it — dates read as text in one
locale and as numbers in another, a UGX figure loses its grouping, and the
column widths are whatever the reader drags them to. The Work Plan export had
already solved this with openpyxl, and the styling here is that export's,
lifted so the plan a Lead sends to a Country Director looks like the plan a
CCEO sends to a Lead.

The rows are not built here. Each export hands over exactly what its own page
rendered — the same list, in the same order — because an export that re-queries
is an export that can disagree with the screen it was taken from.
"""

from __future__ import annotations

from django.http import HttpResponse

#: The Work Plan export's palette, kept identical on purpose.
HEADER_FILL = "102A43"
HEADER_RULE = "1677FF"
ROW_RULE = "D9E2EC"
BAND = "F7FAFC"

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


#: Above this many rows in any sheet the workbook is streamed (see
#: `_streamed_workbook`): same look, a fraction of the memory.
STREAM_ROWS = 5_000

#: What a spreadsheet reads as the start of a formula (OWASP CSV injection).
#: A school name, a purpose or a note typed as "=HYPERLINK(...)" would run
#: when the export is opened (owner, 2026-10-19: "Protect against spreadsheet
#: formula injection when exporting untrusted text").
_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")


def _is_plain_number(text: str) -> bool:
    import re

    return bool(re.fullmatch(r"[+-]?\d[\d,]*(?:\.\d+)?", text))


def spreadsheet_text(value):
    """A CSV cell that cannot run as a formula: text that starts like one is
    written with a leading apostrophe, the form spreadsheets read as text. A
    plain signed number ("-5", "-1,200") and a lone "-" placeholder stay as
    they are; anything else is untouched."""
    if (
        isinstance(value, str)
        and value[:1] in _FORMULA_LEADS
        and value not in ("-", "+")
        and not _is_plain_number(value)
    ):
        return "'" + value
    return value


class _SafeCsvWriter:
    """``csv.writer`` whose rows pass through `spreadsheet_text`."""

    def __init__(self, writer):
        self._writer = writer

    def writerow(self, row):
        return self._writer.writerow([spreadsheet_text(value) for value in row])

    def writerows(self, rows):
        for row in rows:
            self.writerow(row)


def safe_csv_writer(stream, *args, **kwargs) -> _SafeCsvWriter:
    """The CSV writer every export uses: ``csv.writer`` with formula-shaped
    text neutralised."""
    import csv

    return _SafeCsvWriter(csv.writer(stream, *args, **kwargs))


def _keep_as_text(cell):
    """An .xlsx cell given text that starts with "=" is stored as a formula by
    openpyxl. It is stored as the text it is, as typed, with nothing added."""
    if getattr(cell, "data_type", None) == "f":
        cell.data_type = "s"
    return cell


def workbook_response(filename: str, sheets: list[dict]) -> HttpResponse:
    """A styled workbook as a download.

    Each sheet is ``{title, headers, rows, widths?, number_formats?,
    row_styles?, filter?, freeze?}``. ``number_formats`` maps a 1-based column
    index to an Excel format, for the money and count columns that otherwise
    arrive as bare integers. ``row_styles`` maps a 1-based row index to one of
    `LEDGER_STYLES`, for a sheet laid out as a sectioned ledger (the Budget's
    Activity cost plan: a band naming each activity type, that table's own
    column headings, its rows, a subtotal, a grand total) rather than one
    heading row over a flat table; such a sheet passes ``filter=False`` and
    ``freeze=None`` so the first row is not read as a filterable heading, and
    may pass ``align`` (1-based column index to "center" or "right") so its
    figures sit under their headings as they do on the page.
    """
    import datetime as _dt
    import openpyxl.cell.cell as _cell

    for _cls, _fmt in (
        (_dt.date, "yyyy-mm-dd"),
        (_dt.datetime, "yyyy-mm-dd h:mm:ss"),
        (_dt.time, "h:mm:ss"),
        (_dt.timedelta, "[hh]:mm:ss"),
    ):
        _cell.TIME_FORMATS[_cls] = _fmt
        _cell._TYPES[_cls] = "d"
        _now_d = type(_dt.date.today())
        _cell.TIME_FORMATS[_now_d] = "yyyy-mm-dd"
        _cell._TYPES[_now_d] = "d"
        _now_dt = type(_dt.datetime.now())
        _cell.TIME_FORMATS[_now_dt] = "yyyy-mm-dd h:mm:ss"
        _cell._TYPES[_now_dt] = "d"

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    if any(
        isinstance(spec.get("rows"), (list, tuple)) and len(spec["rows"]) > STREAM_ROWS
        for spec in sheets or []
    ):
        workbook = _streamed_workbook(sheets)
        response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        workbook.save(response)
        return response

    workbook = Workbook()
    workbook.remove(workbook.active)

    for spec in sheets or [{"title": "Export", "headers": [], "rows": []}]:
        # Excel refuses a sheet title over 31 characters or carrying []:*?/\\.
        sheet = workbook.create_sheet(_sheet_title(spec.get("title") or "Export"))
        headers = list(spec.get("headers") or [])
        if headers:
            sheet.append(headers)
        for row in spec.get("rows") or []:
            sheet.append(list(row))
            for cell in sheet[sheet.max_row]:
                _keep_as_text(cell)

        if headers:
            style_header(sheet)
            sheet.freeze_panes = spec.get("freeze", "A2")
            if sheet.max_row > 1 and spec.get("filter", True):
                sheet.auto_filter.ref = sheet.dimensions
        style_body(sheet, spec.get("number_formats") or {})
        style_ledger_rows(
            sheet,
            spec.get("row_styles") or {},
            spec.get("align") or {},
            spec.get("number_formats") or {},
        )

        widths = spec.get("widths") or _widths_for(headers)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.sheet_view.showGridLines = False

    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def table_download(request, stem: str, sheets: list[dict]) -> HttpResponse:
    """An export as a workbook, or as CSV for a link that asks for it.

    Owner, 2026-09-27: "It should be export in excel not csv since the team
    use excel more." A page's one Export button downloads ``<stem>.xlsx``;
    ``?format=csv`` still answers with the first sheet as ``<stem>.csv`` for
    anything that already links to the CSV.
    """
    if (request.GET.get("format") or "").strip().lower() == "csv":
        sheet = (sheets or [{}])[0]
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{stem}.csv"'
        writer = safe_csv_writer(response)
        if sheet.get("headers"):
            writer.writerow(sheet["headers"])
        for row in sheet.get("rows") or []:
            writer.writerow(row)
        return response
    return workbook_response(f"{stem}.xlsx", sheets)


def style_header(sheet) -> None:
    """The navy heading row every plan export opens with."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    fill = PatternFill("solid", fgColor=HEADER_FILL)
    font = Font(color="FFFFFF", bold=True, size=10)
    alignment = Alignment(vertical="center")
    border = Border(bottom=Side(style="medium", color=HEADER_RULE))
    for row in sheet.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment
            cell.border = border
    sheet.row_dimensions[1].height = 28


def style_body(sheet, number_formats: dict[int, str] | None = None) -> None:
    """Banded, ruled, top-aligned body rows, and each column's number format.

    One pass with `iter_rows`, and one style object per kind shared by every
    cell. `sheet[row_index]` recomputes the sheet's width on every call, so
    styling row by row that way was quadratic: 231 million iterations and
    11-15 s for a 4,000-row Work Plan at 50,000 schools (2026-09-24 audit).
    """
    from openpyxl.styles import Alignment, Border, PatternFill, Side
    from openpyxl.styles.cell_style import StyleArray

    fills = (
        PatternFill("solid", fgColor="FFFFFF"),
        PatternFill("solid", fgColor=BAND),
    )
    border = Border(bottom=Side(style="thin", color=ROW_RULE))
    alignment = Alignment(vertical="top", wrap_text=True)
    formats = number_formats or {}
    # Assigning a style through the cell looks it up in the workbook's style
    # list by hashing it, three times per cell: ~250,000 lookups and most of a
    # 4,000-row export's time (2026-09-24 A+ audit). The first cell of each
    # band still goes through the assignment, which registers the style in
    # the order it always did; every later cell takes the same list indices
    # directly, which is exactly what the assignment stores.
    fill_ids: dict[bool, int] = {}
    border_id = alignment_id = None
    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        banded = row_index % 2 == 0
        fill_id = fill_ids.get(banded)
        for cell in row:
            if fill_id is None:
                cell.fill = fills[banded]
                cell.border = border
                cell.alignment = alignment
                fill_id = fill_ids[banded] = cell._style.fillId
                border_id = cell._style.borderId
                alignment_id = cell._style.alignmentId
                continue
            style = cell._style
            if not style:
                # What openpyxl's style descriptor does for an unstyled cell.
                style = cell._style = StyleArray()
            style.fillId = fill_id
            style.borderId = border_id
            style.alignmentId = alignment_id
        for column, number_format in formats.items():
            sheet.cell(row_index, column).number_format = number_format


#: The Budget page's ledger, in Excel, in the page's light-theme colours:
#: the band that names each activity type and carries that table's column
#: headings (`.budget-ledger-band`: --edify-band-surface, --edify-band-ink),
#: the muted bold subtotal (`.budget-ledger-subtotal`: --edify-surface-muted),
#: the grand total, and the plain caption under the sheet's title.
LEDGER_BAND = "DCE6EC"
LEDGER_INK = "17232B"
LEDGER_SUBTOTAL = "E5EAED"
LEDGER_STYLES = ("title", "caption", "band", "columns", "subtotal", "total")


def ledger_styles() -> dict[str, dict]:
    """One style set per `LEDGER_STYLES` name, as keyword arguments a cell
    takes (`font`, `fill`, `border`, `alignment`)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    band = PatternFill("solid", fgColor=LEDGER_BAND)
    muted = PatternFill("solid", fgColor=LEDGER_SUBTOTAL)
    ink = Font(color=LEDGER_INK, bold=True, size=10)
    rule = Border(bottom=Side(style="thin", color=ROW_RULE))
    middle = Alignment(vertical="center")
    paper = PatternFill("solid", fgColor="FFFFFF")
    return {
        "title": {
            "font": Font(bold=True, size=13, color=HEADER_FILL),
            "fill": paper,
            "border": Border(),
            "alignment": middle,
        },
        "caption": {
            "font": Font(size=9, color="52606D"),
            "fill": paper,
            "border": Border(),
            "alignment": Alignment(vertical="center", wrap_text=False),
        },
        "band": {
            "font": Font(color=LEDGER_INK, bold=True, italic=True, size=10),
            "fill": band,
            "border": Border(),
            "alignment": middle,
        },
        "columns": {"font": ink, "fill": band, "border": Border(), "alignment": middle},
        "subtotal": {
            "font": Font(bold=True, size=10),
            "fill": muted,
            "border": rule,
            "alignment": middle,
        },
        "total": {
            "font": Font(bold=True, size=11, color=LEDGER_INK),
            "fill": band,
            "border": Border(),
            "alignment": middle,
        },
    }


#: Ledger rows whose text runs across the sheet, so no column alignment.
_LEDGER_RUNNING_TEXT = ("title", "caption")


def _aligned(alignment, horizontal):
    from openpyxl.styles import Alignment

    return Alignment(
        horizontal=horizontal,
        vertical=alignment.vertical if alignment else "top",
        wrap_text=alignment.wrap_text if alignment else None,
    )


def style_ledger_rows(
    sheet, row_styles: dict[int, str], align=None, number_formats=None
) -> None:
    """Restyle the named rows of a sectioned ledger after `style_body` has
    banded the sheet, and set each column's alignment; a few dozen rows, so
    plain cell assignment is fine."""
    align = align or {}
    if not row_styles and not align:
        return
    styles = ledger_styles()
    width = sheet.max_column
    for row_index in range(1, sheet.max_row + 1):
        name = row_styles.get(row_index)
        style = styles[name] if name else {}
        for column in range(1, width + 1):
            cell = sheet.cell(row_index, column)
            for attribute, value in style.items():
                setattr(cell, attribute, value)
            if column in align and name not in _LEDGER_RUNNING_TEXT:
                cell.alignment = _aligned(cell.alignment, align[column])
            # A headerless sheet's first row is not a body row to style_body.
            if name and column in (number_formats or {}):
                cell.number_format = number_formats[column]
        if name in ("band", "title", "total"):
            sheet.row_dimensions[row_index].height = 22


def _streamed_workbook(sheets: list[dict]):
    """The same workbook, written row by row instead of held in memory.

    An in-memory sheet keeps an object per cell: the Country Planning
    Oversight export at 50,000 schools is a million cells, ~500 MB and 12 s
    inside a web worker (2026-09-28). openpyxl's write-only mode streams each
    row to the file as it is appended; every cell carries the style indices
    `style_header` and `style_body` give an in-memory sheet — header, bands,
    rules, alignment, number formats — and the sheet keeps its widths, frozen
    heading, filter and hidden gridlines.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.styles.cell_style import StyleArray
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    for spec in sheets:
        sheet = workbook.create_sheet(_sheet_title(spec.get("title") or "Export"))
        headers = list(spec.get("headers") or [])
        rows = spec.get("rows") or []
        formats = spec.get("number_formats") or {}
        width = max([len(headers), *(len(row) for row in rows)] or [0])

        # Everything written ahead of the rows must be set before the first.
        widths = spec.get("widths") or _widths_for(headers)
        for index, column_width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = column_width
        sheet.sheet_view.showGridLines = False
        row_styles = spec.get("row_styles") or {}
        if headers:
            if spec.get("freeze", "A2"):
                sheet.freeze_panes = spec.get("freeze", "A2")
            sheet.row_dimensions[1].height = 28
        for row_index, name in row_styles.items():
            if name in ("band", "title", "total"):
                sheet.row_dimensions[row_index].height = 22

        def styled(value, **styles):
            cell = _keep_as_text(WriteOnlyCell(sheet, value))
            for name, style in styles.items():
                setattr(cell, name, style)
            return cell

        if headers:
            font = Font(color="FFFFFF", bold=True, size=10)
            fill = PatternFill("solid", fgColor=HEADER_FILL)
            alignment = Alignment(vertical="center")
            border = Border(bottom=Side(style="medium", color=HEADER_RULE))
            sheet.append(
                [
                    styled(
                        value, fill=fill, font=font, alignment=alignment, border=border
                    )
                    for value in headers
                ]
            )
            if rows and spec.get("filter", True):
                sheet.auto_filter.ref = f"A1:{get_column_letter(width)}{len(rows) + 1}"

        # Register each body style once and copy its indices into every cell,
        # as style_body does; a date keeps the format its value gave it.
        band = []
        for banded in (False, True):
            probe = styled(
                None,
                fill=PatternFill("solid", fgColor=BAND if banded else "FFFFFF"),
                border=Border(bottom=Side(style="thin", color=ROW_RULE)),
                alignment=Alignment(vertical="top", wrap_text=True),
            )
            band.append(
                (probe._style.fillId, probe._style.borderId, probe._style.alignmentId)
            )
        format_ids = {}
        for column, number_format in formats.items():
            probe = styled(None)
            probe.number_format = number_format
            format_ids[column] = probe._style.numFmtId

        ledger = ledger_styles() if row_styles else {}
        align = spec.get("align") or {}
        plain = {
            banded: {
                "fill": PatternFill("solid", fgColor=BAND if banded else "FFFFFF"),
                "border": Border(bottom=Side(style="thin", color=ROW_RULE)),
                "alignment": Alignment(vertical="top", wrap_text=True),
            }
            for banded in (False, True)
        }

        # The body starts on row 2; without a heading, row 1 is a body row
        # left unstyled, as style_body leaves it.
        for row_index, row in enumerate(rows, start=2 if headers else 1):
            values = list(row)
            values += [None] * (width - len(values))
            if row_index in row_styles or (align and row_index >= 2):
                # A ledger sheet is small: each cell takes its styles whole,
                # a named row its own look, every row its columns' formats
                # and alignment.
                name = row_styles.get(row_index)
                style = ledger[name] if name else plain[row_index % 2 == 0]
                cells = []
                for column, value in enumerate(values, start=1):
                    cell = styled(value, **style)
                    if column in formats:
                        cell.number_format = formats[column]
                    if column in align and name not in _LEDGER_RUNNING_TEXT:
                        cell.alignment = _aligned(cell.alignment, align[column])
                    cells.append(cell)
                sheet.append(cells)
                continue
            if row_index < 2:
                sheet.append(
                    [_keep_as_text(WriteOnlyCell(sheet, value)) for value in values]
                )
                continue
            fill_id, border_id, alignment_id = band[row_index % 2 == 0]
            cells = []
            for column, value in enumerate(values, start=1):
                cell = _keep_as_text(WriteOnlyCell(sheet, value))
                style = cell._style
                if style is None:  # unstyled until now, as in style_body
                    style = cell._style = StyleArray()
                style.fillId = fill_id
                style.borderId = border_id
                style.alignmentId = alignment_id
                if column in format_ids:
                    style.numFmtId = format_ids[column]
                cells.append(cell)
            sheet.append(cells)
    return workbook


def _sheet_title(title: str) -> str:
    for bad in "[]:*?/\\":
        title = title.replace(bad, " ")
    return (title.strip() or "Export")[:31]


def _widths_for(headers) -> list[int]:
    """A readable width from the header alone — wide enough to read, capped so
    one long heading does not push every other column off the screen."""
    return [min(42, max(14, len(str(header)) + 6)) for header in headers]
