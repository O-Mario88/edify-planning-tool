"""Every plan a role holds can leave the platform as a workbook.

Owner, 2026-09-22: "IA and PL and CD and Regional Programme Leads, and CCEO
should be able to export all of their plans into Excel."

Four of those five already had an export and got a CSV, which opens in Excel
and then argues with it about dates and money. The rows are unchanged — the
same `export_rows` the CSV has always used — so these tests hold the two things
that could go wrong: that the workbook carries exactly those rows, and that the
CSV every existing link points at still answers.
"""

from __future__ import annotations

import io

from django.test import SimpleTestCase

from apps.core.excel import XLSX_CONTENT_TYPE, workbook_response


def _sheets(response):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(response.content))
    return {
        name: [list(row) for row in book[name].iter_rows(values_only=True)]
        for name in book.sheetnames
    }


class WorkbookResponseTest(SimpleTestCase):
    def test_it_is_a_download_a_spreadsheet_will_open(self):
        response = workbook_response(
            "plan.xlsx", [{"title": "Plan", "headers": ["A"], "rows": [[1]]}]
        )
        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertIn('filename="plan.xlsx"', response["Content-Disposition"])

    def test_the_rows_arrive_exactly_as_they_were_handed_over(self):
        response = workbook_response(
            "plan.xlsx",
            [
                {
                    "title": "Plan",
                    "headers": ["School", "Cost"],
                    "rows": [["Hope Primary", 120000], ["Faith Junior", 90000]],
                }
            ],
        )
        self.assertEqual(
            _sheets(response)["Plan"],
            [
                ["School", "Cost"],
                ["Hope Primary", 120000],
                ["Faith Junior", 90000],
            ],
        )

    def test_a_plan_with_no_rows_is_still_a_readable_workbook(self):
        # An officer who has planned nothing yet must get a file with the
        # column headings, not a broken download.
        response = workbook_response(
            "plan.xlsx", [{"title": "Plan", "headers": ["School"], "rows": []}]
        )
        self.assertEqual(_sheets(response)["Plan"], [["School"]])

    def test_a_sheet_name_excel_would_refuse_is_made_acceptable(self):
        response = workbook_response(
            "plan.xlsx",
            [{"title": "Team/Plan: 2026 " + "x" * 40, "headers": ["A"], "rows": []}],
        )
        name = next(iter(_sheets(response)))
        self.assertLessEqual(len(name), 31)
        for bad in "[]:*?/\\":
            self.assertNotIn(bad, name)

    def test_several_sheets_keep_their_own_rows(self):
        response = workbook_response(
            "plan.xlsx",
            [
                {"title": "Plan", "headers": ["A"], "rows": [["one"]]},
                {"title": "Summary", "headers": ["B"], "rows": [["two"]]},
            ],
        )
        sheets = _sheets(response)
        self.assertEqual(sheets["Plan"], [["A"], ["one"]])
        self.assertEqual(sheets["Summary"], [["B"], ["two"]])


class TheExportsOfferBothFilesTest(SimpleTestCase):
    """The format is asked for in the URL, and CSV stays the default."""

    def test_the_oversight_exports_read_the_format_parameter(self):
        from django.test import RequestFactory

        from apps.frontend.views.oversight_views import _wants_excel

        rf = RequestFactory()
        for query, expected in (
            ("?format=xlsx", True),
            ("?format=excel", True),
            ("?format=XLSX", True),
            ("?format=csv", False),
            ("", False),
        ):
            with self.subTest(query=query):
                self.assertEqual(_wants_excel(rf.get(f"/export{query}")), expected)

    def test_my_plan_offers_both_files(self):
        from pathlib import Path

        page = (
            Path(__file__).resolve().parents[3] / "templates/pages/my_plan/index.html"
        ).read_text()
        self.assertIn("{% querystring export='csv' %}", page)
        self.assertIn("{% querystring export='xlsx' %}", page)

    def test_every_oversight_workspace_offers_both_files(self):
        from pathlib import Path

        root = Path(__file__).resolve().parents[3]
        for name in (
            "pl_workspace",  # Programme Lead, Impact Assessment, Regional Lead
            "team_country_workspace",
            "cd_workspace",  # Country Director
        ):
            source = (root / f"templates/partials/oversight/{name}.html").read_text()
            with self.subTest(workspace=name):
                self.assertIn("Export CSV", source)
                self.assertIn("Export Excel", source)
                self.assertIn("format=xlsx", source)


class WorkbookStylingCostTest(SimpleTestCase):
    """Styling a body is one pass, not one scan of the sheet per row (PERF-07).

    `sheet[row_index]` recomputes the sheet's width on every call, so styling
    row by row that way was quadratic: 231 million iterations and 11-15 s for
    a 4,000-row Work Plan at 50,000 schools (2026-09-24 audit).
    """

    def _width_scans(self, rows):
        from unittest import mock

        from openpyxl.worksheet.worksheet import Worksheet

        calls = []
        real = Worksheet.max_column.fget

        def counted(sheet):
            calls.append(1)
            return real(sheet)

        with mock.patch.object(Worksheet, "max_column", property(counted)):
            workbook_response(
                "plan.xlsx",
                [
                    {
                        "title": "Plan",
                        "headers": ["School", "Cost"],
                        "rows": [[f"School {n}", n] for n in range(rows)],
                    }
                ],
            )
        return len(calls)

    def test_the_width_is_not_rescanned_per_row(self):
        self.assertEqual(self._width_scans(20), self._width_scans(400))

    def test_the_body_keeps_its_banding_rules_and_formats(self):
        from openpyxl import load_workbook

        response = workbook_response(
            "plan.xlsx",
            [
                {
                    "title": "Plan",
                    "headers": ["School", "Cost"],
                    "rows": [["A", 1000], ["B", 2000], ["C", 3000]],
                    "number_formats": {2: "#,##0"},
                }
            ],
        )
        sheet = load_workbook(io.BytesIO(response.content))["Plan"]
        header, even, odd = sheet["A1"], sheet["A2"], sheet["A3"]
        self.assertEqual(header.fill.fgColor.rgb, "00102A43")
        self.assertTrue(header.font.b)
        self.assertEqual(header.border.bottom.style, "medium")
        self.assertEqual(even.fill.fgColor.rgb, "00F7FAFC")
        self.assertEqual(odd.fill.fgColor.rgb, "00FFFFFF")
        self.assertEqual(even.border.bottom.style, "thin")
        self.assertTrue(even.alignment.wrap_text)
        self.assertEqual(even.alignment.vertical, "top")
        self.assertEqual(sheet["B3"].number_format, "#,##0")
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_the_body_styling_writes_the_same_workbook_as_before(self):
        """The direct style indices save exactly what cell assignment saved.

        Styled through the cell, every body cell hashed three styles into the
        workbook's lists (2026-09-24 A+ audit); `style_body` now assigns once
        per band and copies the indices. This frozen copy of the old loop and
        the new one must save byte-identical sheets and styles, dates (which
        carry their own number format) and blank cells included.
        """
        import datetime
        import zipfile

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, PatternFill, Side

        from apps.core import excel

        def frozen_style_body(sheet, number_formats=None):
            fills = (
                PatternFill("solid", fgColor="FFFFFF"),
                PatternFill("solid", fgColor=excel.BAND),
            )
            border = Border(bottom=Side(style="thin", color=excel.ROW_RULE))
            alignment = Alignment(vertical="top", wrap_text=True)
            for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                fill = fills[row_index % 2 == 0]
                for cell in row:
                    cell.fill = fill
                    cell.border = border
                    cell.alignment = alignment
                for column, number_format in (number_formats or {}).items():
                    sheet.cell(row_index, column).number_format = number_format

        def build(style_body):
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["School", "Date", "Cost", "Note"])
            excel.style_header(sheet)
            for n in range(57):
                sheet.append(
                    [
                        f"School {n}",
                        datetime.date(2026, 1, 1) + datetime.timedelta(days=n),
                        n * 1000,
                        None if n % 5 else "note",
                    ]
                )
            style_body(sheet, {3: "#,##0"})
            out = io.BytesIO()
            workbook.save(out)
            with zipfile.ZipFile(out) as archive:
                return {
                    name: archive.read(name)
                    for name in archive.namelist()
                    if not name.startswith("docProps/")
                }

        self.assertEqual(build(excel.style_body), build(frozen_style_body))
