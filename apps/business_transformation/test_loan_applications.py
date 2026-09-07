from datetime import timedelta
from io import BytesIO

from django.db import IntegrityError, transaction
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import (
    StaffGeographyAssignment,
    StaffProfile,
    StaffSchoolAssignment,
    User,
)
from apps.core.exceptions import BadRequest
from apps.core.rbac import EdifyRole
from apps.core.throttling import reset_throttle_state
from apps.geography.models import Region
from apps.notifications.models import Notification
from apps.schools.models import School

from .loan_tracking import (
    send_application_follow_up_reminders,
    send_monthly_rvp_report,
    send_repayment_follow_up_reminders,
    update_application,
)
from .models import (
    LoanApplication,
    LoanApplicationStatus,
    LoanPurpose,
    LoanTrackingDeliveryLog,
    LoanStatus,
    MfiLoan,
    MfiOrganization,
    RepaymentTransaction,
    RepaymentTransactionKind,
    TransformationCase,
)


class LoanApplicationWorkflowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.region = Region.objects.create(
            name="Loan Test Central Region", country="Uganda"
        )
        cls.other_region = Region.objects.create(
            name="Loan Test Eastern Region", country="Uganda"
        )
        cls.school = School.objects.create(
            school_id="UG-LOAN-001",
            name="Bright Future School",
            region=cls.region,
        )
        cls.other_school = School.objects.create(
            school_id="UG-LOAN-002",
            name="Another School",
            region=cls.other_region,
        )
        cls.purpose = LoanPurpose.objects.create(
            code="LOAN_APPLICATION_TEST",
            label="Classroom improvement",
            active=True,
        )
        cls.mfi = MfiOrganization.objects.create(
            code="LOAN-TEST-MFI", name="Test Lending Partner", active=True
        )

        def staff(role, prefix):
            user = User.objects.create_user(
                email=f"{prefix}@example.org",
                name=prefix.title(),
                roles=[role.value],
                active_role=role.value,
            )
            profile = StaffProfile.objects.create(
                user=user, country="Uganda", onboarding_state="active"
            )
            return user, profile

        cls.cceo, cls.cceo_profile = staff(EdifyRole.CCEO, "loan-owner")
        cls.bt, cls.bt_profile = staff(
            EdifyRole.BUSINESS_TRANSFORMATION_OFFICER, "loan-bt"
        )
        cls.cd, cls.cd_profile = staff(EdifyRole.COUNTRY_DIRECTOR, "loan-cd")
        cls.ia, cls.ia_profile = staff(EdifyRole.IMPACT_ASSESSMENT, "loan-ia")
        cls.rvp, cls.rvp_profile = staff(EdifyRole.REGIONAL_VICE_PRESIDENT, "loan-rvp")
        cls.school.account_owner_id = cls.cceo_profile.id
        cls.school.account_owner_name_raw = cls.cceo.name
        cls.school.save(update_fields=["account_owner_id", "account_owner_name_raw"])
        StaffSchoolAssignment.objects.create(
            staff=cls.cceo_profile, school_id=cls.school.id
        )
        StaffGeographyAssignment.objects.create(
            staff=cls.rvp_profile, region_id=cls.region.id
        )

    def setUp(self):
        reset_throttle_state(["public.loan_application:127.0.0.1"])

    def _payload(self):
        return {
            "school_id": self.school.school_id,
            "school_name": self.school.name,
            "applicant_name": "Grace Headteacher",
            "applicant_role": "Headteacher",
            "applicant_phone": "+256700000001",
            "applicant_email": "grace@example.org",
            "purpose": self.purpose.id,
            "preferred_mfi": self.mfi.id,
            "requested_amount": "15000000",
            "intended_use": "Complete two classrooms.",
            "requested_term_months": "12",
            "repayment_frequency": "monthly",
            "consent": "on",
        }

    def _application(self, **overrides):
        values = {
            "school": self.school,
            "purpose": self.purpose,
            "preferred_mfi": self.mfi,
            "applicant_name": "Grace Headteacher",
            "applicant_role": "Headteacher",
            "applicant_phone": "+256700000001",
            "requested_amount": 15000000,
            "intended_use": "Complete two classrooms.",
            "requested_term_months": 12,
            "repayment_frequency": "monthly",
            "consent_recorded_at": timezone.now(),
            "submitted_at": timezone.now(),
            "next_follow_up_on": timezone.localdate(),
        }
        values.update(overrides)
        return LoanApplication.objects.create(**values)

    def test_school_can_submit_and_the_four_responsible_roles_are_notified(self):
        response = self.client.post("/loan-applications/apply", self._payload())

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Application received")
        application = LoanApplication.objects.get()
        self.assertEqual(application.school, self.school)
        self.assertEqual(application.status, LoanApplicationStatus.SUBMITTED)
        self.assertEqual(application.follow_ups.count(), 1)
        self.assertSetEqual(
            set(
                Notification.objects.filter(
                    source_event_type="bt.loan_application.submitted"
                ).values_list("recipient_id", flat=True)
            ),
            {self.cceo.id, self.bt.id, self.cd.id, self.ia.id},
        )

    def test_school_identity_must_match_and_an_open_duplicate_is_blocked(self):
        bad = self._payload()
        bad["school_name"] = "Different name"
        response = self.client.post("/loan-applications/apply", bad)
        self.assertContains(response, "do not match our records")
        self.assertFalse(LoanApplication.objects.exists())

        self._application()
        response = self.client.post("/loan-applications/apply", self._payload())
        self.assertContains(response, "already has an open loan application")
        self.assertEqual(LoanApplication.objects.count(), 1)
        with self.assertRaises(IntegrityError), transaction.atomic():
            self._application(applicant_phone="+256700000099")

    def test_cceo_sees_only_assigned_school_and_can_record_follow_up(self):
        mine = self._application()
        self._application(school=self.other_school, applicant_phone="+256700000002")
        self.client.force_login(self.cceo)

        page = self.client.get("/loans")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, self.school.name)
        self.assertNotContains(page, self.other_school.name)

        next_date = timezone.localdate() + timedelta(days=5)
        response = self.client.post(
            f"/loan-applications/{mine.id}/follow-up",
            {
                "status": LoanApplicationStatus.UNDER_REVIEW,
                "note": "Called the headteacher and confirmed the documents needed.",
                "nextFollowUpOn": next_date.isoformat(),
            },
        )
        self.assertRedirects(response, "/loans")
        mine.refresh_from_db()
        self.assertEqual(mine.status, LoanApplicationStatus.UNDER_REVIEW)
        self.assertEqual(mine.next_follow_up_on, next_date)
        self.assertEqual(mine.follow_ups.count(), 1)

    def test_due_application_reminders_are_permanently_idempotent(self):
        application = self._application(
            next_follow_up_on=timezone.localdate() - timedelta(days=2)
        )

        first = send_application_follow_up_reminders()
        second = send_application_follow_up_reminders()

        self.assertEqual(first, 4)
        self.assertEqual(second, 0)
        self.assertEqual(
            LoanTrackingDeliveryLog.objects.filter(
                delivery_kind="application_follow_up", subject_id=application.id
            ).count(),
            4,
        )

    def test_application_lifecycle_rejects_skips_and_locks_terminal_state(self):
        application = self._application()
        with self.assertRaisesMessage(BadRequest, "cannot move directly"):
            update_application(
                application.id,
                {
                    "status": LoanApplicationStatus.APPROVED,
                    "note": "Attempted to skip review and lender referral.",
                    "nextFollowUpOn": (
                        timezone.localdate() + timedelta(days=2)
                    ).isoformat(),
                },
                self.cceo,
            )

        application.status = LoanApplicationStatus.APPROVED
        application.save(update_fields=["status", "updated_at"])
        case = TransformationCase.objects.create(
            school=self.school,
            opened_fy="2026",
            owner_staff_id=self.cceo_profile.id,
        )
        loan = MfiLoan.objects.create(
            mfi=self.mfi,
            school=self.school,
            case=case,
            purpose=self.purpose,
            external_loan_reference="conversion-loan",
            status=LoanStatus.PROCESSING,
            registered_by=self.bt.id,
        )
        update_application(
            application.id,
            {
                "status": LoanApplicationStatus.CONVERTED,
                "note": "Matched to the governed lender record.",
                "linkedLoanId": loan.id,
            },
            self.cceo,
        )
        application.refresh_from_db()
        self.assertEqual(application.linked_loan, loan)
        self.assertIsNotNone(application.processed_at)
        with self.assertRaisesMessage(BadRequest, "cannot move directly"):
            update_application(
                application.id,
                {
                    "status": LoanApplicationStatus.UNDER_REVIEW,
                    "note": "Attempted reopen.",
                    "nextFollowUpOn": timezone.localdate().isoformat(),
                },
                self.cceo,
            )

    def test_monthly_and_termly_loans_generate_frequency_based_reminders(self):
        today = timezone.localdate()
        for frequency, months in (("monthly", 1), ("termly", 3)):
            school = (
                self.school
                if frequency == "monthly"
                else School.objects.create(
                    school_id="UG-LOAN-TERMLY",
                    name="Termly School",
                    account_owner_id=self.cceo_profile.id,
                )
            )
            case = TransformationCase.objects.create(
                school=school, opened_fy="2026", owner_staff_id=self.cceo_profile.id
            )
            disbursed_on = today
            for _ in range(months):
                disbursed_on = (
                    disbursed_on.replace(day=1) - timedelta(days=1)
                ).replace(day=min(today.day, 28))
            MfiLoan.objects.create(
                mfi=self.mfi,
                school=school,
                case=case,
                purpose=self.purpose,
                external_loan_reference=f"{frequency}-loan",
                disbursed_amount=15000000,
                disbursement_date=disbursed_on,
                disbursement_confirmed_at=timezone.now(),
                term_months=12,
                repayment_frequency=frequency,
                status=LoanStatus.ACTIVE,
                registered_by=self.bt.id,
            )

        first = send_repayment_follow_up_reminders(today)
        second = send_repayment_follow_up_reminders(today)

        self.assertEqual(first, 8)
        self.assertEqual(second, 0)
        self.assertEqual(
            Notification.objects.filter(
                source_event_type="bt.loan_repayment.follow_up"
            ).count(),
            8,
        )

    def test_rvp_gets_one_monthly_report_and_can_export_excel(self):
        report_date = timezone.localdate().replace(day=1) - timedelta(days=1)
        submitted_at = timezone.now().replace(
            year=report_date.year, month=report_date.month, day=report_date.day
        )
        self._application(
            applicant_name='=HYPERLINK("https://invalid","Open")',
            submitted_at=submitted_at,
        )
        self._application(
            school=self.other_school,
            applicant_phone="+256700000002",
            submitted_at=submitted_at,
        )
        case = TransformationCase.objects.create(
            school=self.school,
            opened_fy="2026",
            owner_staff_id=self.cceo_profile.id,
        )
        loan = MfiLoan.objects.create(
            mfi=self.mfi,
            school=self.school,
            case=case,
            purpose=self.purpose,
            external_loan_reference="monthly-report-loan",
            disbursed_amount=15000000,
            disbursement_date=report_date,
            disbursement_confirmed_at=timezone.now(),
            term_months=12,
            repayment_frequency="monthly",
            status=LoanStatus.ACTIVE,
            registered_by=self.bt.id,
        )
        RepaymentTransaction.objects.create(
            loan=loan,
            kind=RepaymentTransactionKind.PAYMENT,
            external_reference="PAY-REPORT-001",
            idempotency_key="pay-report-001",
            amount=500000,
            received_on=report_date,
            value_date=report_date,
            evidence_reference="BANK-REPORT-001",
            posted_by=self.bt.id,
            posted_at=timezone.now(),
        )
        first = send_monthly_rvp_report()
        second = send_monthly_rvp_report()
        self.assertEqual(first, 1)
        self.assertEqual(second, 0)
        notice = Notification.objects.get(source_event_type="bt.loan_report.monthly")
        self.assertEqual(notice.recipient_id, self.rvp.id)
        self.assertIn("1 application(s)", notice.body)
        self.assertIn("/loans?report_month=", notice.target_route)

        self.client.force_login(self.rvp)
        response = self.client.get(
            f"/loans/export.xlsx?report_month={report_date:%Y-%m}"
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Applications", "Processed loans", "Repayments"],
        )
        self.assertEqual(workbook["Applications"]["A1"].value, "Reference")
        self.assertTrue(workbook["Applications"]["G2"].value.startswith("'="))
        self.assertEqual(workbook["Applications"]["G2"].data_type, "s")
        application_schools = {
            row[3].value for row in workbook["Applications"].iter_rows(min_row=2)
        }
        self.assertSetEqual(application_schools, {self.school.name})
        self.assertEqual(workbook["Repayments"]["G2"].value, "PAY-REPORT-001")
        self.assertEqual(workbook["Repayments"]["I2"].value, 500000)
