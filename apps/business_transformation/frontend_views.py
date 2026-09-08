import csv
from datetime import timedelta
from urllib.parse import urlencode

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_GET, require_POST, require_http_methods

from apps.core.permissions import (
    has_permission,
    render_access_denied,
    require_page_permission,
)
from apps.core.rbac import Permission
from apps.core.exceptions import BadRequest

from . import services
from .loan_forms import PublicLoanApplicationForm
from .loan_tracking import scoped_applications, submit_application, update_application


def _topbar_search(request, form_id: str, placeholder: str) -> dict:
    """Bind the shell's one top-bar search to this page's filter form.

    The pages carry no body search of their own (the search-consolidation
    guard holds that line); the top-bar input joins the filter form via the
    HTML5 form= attribute, so q submits and swaps with the other filters.
    """
    return {
        "attach_to": form_id,
        "input_id": f"{form_id}-q",
        "name": "q",
        "value": (request.GET.get("q") or "").strip(),
        "placeholder": placeholder,
    }


def _filters(request) -> dict:
    return {
        key: (request.GET.get(key) or "").strip()
        for key in (
            "fy",
            "q",
            "region",
            "district",
            "mfi",
            "status",
            "period_type",
            "quarter",
            "month",
            "custom_from",
            "custom_to",
            "school",
            "purpose",
            "edtech",
            "repayment_health",
            "salesforce_status",
            "ia_status",
            "impact_status",
            "case_status",
            "partner",
            "support_status",
            "page",
            "per_page",
        )
    }


def _loan_filters(request) -> dict:
    """Return only filters represented by controls on the loan dashboard."""

    return {
        key: (request.GET.get(key) or "").strip()
        for key in (
            "fy",
            "q",
            "mfi",
            "district",
            "status",
            "repayment_health",
            "salesforce_status",
            "application_status",
            "report_month",
            "page",
            "per_page",
        )
    }


@require_GET
@require_page_permission("business_transformation")
def workspace_page(request):
    context = services.workspace_context(request.user, _filters(request))
    context["filter_query"] = request.GET.urlencode()
    context["topbar_search"] = _topbar_search(
        request, "bt-workspace-filters", "Search Uganda portfolio"
    )
    return render(request, "pages/business_transformation/index.html", context)


@require_GET
@require_page_permission("business_transformation_finance")
def financial_health_page(request):
    context = services.financial_health_context(request.user, _filters(request))
    context["topbar_search"] = _topbar_search(
        request, "bt-school-portfolio-filters", "Name, ID, district or owner"
    )
    context["base_template"] = (
        "layouts/blank.html"
        if request.headers.get("HX-Request") == "true"
        else "layouts/shell.html"
    )
    return render(
        request,
        "pages/business_transformation/financial_health.html",
        context,
    )


@require_GET
@require_page_permission("business_transformation_government")
def government_requirements_page(request):
    context = services.government_requirements_context(request.user, _filters(request))
    context["topbar_search"] = _topbar_search(
        request, "bt-school-portfolio-filters", "Name, ID, district or owner"
    )
    context["base_template"] = (
        "layouts/blank.html"
        if request.headers.get("HX-Request") == "true"
        else "layouts/shell.html"
    )
    return render(
        request,
        "pages/business_transformation/government_requirements.html",
        context,
    )


@require_GET
@require_page_permission("business_transformation_reports")
def impact_reports_page(request):
    return render(
        request,
        "pages/business_transformation/impact_reports.html",
        services.impact_reports_context(request.user, _filters(request)),
    )


@require_GET
@require_page_permission("mfi_portal")
def mfi_portal_page(request, section: str = "dashboard"):
    return render(
        request,
        "pages/business_transformation/mfi_portal.html",
        services.mfi_portal_context(request.user, section, _filters(request)),
    )


@require_GET
@require_page_permission("loans")
def loan_page(request):
    loan_filters = _loan_filters(request)
    context = services.loan_register_context(request.user, loan_filters)
    context["filter_query"] = urlencode(
        {
            key: value
            for key, value in loan_filters.items()
            if value and key not in {"page", "per_page"}
        }
    )
    context["topbar_search"] = _topbar_search(
        request, "bt-loan-filters", "Search Uganda loans"
    )
    context["public_application_url"] = request.build_absolute_uri(
        "/loan-applications/apply"
    )
    return render(request, "pages/business_transformation/loans.html", context)


@require_http_methods(["GET", "POST"])
def public_loan_application_page(request):
    submitted = False
    reference = ""
    throttle_error = ""
    form = PublicLoanApplicationForm(request.POST or None)
    throttled = False
    if request.method == "POST":
        from apps.core.throttling import throttle_by_ip

        throttled = not throttle_by_ip(
            request,
            name="public.loan_application",
            limit=10,
            window_ms=60 * 60 * 1000,
        )
        if throttled:
            throttle_error = (
                "Too many applications have been submitted from this connection. "
                "Try again later or contact your Edify portfolio owner."
            )
    if request.method == "POST" and not throttled and form.is_valid():
        try:
            application = submit_application(form)
        except BadRequest as exc:
            form.add_error(None, str(exc))
        else:
            submitted = True
            reference = str(application.id)
            form = PublicLoanApplicationForm()
    return render(
        request,
        "pages/business_transformation/loan_application_public.html",
        {
            "form": form,
            "submitted": submitted,
            "reference": reference,
            "throttle_error": throttle_error,
        },
        status=429 if throttled else 200,
    )


@require_POST
@require_page_permission("loans")
def loan_application_follow_up_action(request, application_id: str):
    update_application(application_id, request.POST.dict(), request.user)
    messages.success(request, "Loan application follow-up recorded.")
    return redirect("/loans")


@require_GET
@require_page_permission("loans")
def loan_drawer(request, loan_id: str):
    loan = services.scoped_loans(request.user).filter(id=loan_id).first()
    if loan is None:
        return render_access_denied(
            request, "This loan is outside your governed scope."
        )
    return render(
        request,
        "partials/business_transformation/loan_drawer.html",
        {
            "loan": loan,
            "can_confirm_salesforce": has_permission(
                request.user,
                Permission.BUSINESS_TRANSFORMATION_SALESFORCE_CONFIRM.value,
            ),
            "can_validate_ia": has_permission(
                request.user, Permission.BUSINESS_TRANSFORMATION_IA_VALIDATE.value
            ),
        },
    )


def _spreadsheet_safe(value) -> str:
    text = "" if value is None else str(value)
    candidate = text.lstrip()
    dangerous_prefixes = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")
    return f"'{text}" if candidate.startswith(dangerous_prefixes) else text


def _csv_safe(value) -> str:
    return _spreadsheet_safe(value)


def _xlsx_row(values):
    return [
        _spreadsheet_safe(value) if isinstance(value, str) else value
        for value in values
    ]


@require_GET
@require_page_permission("loans")
def loan_export_action(request):
    if not has_permission(
        request.user, Permission.BUSINESS_TRANSFORMATION_EXPORT.value
    ):
        return render_access_denied(
            request, "Your active role cannot export the Uganda loan register."
        )
    loans = services.loan_export_rows(request.user, _loan_filters(request))
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = (
        'attachment; filename="uganda-governed-loan-register.csv"'
    )
    writer = csv.writer(response)
    writer.writerow(
        [
            "Lending partner",
            "School ID",
            "School",
            "Current enrolment",
            "Loan purpose",
            "MFI loan reference",
            "Salesforce Loan ID",
            "Salesforce status",
            "Salesforce confirmed at",
            "Loan status",
            "IA validation",
            "Impact status",
            "Approved amount (UGX)",
            "Disbursed amount (UGX)",
            "Disbursement date",
            "Term months",
            "Repayment frequency",
            "Last repayment data date",
        ]
    )
    for loan in loans.iterator(chunk_size=500):
        writer.writerow(
            [
                _csv_safe(loan.mfi.name),
                _csv_safe(loan.school.school_id),
                _csv_safe(loan.school.name),
                loan.school.enrollment or "",
                _csv_safe(loan.purpose.label),
                _csv_safe(loan.external_loan_reference),
                loan.salesforce_loan_id or "",
                loan.get_salesforce_status_display(),
                loan.salesforce_confirmed_at.isoformat()
                if loan.salesforce_confirmed_at
                else "",
                loan.get_status_display(),
                loan.get_ia_validation_status_display(),
                loan.get_impact_status_display(),
                loan.approved_amount or "",
                loan.disbursed_amount or "",
                loan.disbursement_date.isoformat() if loan.disbursement_date else "",
                loan.term_months or "",
                _csv_safe(loan.repayment_frequency),
                loan.last_repayment_data_date.isoformat()
                if loan.last_repayment_data_date
                else "",
            ]
        )
    return response


def _application_export_rows(principal, filters):
    rows = scoped_applications(principal)
    if filters.get("application_status"):
        rows = rows.filter(status=filters["application_status"])
    report_month = filters.get("report_month") or ""
    if len(report_month) == 7:
        start = parse_date(f"{report_month}-01")
        if start:
            end = (start + timedelta(days=32)).replace(day=1)
            rows = rows.filter(
                submitted_at__date__gte=start, submitted_at__date__lt=end
            )
    return rows.order_by("-submitted_at", "school__name")


@require_GET
@require_page_permission("loans")
def loan_application_export_action(request):
    if not has_permission(
        request.user, Permission.BUSINESS_TRANSFORMATION_EXPORT.value
    ):
        return render_access_denied(
            request, "Your active role cannot export loan applications."
        )
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = (
        'attachment; filename="uganda-loan-applications.csv"'
    )
    writer = csv.writer(response)
    writer.writerow(
        [
            "Application reference",
            "Submitted",
            "School ID",
            "School",
            "District",
            "Portfolio owner",
            "Applicant",
            "Applicant role",
            "Phone",
            "Email",
            "Purpose",
            "Preferred lender",
            "Requested amount (UGX)",
            "Term months",
            "Repayment frequency",
            "Status",
            "Next follow-up",
        ]
    )
    for application in _application_export_rows(
        request.user, _loan_filters(request)
    ).iterator(chunk_size=500):
        writer.writerow(
            [
                application.id,
                application.submitted_at.isoformat(),
                _csv_safe(application.school.school_id),
                _csv_safe(application.school.name),
                _csv_safe(application.school.district.name)
                if application.school.district
                else "",
                _csv_safe(application.school.account_owner_name_raw),
                _csv_safe(application.applicant_name),
                _csv_safe(application.applicant_role),
                _csv_safe(application.applicant_phone),
                _csv_safe(application.applicant_email),
                _csv_safe(application.purpose.label),
                _csv_safe(application.preferred_mfi.name)
                if application.preferred_mfi
                else "",
                application.requested_amount,
                application.requested_term_months,
                application.get_repayment_frequency_display(),
                application.get_status_display(),
                application.next_follow_up_on.isoformat()
                if application.next_follow_up_on
                else "",
            ]
        )
    return response


@require_GET
@require_page_permission("loans")
def loan_excel_export_action(request):
    if not has_permission(
        request.user, Permission.BUSINESS_TRANSFORMATION_EXPORT.value
    ):
        return render_access_denied(
            request, "Your active role cannot export the Uganda loan report."
        )
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .loan_tracking import monthly_report_summary
    from .models import RepaymentTransaction, RepaymentTransactionKind

    filters = _loan_filters(request)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    report_month = filters.get("report_month") or ""
    report_start = parse_date(f"{report_month}-01") if len(report_month) == 7 else None
    report_end = (
        (report_start + timedelta(days=32)).replace(day=1) if report_start else None
    )
    summary_sheet.append(["Monthly loan report", report_month or "All dates"])
    if report_start and report_end:
        summary = monthly_report_summary(report_start, report_end, request.user)
        summary_sheet.append(["Applications submitted", summary["applications"]])
        summary_sheet.append(["Loans disbursed", summary["disbursed_loans"]])
        summary_sheet.append(["Disbursed value (UGX)", summary["disbursed_value"]])
        summary_sheet.append(["Net repayments (UGX)", summary["repayments"]])
    summary_sheet.append(["Generated", timezone.now().isoformat()])

    applications_sheet = workbook.create_sheet("Applications")
    application_headers = [
        "Reference",
        "Submitted",
        "School ID",
        "School",
        "District",
        "Portfolio owner",
        "Applicant",
        "Phone",
        "Purpose",
        "Preferred lender",
        "Requested amount (UGX)",
        "Term months",
        "Frequency",
        "Status",
        "Next follow-up",
    ]
    applications_sheet.append(application_headers)
    for application in _application_export_rows(request.user, filters).iterator(
        chunk_size=500
    ):
        applications_sheet.append(
            _xlsx_row(
                [
                    str(application.id),
                    application.submitted_at.date(),
                    application.school.school_id,
                    application.school.name,
                    application.school.district.name
                    if application.school.district
                    else "",
                    application.school.account_owner_name_raw or "",
                    application.applicant_name,
                    application.applicant_phone,
                    application.purpose.label,
                    application.preferred_mfi.name if application.preferred_mfi else "",
                    application.requested_amount,
                    application.requested_term_months,
                    application.get_repayment_frequency_display(),
                    application.get_status_display(),
                    application.next_follow_up_on,
                ]
            )
        )

    loans_sheet = workbook.create_sheet("Processed loans")
    loan_headers = [
        "Lending partner",
        "School ID",
        "School",
        "Purpose",
        "Loan reference",
        "Status",
        "Approved amount (UGX)",
        "Disbursed amount (UGX)",
        "Disbursement date",
        "Term months",
        "Repayment frequency",
        "Last repayment update",
    ]
    loans_sheet.append(loan_headers)
    for loan in services.loan_export_rows(request.user, filters).iterator(
        chunk_size=500
    ):
        loans_sheet.append(
            _xlsx_row(
                [
                    loan.mfi.name,
                    loan.school.school_id,
                    loan.school.name,
                    loan.purpose.label,
                    loan.external_loan_reference,
                    loan.get_status_display(),
                    loan.approved_amount,
                    loan.disbursed_amount,
                    loan.disbursement_date,
                    loan.term_months,
                    loan.repayment_frequency,
                    loan.last_repayment_data_date,
                ]
            )
        )

    repayments_sheet = workbook.create_sheet("Repayments")
    repayments_sheet.append(
        [
            "Value date",
            "Received on",
            "School ID",
            "School",
            "Lending partner",
            "Loan reference",
            "Transaction reference",
            "Type",
            "Signed amount (UGX)",
            "Evidence reference",
            "Reason",
        ]
    )
    repayments = RepaymentTransaction.objects.filter(
        loan__in=services.scoped_loans(request.user)
    ).select_related("loan", "loan__school", "loan__mfi")
    if report_start and report_end:
        repayments = repayments.filter(
            value_date__gte=report_start, value_date__lt=report_end
        )
    for repayment in repayments.order_by("value_date", "created_at").iterator(
        chunk_size=500
    ):
        signed_amount = (
            -repayment.amount
            if repayment.kind == RepaymentTransactionKind.REVERSAL
            else repayment.amount
        )
        repayments_sheet.append(
            _xlsx_row(
                [
                    repayment.value_date,
                    repayment.received_on,
                    repayment.loan.school.school_id,
                    repayment.loan.school.name,
                    repayment.loan.mfi.name,
                    repayment.loan.external_loan_reference,
                    repayment.external_reference,
                    repayment.get_kind_display(),
                    signed_amount,
                    repayment.evidence_reference,
                    repayment.reason,
                ]
            )
        )

    header_fill = PatternFill("solid", fgColor="17365D")
    for sheet in (summary_sheet, applications_sheet, loans_sheet, repayments_sheet):
        if sheet is not summary_sheet:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for column in sheet.columns:
            width = min(
                36, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )
            sheet.column_dimensions[column[0].column_letter].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        'attachment; filename="uganda-monthly-loan-report.xlsx"'
    )
    workbook.save(response)
    return response


@require_POST
@require_page_permission("loans")
def loan_save_action(request):
    services.register_or_update_loan(request.POST.dict(), request.user)
    messages.success(request, "Loan record saved.")
    return redirect("/loans")


@require_POST
@require_page_permission("loans")
def salesforce_confirmation_action(request, loan_id: str):
    result = services.confirm_salesforce_loan(
        loan_id, request.POST.dict(), request.user
    )
    messages.success(
        request,
        f"Salesforce Loan ID {result['salesforceLoanId']} confirmed.",
    )
    return redirect("/loans")


@require_POST
@require_page_permission("loans")
def salesforce_return_action(request, loan_id: str):
    services.return_salesforce_loan(loan_id, request.POST.dict(), request.user)
    messages.success(request, "Loan returned to the MFI for correction.")
    return redirect("/loans")


@require_POST
@require_page_permission("loans")
def ia_validation_action(request, loan_id: str):
    result = services.validate_loan_by_ia(loan_id, request.POST.dict(), request.user)
    message = (
        "Loan verified as a programme record."
        if result["iaValidationStatus"] == "verified"
        else "Loan returned for correction."
    )
    messages.success(request, message)
    return redirect("/loans")


@require_POST
@require_page_permission("loans")
def repayment_snapshot_action(request, loan_id: str):
    services.add_repayment_snapshot(loan_id, request.POST.dict(), request.user)
    messages.success(request, "Repayment snapshot recorded.")
    return redirect("/loans")


@require_POST
@require_page_permission("loans")
def repayment_snapshot_save_action(request):
    services.add_repayment_snapshot(
        request.POST.get("loanId", ""), request.POST.dict(), request.user
    )
    messages.success(request, "Repayment snapshot recorded.")
    return redirect("/loans")


@require_GET
@require_page_permission("business_transformation")
def portfolio_partial(request):
    context = services.workspace_context(request.user, _filters(request))
    return render(request, "partials/business_transformation/portfolio.html", context)


@require_POST
@require_page_permission("business_transformation")
def triage_action(request, case_id: str):
    services.triage_case(case_id, request.POST.dict(), request.user)
    messages.success(request, "Business Transformation triage decision saved.")
    if request.headers.get("HX-Request") == "true":
        response = HttpResponse(status=204)
        response["HX-Redirect"] = "/business-transformation"
        return response
    return redirect("/business-transformation")
