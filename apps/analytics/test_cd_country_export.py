"""One period, five country CSVs (owner, 2026-09-03; impact added 2026-09-13).

The CD used to get one export: the PL roster. Risk, finance and core health
lived on screen only, so any figure carried into a board meeting was typed
by hand from a browser window.
"""

from __future__ import annotations

from django.contrib.auth import get_user_model
from django.test import TestCase

from apps.accounts.models import StaffProfile
from apps.analytics.cd_export_service import DATASETS, country_export, normalise_dataset
from apps.core.fy import get_operational_fy
from apps.core.rbac import EdifyRole

User = get_user_model()


def _person(email, role):
    u = User.objects.create_user(
        email=email,
        name=email,
        roles=[role],
        active_role=role,
        password="x",
        is_active=True,
    )
    StaffProfile.objects.create(user=u, title=role, country="Uganda")
    return u


class CountryExportSetTest(TestCase):
    def setUp(self):
        self.cd = _person("exp-cd@t.org", EdifyRole.COUNTRY_DIRECTOR.value)
        self.cceo = _person("exp-cceo@t.org", EdifyRole.CCEO.value)

    def test_every_dataset_still_downloads_as_csv_for_the_selected_period(self):
        self.client.force_login(self.cd)
        fy = get_operational_fy()
        for dataset in DATASETS:
            response = self.client.get(
                f"/analytics/country-director/export?set={dataset}&fy={fy}"
                "&quarter=Q1&format=csv"
            )
            self.assertEqual(response.status_code, 200, dataset)
            self.assertTrue(response["Content-Type"].startswith("text/csv"))
            self.assertIn(f"-{fy}.csv", response["Content-Disposition"])
            first_line = response.content.decode().splitlines()[0]
            self.assertTrue(first_line, dataset)

    def test_one_export_is_one_workbook_with_a_sheet_per_dataset(self):
        """Owner, 2026-09-27: "export buttons dont have to be two buttons. It
        should be export in excel not csv since the team use excel more"."""
        from io import BytesIO

        from openpyxl import load_workbook

        from apps.analytics.cd_export_service import dataset_label
        from apps.core.excel import XLSX_CONTENT_TYPE

        self.client.force_login(self.cd)
        fy = get_operational_fy()
        response = self.client.get(f"/analytics/country-director/export?fy={fy}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertIn(f"cd-country-{fy}.xlsx", response["Content-Disposition"])
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, [dataset_label(d) for d in DATASETS])
        one = self.client.get(f"/analytics/country-director/export?set=risk&fy={fy}")
        self.assertEqual(one["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertEqual(
            load_workbook(BytesIO(one.content)).sheetnames, [dataset_label("risk")]
        )

    def test_the_five_datasets_carry_distinct_headers(self):
        seen = set()
        for dataset in DATASETS:
            slug, header, rows = country_export(
                self.cd, dataset, fy=get_operational_fy()
            )
            self.assertIsInstance(rows, list)
            seen.add((slug, tuple(header)))
        self.assertEqual(len(seen), 5)

    def test_an_unknown_set_falls_back_to_delivery(self):
        self.assertEqual(normalise_dataset("wat"), "delivery")
        self.assertEqual(normalise_dataset(None), "delivery")
        self.assertEqual(normalise_dataset(" Risk "), "risk")

    def test_the_page_offers_one_export_and_a_cceo_gets_nothing(self):
        self.client.force_login(self.cd)
        page = self.client.get("/analytics/country-director")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'data-component="cd-export"')
        self.assertContains(page, "/analytics/country-director/export?fy=")
        self.assertNotContains(page, "/analytics/country-director/export?set=")
        self.client.force_login(self.cceo)
        response = self.client.get("/analytics/country-director/export?set=finance")
        self.assertNotEqual(response.status_code, 200)
        self.assertNotEqual(response.get("Content-Type"), "text/csv")
