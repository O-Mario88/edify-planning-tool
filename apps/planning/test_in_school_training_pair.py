"""One in-school Training decision creates and completes two governed records."""

from __future__ import annotations

import datetime
from unittest.mock import patch

from django.test import Client

from apps.accounts.models import StaffSchoolAssignment
from apps.activities.models import (
    Activity,
    ActivitySalesforceReference,
    ActivityScheduleCostLine,
    SchoolVisitFeedback,
)
from apps.activities.services import (
    _apply_schedule_cost_snapshot as apply_real_cost_snapshot,
    complete_in_school_training_pair,
    start_in_school_training_pair,
)
from apps.activity_catalogue.availability import in_school_training_course_options
from apps.core.enums import EvidenceKind, SsaIntervention
from apps.core.exceptions import BadRequest
from apps.evidence.models import EvidenceRecord
from apps.planning.services import schedule_in_school_training_pair
from apps.system_health.services import missing_cost_lines_count

from .test_standard_support_scheduling import (
    StandardSupportBase,
    _at,
    _schedulable_date,
)


class InSchoolTrainingPairTest(StandardSupportBase):
    def payload(self, course_code="TAM_I"):
        return {
            "schoolId": self.school.school_id,
            "catalogueItemId": self.item(course_code).id,
            "scheduledDate": _at(_schedulable_date()).isoformat(),
            "responsibleStaffId": self.staff.id,
            "deliveryType": "staff",
            "executorType": "staff",
            "requireCatalogue": True,
            "activityPurposeText": "Deliver the selected in-school course",
        }

    def test_in_school_picker_contains_all_21_courses_with_ssa_metadata(self):
        options = in_school_training_course_options(school=self.school)
        self.assertEqual(len(options), 21)
        by_code = {option["stableCode"]: option for option in options}
        # SSA Training is linked to Leadership (owner, 2026-10-06); it was
        # linked to no intervention until then.
        self.assertEqual(
            by_code["SSA_TRAINING"]["ssaIntervention"],
            SsaIntervention.LEADERSHIP,
        )
        self.assertEqual(
            by_code["TAM_I"]["ssaIntervention"],
            SsaIntervention.EXPOSURE_TO_WORD_OF_GOD,
        )
        self.assertEqual(by_code["TAM_I"]["category"], "Christian Transformation")
        self.assertEqual(by_code["NEW_SCHOOL_ORIENTATION"]["ssaIntervention"], "")

    def test_school_drawer_renders_cluster_and_school_course_titles(self):
        client = Client()
        client.force_login(self.user)
        response = client.get(
            f"/planning/schedule-modal?school_id={self.school.school_id}",
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertIn("Teaching as Mission (TAM)", html)
        self.assertIn("EdTech Foundations Training", html)
        self.assertIn("New School Orientation", html)

    def test_schedule_action_routes_in_school_training_to_pair_service(self):
        client = Client()
        client.force_login(self.user)
        with patch(
            "apps.frontend.views.planning_views.schedule_in_school_training_pair",
            return_value={
                "id": "training-id",
                "pairedSchoolVisitId": "visit-id",
            },
        ) as schedule_pair:
            response = client.post(
                "/planning/schedule-action",
                {
                    "school_id": self.school.school_id,
                    "purpose_of_visit": "in_school_training",
                    "scheduled_date": str(_schedulable_date()),
                    "catalogue_item_id": self.item("TAM_I").id,
                    "require_catalogue": "yes",
                    "delivery_type": "staff",
                    "executor_type": "staff",
                },
                HTTP_HX_REQUEST="true",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(schedule_pair.call_count, 1)
        payload = schedule_pair.call_args.args[0]
        self.assertEqual(payload["catalogueItemId"], self.item("TAM_I").id)
        self.assertEqual(
            payload["focusIntervention"],
            SsaIntervention.EXPOSURE_TO_WORD_OF_GOD,
        )

    def test_one_decision_creates_training_and_visit_with_course_identity(self):
        result = schedule_in_school_training_pair(self.payload(), self.user)

        training = Activity.objects.get(id=result["id"])
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])
        self.assertEqual(Activity.objects.filter(school=self.school).count(), 2)
        self.assertEqual(training.activity_type, "in_school_training")
        self.assertEqual(training.training_course.stable_code, "TAM_I")
        self.assertEqual(
            training.catalogue_item.stable_code, "STANDARD_IN_SCHOOL_TRAINING"
        )
        self.assertEqual(training.activity_name_snapshot, "Teaching as Mission (TAM)")
        self.assertEqual(
            training.focus_intervention,
            SsaIntervention.EXPOSURE_TO_WORD_OF_GOD,
        )
        self.assertEqual(training.paired_school_visit_id, visit.id)
        self.assertEqual(visit.activity_type, "school_visit")
        self.assertEqual(visit.catalogue_item.stable_code, "STANDARD_SCHOOL_VISIT")
        self.assertEqual(visit.school_id, training.school_id)
        self.assertEqual(visit.planned_date, training.planned_date)
        self.assertEqual(visit.responsible_staff_id, training.responsible_staff_id)
        self.assertEqual(visit.focus_intervention, training.focus_intervention)

    def test_the_training_is_facilitated_and_its_visit_is_not(self):
        """Owner, 2026-09-29: "Facilitated by" on scheduling training. The
        school drawer offers it with the course; the named partner lands on
        the Training only, since the visit is the officer's journey."""
        from apps.partners.models import Partner

        partner = Partner.objects.create(
            name="Pair Facilitator Org", active_status=True
        )
        client = Client()
        client.force_login(self.user)
        html = client.get(
            f"/planning/schedule-modal?school_id={self.school.school_id}",
            HTTP_HX_REQUEST="true",
        ).content.decode("utf-8")
        self.assertIn('name="facilitating_partner_id"', html)
        self.assertIn("Pair Facilitator Org", html)

        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(
            {**self.payload(), "facilitatingPartnerId": partner.id}, self.user
        )
        training = Activity.objects.get(id=result["id"])
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])
        self.assertEqual(training.facilitating_partner_id, partner.id)
        self.assertEqual(training.delivery_type, "staff")
        self.assertIsNone(visit.facilitating_partner_id)
        # The facilitation fee comes from the rate card, on the Training
        # only, and it is the partner's.
        fee = ActivityScheduleCostLine.objects.get(
            activity=training, line_item_type="facilitation"
        )
        self.assertGreater(fee.amount, 0)
        self.assertEqual(fee.partner_id, partner.id)
        self.assertEqual(
            ActivityScheduleCostLine.objects.filter(activity=training).count(), 1
        )
        self.assertFalse(
            ActivityScheduleCostLine.objects.filter(
                activity=visit, line_item_type="facilitation"
            ).exists()
        )
        # The fee is not visit cost stranded on the Training.
        from apps.activities.pair_costing import find_pair_trainings_carrying_cost

        self.assertNotIn(training.id, find_pair_trainings_carrying_cost())

    def test_administrative_course_does_not_invent_an_ssa_intervention(self):
        result = schedule_in_school_training_pair(
            self.payload("NEW_SCHOOL_ORIENTATION"), self.user
        )
        training = Activity.objects.get(id=result["id"])
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])
        self.assertIsNone(training.focus_intervention)
        self.assertIsNone(visit.focus_intervention)
        # Not school-improvement work, and the companion visit says the same
        # rather than being handed the school's weakest intervention.
        self.assertEqual(training.ssa_alignment, "not_applicable")
        self.assertEqual(visit.ssa_alignment, "not_applicable")

    def test_pair_carries_one_visit_cost_on_the_visit(self):
        """Owner, 2026-09-28: the in-school training "use[s] the visit cost"
        and its own cost is 0, "since it is part of school visit"."""
        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])

        visit_lines = list(ActivityScheduleCostLine.objects.filter(activity=visit))
        self.assertTrue(visit_lines)
        self.assertEqual(
            {line.cost_setting_key for line in visit_lines},
            {"primary_transport_per_day", "lunch_per_day"},
        )
        self.assertGreater(visit.est_cost_cents, 0)
        self.assertIsNotNone(visit.daily_visit_batch_id)
        self.assertFalse(
            ActivityScheduleCostLine.objects.filter(activity=training).exists()
        )
        self.assertEqual(training.est_cost_cents, 0)
        self.assertFalse(training.cost_missing)
        self.assertIsNone(training.daily_visit_batch_id)
        self.assertEqual(visit.planned_date, training.planned_date)
        self.assertEqual(missing_cost_lines_count(), 0)

    def _plain_visit(self, school, day):
        from apps.activities.services import create

        result = create(
            {
                "schoolId": school.school_id,
                "catalogueItemId": self.item("STANDARD_SCHOOL_VISIT").id,
                "focusIntervention": SsaIntervention.LEADERSHIP,
                "activityPurposeText": "Coach the head teacher",
                "scheduledDate": _at(day).isoformat(),
                "responsibleStaffId": self.staff.id,
                "requireCatalogue": True,
            },
            self.user,
        )
        return Activity.objects.get(id=result["id"])

    def _cost(self, activity):
        return sorted(
            ActivityScheduleCostLine.objects.filter(activity=activity).values_list(
                "cost_setting_key", "amount"
            )
        )

    def _assert_pair_costs_what_a_visit_costs(self, day_keys):
        """An in-school training and a plain School Visit by the same
        officer, on separate trips, at neighbouring schools of one district."""
        from apps.schools.models import School

        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        neighbour = School.objects.get(school_id="STD-MEM-0")
        StaffSchoolAssignment.objects.create(staff=self.staff, school_id=neighbour.id)
        # Two days apart: on consecutive days in a secondary district the
        # second is the day the officer comes home, and costs less.
        first = _schedulable_date(room=4)
        second = first + datetime.timedelta(days=2)
        while second.weekday() == 6:
            second += datetime.timedelta(days=1)

        result = schedule_in_school_training_pair(
            {**self.payload(), "scheduledDate": _at(first).isoformat()}, self.user
        )
        training = Activity.objects.get(id=result["id"])
        companion = Activity.objects.get(id=result["pairedSchoolVisitId"])
        plain = self._plain_visit(neighbour, second)

        self.assertEqual({key for key, _amount in self._cost(plain)}, day_keys)
        self.assertEqual(self._cost(companion), self._cost(plain))
        self.assertEqual(companion.est_cost_cents, plain.est_cost_cents)
        self.assertGreater(companion.est_cost_cents, 0)
        # The Training adds nothing on top of the one visit day.
        self.assertEqual(self._cost(training), [])
        self.assertEqual(training.est_cost_cents, 0)

    def test_in_a_primary_district_the_pair_costs_a_primary_district_visit(self):
        """Owner, 2026-10-05: "make sure in-school training uses the same cost
        as school visits. If it is from primary district, it should fetch cost
        for primary district visit"."""
        self._assert_pair_costs_what_a_visit_costs(
            {"primary_transport_per_day", "lunch_per_day"}
        )

    def test_in_a_secondary_district_the_pair_costs_a_secondary_district_visit(self):
        """... "and if it from secondary district, it should fetch cost from
        secondary District." The officer's own district is another one, so
        this school is a night away."""
        from apps.geography.models import District

        home = District.objects.create(
            name="Standard Home District", region=self.region, district_type="primary"
        )
        self.staff.primary_district_id = home.id
        self.staff.save(update_fields=["primary_district_id"])
        self._assert_pair_costs_what_a_visit_costs(
            {
                "secondary_transport_per_day",
                "lunch_per_day",
                "secondary_breakfast_per_day",
                "secondary_overnight_dinner_per_day",
                "secondary_accommodation_per_night",
            }
        )

    def _work_plan(self, day, **params):
        from apps.core.fy import get_operational_fy
        from apps.frontend.views.work_plan_page import build_work_plan_context

        return build_work_plan_context(
            self.user,
            {
                "fy": get_operational_fy(day),
                "view": "month",
                "period": str(day.month),
                **params,
            },
        )

    def test_the_work_plan_keeps_the_training_at_zero_and_says_where_its_cost_is(
        self,
    ):
        """Owner, 2026-10-05: "On the work plan the cost is missing", then
        "since in-school training is done during school visit, it is ok to
        keep it at UGX 0 just add (captured in the school visit cost) to UGX
        0". The Training has no cost lines of its own and the page read that
        as a cost nobody had set up."""
        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])

        context = self._work_plan(training.planned_date)
        rows = {row["id"]: row for row in context["rows"]}
        training_row, visit_row = rows[training.id], rows[visit.id]

        self.assertFalse(training_row["cost_missing"])
        self.assertEqual(training_row["cost"], 0)
        self.assertEqual(
            training_row["cost_note"], "(captured in the school visit cost)"
        )
        # The day's money is the visit's, counted once.
        self.assertFalse(visit_row["cost_missing"])
        self.assertEqual(visit_row["cost"], visit.est_cost_cents)
        self.assertEqual(visit_row["cost_note"], "")
        self.assertEqual(context["plan_summary"]["total"]["cost"], visit.est_cost_cents)
        self.assertEqual(context["plan_summary"]["total"]["cost_missing_count"], 0)

        summary = {
            row["label"]: row
            for section in context["plan_summary_sections"]
            for row in section["rows"]
        }
        pair_row = summary["Staff In-school Training (with School Visit)"]
        self.assertEqual(pair_row["count"], 1)
        self.assertEqual(pair_row["cost"], 0)
        self.assertEqual(pair_row["unit_cost_display"], "0")
        self.assertEqual(
            pair_row["unit_cost_note"], "(captured in the school visit cost)"
        )
        self.assertEqual(summary["Staff School Visit"]["cost"], visit.est_cost_cents)
        self.assertEqual(summary["Staff School Visit"]["unit_cost_note"], "")

        # Nothing is left for the "Cost setup required" drill-down to list.
        flagged = self._work_plan(training.planned_date, flag="cost_missing")
        self.assertEqual(flagged["rows"], [])

    def test_the_work_plan_page_and_export_say_where_the_cost_is(self):
        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        from apps.core.fy import get_operational_fy

        query = (
            f"?fy={get_operational_fy(training.planned_date)}&view=month"
            f"&period={training.planned_date.month}"
        )
        client = Client()
        client.force_login(self.user)
        response = client.get(f"/work-plan{query}")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertNotIn("Cost setup required", html)
        self.assertRegex(
            html,
            r"UGX 0 <span[^>]*data-cost-note>\(captured in the school visit cost\)",
        )

        from io import BytesIO

        from openpyxl import load_workbook

        # A CCEO reads the page; the export belongs to the roles above them.
        from apps.accounts.models import User

        director = User.objects.create_user(
            email="standard-cd@edify.org",
            name="Standard CD",
            roles=["CountryDirector"],
            active_role="CountryDirector",
            password="x",
            is_active=True,
        )
        client.force_login(director)
        export = client.get(f"/work-plan/export.xlsx{query}")
        self.assertEqual(export.status_code, 200)
        summary = load_workbook(BytesIO(export.content))["Plan Summary"]
        cells = {row[0]: row for row in summary.iter_rows(values_only=True)}
        self.assertEqual(
            cells["Staff In-school Training (with School Visit)"][2],
            "0 (captured in the school visit cost)",
        )
        self.assertEqual(cells["Staff In-school Training (with School Visit)"][3], 0)
        self.assertNotIn("Cost setup required", cells)

    def test_my_plan_keeps_the_training_at_zero_and_says_where_its_cost_is(self):
        from apps.core.fy import get_operational_fy
        from apps.my_plan.services import get_frontend_context

        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        fy = get_operational_fy(training.planned_date)

        context = get_frontend_context(self.user, {"fy": fy, "period": "fy"})
        row = next(r for r in context["cluster_trainings"] if r["id"] == training.id)
        # Never priced, so it holds no estimate: the page writes UGX 0.
        self.assertFalse(row["budget_total"])
        self.assertEqual(row["cost_note"], "(captured in the school visit cost)")
        self.assertEqual(row["budget_status"], "In School Visit")

        client = Client()
        client.force_login(self.user)
        html = client.get(f"/my-plan?fy={fy}&period=fy").content.decode("utf-8")
        self.assertRegex(
            html,
            r"UGX 0 <span[^>]*data-cost-note>\(captured in the school visit cost\)",
        )

        # A pair whose visit has no price is not papered over.
        ActivityScheduleCostLine.objects.filter(
            activity_id=result["pairedSchoolVisitId"]
        ).delete()
        context = get_frontend_context(self.user, {"fy": fy, "period": "fy"})
        row = next(r for r in context["cluster_trainings"] if r["id"] == training.id)
        self.assertEqual(row["cost_note"], "")
        self.assertEqual(row["budget_status"], "No Budget")

    def test_the_work_plan_still_flags_a_pair_whose_visit_has_no_price(self):
        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        ActivityScheduleCostLine.objects.filter(
            activity_id=result["pairedSchoolVisitId"]
        ).delete()

        context = self._work_plan(training.planned_date)
        rows = {row["id"]: row for row in context["rows"]}
        self.assertTrue(rows[training.id]["cost_missing"])
        self.assertTrue(rows[result["pairedSchoolVisitId"]]["cost_missing"])

    def test_planning_oversight_does_not_call_the_training_uncosted(self):
        from apps.core.fy import get_operational_fy
        from apps.planning.oversight_service import build_items

        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        items = {
            item.activity_id: item
            for item in build_items(
                self.user, fy=get_operational_fy(training.planned_date)
            )
        }
        for activity_id in (training.id, result["pairedSchoolVisitId"]):
            with self.subTest(activity=activity_id):
                self.assertNotIn(
                    "scheduled_without_cost",
                    {risk["key"] for risk in items[activity_id].risks},
                )
        self.assertTrue(items[training.id].cost_on_school_visit)
        self.assertFalse(items[result["pairedSchoolVisitId"]].cost_on_school_visit)

    def test_training_needs_no_participant_count(self):
        """The in-school drawer asks for no participants. The Training used
        to be priced, and a Training priced as a group session was refused
        with "Enter the planned participant count"; unpriced, it cannot be."""
        from apps.budget import costing_service

        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        real_preview = costing_service.preview
        gated = []

        def record(input, **kwargs):
            gated.append(input["activityType"])
            return real_preview(input, **kwargs)

        with patch.object(costing_service, "preview", side_effect=record):
            result = schedule_in_school_training_pair(self.payload(), self.user)
        self.assertEqual(gated, ["school_visit"])
        self.assertTrue(result["pairedSchoolVisitId"])

    def test_repricing_the_training_keeps_it_at_zero(self):
        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        visit_total = Activity.objects.get(
            id=result["pairedSchoolVisitId"]
        ).est_cost_cents

        apply_real_cost_snapshot(training, {})

        training.refresh_from_db()
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])
        self.assertEqual(training.est_cost_cents, 0)
        self.assertFalse(
            ActivityScheduleCostLine.objects.filter(activity=training).exists()
        )
        self.assertEqual(visit.est_cost_cents, visit_total)

    def _pair_in_the_old_shape(self):
        """A pair as it was scheduled before 2026-09-28: the visit unpriced,
        the Training priced. Returns (training, visit, the visit's price)."""
        from apps.budget.costing import ActivityCost
        from apps.budget.costing_service import apply_to_activity
        from apps.daily_visit_batches.services import remove_school

        self.cost_snapshot.side_effect = apply_real_cost_snapshot
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])
        visit_total = visit.est_cost_cents
        remove_school(activity_id=visit.id)
        visit.refresh_from_db()
        apply_to_activity(
            visit, {"activityType": "school_visit"}, precomputed_cost=ActivityCost()
        )
        apply_to_activity(
            training,
            {
                "activityType": "in_school_training",
                "deliveryType": "staff",
                "districtType": "primary",
            },
        )
        return training, visit, visit_total

    def test_an_older_pair_reads_the_other_way_round_and_neither_half_is_flagged(
        self,
    ):
        """A pair scheduled before 2026-09-28 that has not been moved carries
        the cost on its Training. The live Work Plan listed every such visit
        as "Cost setup required" and My Plan as "No Budget" (2026-10-05)."""
        from apps.core.fy import get_operational_fy
        from apps.my_plan.services import get_frontend_context

        training, visit, visit_total = self._pair_in_the_old_shape()
        training.refresh_from_db()
        self.assertEqual(training.est_cost_cents, visit_total)

        context = self._work_plan(training.planned_date)
        rows = {row["id"]: row for row in context["rows"]}
        self.assertFalse(rows[visit.id]["cost_missing"])
        self.assertEqual(rows[visit.id]["cost"], 0)
        self.assertEqual(
            rows[visit.id]["cost_note"], "(captured in the in-school training cost)"
        )
        self.assertFalse(rows[training.id]["cost_missing"])
        self.assertEqual(rows[training.id]["cost"], visit_total)
        self.assertEqual(rows[training.id]["cost_note"], "")
        self.assertEqual(context["plan_summary"]["total"]["cost_missing_count"], 0)
        self.assertEqual(context["plan_summary"]["total"]["cost"], visit_total)
        summary = {
            row["label"]: row
            for section in context["plan_summary_sections"]
            for row in section["rows"]
        }
        self.assertEqual(
            summary["Staff School Visit (with In-school Training)"]["unit_cost_note"],
            "(captured in the in-school training cost)",
        )
        # The Training carries the day itself, like one with no visit.
        self.assertEqual(summary["Staff In-school Training"]["cost"], visit_total)
        self.assertEqual(summary["Staff In-school Training"]["unit_cost_note"], "")

        fy = get_operational_fy(training.planned_date)
        my_plan = get_frontend_context(self.user, {"fy": fy, "period": "fy"})
        visit_row = next(r for r in my_plan["school_visits"] if r["id"] == visit.id)
        self.assertEqual(
            visit_row["cost_note"], "(captured in the in-school training cost)"
        )
        self.assertEqual(visit_row["budget_status"], "In Training")

    def test_a_pair_scheduled_before_the_change_moves_its_cost_to_the_visit(self):
        """Pairs made before 2026-09-28 carry the cost on the Training."""
        from apps.activities.pair_costing import (
            find_pair_trainings_carrying_cost,
            move_pair_costs_to_visits,
        )

        training, visit, visit_total = self._pair_in_the_old_shape()
        self.assertEqual(find_pair_trainings_carrying_cost(), [training.id])

        moved = move_pair_costs_to_visits(write=lambda _line: None)

        self.assertEqual(moved, {"moved": [training.id], "skipped": [], "left": []})
        training.refresh_from_db()
        visit.refresh_from_db()
        self.assertEqual(training.est_cost_cents, 0)
        self.assertEqual(visit.est_cost_cents, visit_total)
        self.assertEqual(find_pair_trainings_carrying_cost(), [])
        self.assertEqual(missing_cost_lines_count(), 0)

    def test_a_pair_not_reached_by_the_deadline_is_left_for_the_next_run(self):
        """Migration 0061 runs in App Platform's 30-minute pre-deploy job and
        stops starting pairs at its budget. What it does not reach keeps its
        cost, so the command (or the next deployment) finds it again."""
        import importlib
        import time

        from apps.activities.pair_costing import (
            find_pair_trainings_carrying_cost,
            move_pair_costs_to_visits,
        )

        training, visit, visit_total = self._pair_in_the_old_shape()

        late = move_pair_costs_to_visits(
            write=lambda _line: None, deadline=time.monotonic() - 1
        )

        self.assertEqual(late, {"moved": [], "skipped": [], "left": [training.id]})
        self.assertEqual(find_pair_trainings_carrying_cost(), [training.id])
        self.assertEqual(
            move_pair_costs_to_visits(write=lambda _line: None)["moved"],
            [training.id],
        )
        visit.refresh_from_db()
        self.assertEqual(visit.est_cost_cents, visit_total)

        migration = importlib.import_module(
            "apps.activities.migrations.0061_move_in_school_training_cost_to_visit"
        )
        self.assertFalse(
            migration.Migration.atomic,
            "one transaction for every pair is what outlived the pre-deploy job",
        )
        self.assertLess(migration.BUDGET_SECONDS, 30 * 60)

    def test_a_failure_creating_visit_rolls_back_training(self):
        before = Activity.objects.count()
        real_create = __import__("apps.activities.services", fromlist=["create"]).create
        calls = 0

        def fail_second(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise BadRequest("Visit could not be created.")
            return real_create(*args, **kwargs)

        with patch("apps.activities.services.create", side_effect=fail_second):
            with self.assertRaisesMessage(BadRequest, "Visit could not be created"):
                schedule_in_school_training_pair(self.payload(), self.user)
        self.assertEqual(Activity.objects.count(), before)

    def _scheduled_pair_ready_for_completion(self):
        result = schedule_in_school_training_pair(self.payload(), self.user)
        training = Activity.objects.get(id=result["id"])
        visit = Activity.objects.get(id=result["pairedSchoolVisitId"])
        for activity, kind in (
            (training, EvidenceKind.ATTENDANCE_FORM),
            (visit, EvidenceKind.VISIT_FORM),
        ):
            EvidenceRecord.objects.create(
                activity=activity,
                kind=kind,
                uri=f"{activity.id}.pdf",
                original_name=f"{kind}.pdf",
                mime_type="application/pdf",
                uploaded_by=self.user.id,
                uploader_role=self.user.active_role,
                quarantined=False,
            )
        start_in_school_training_pair(training.id, self.user)
        return training, visit

    def test_both_salesforce_ids_and_statuses_are_committed_together(self):
        training, visit = self._scheduled_pair_ready_for_completion()
        complete_in_school_training_pair(
            training.id,
            {
                "trainingSalesforceId": "TS-PAIR-1001",
                "visitSalesforceId": "SVE-PAIR-1001",
                "teachersAttended": 4,
                "leadersAttended": 1,
                "feedbackFinding": "School leaders applied the training content.",
                "schoolImprovements": [
                    "Lesson observations now happen weekly",
                    "Teacher attendance has improved",
                ],
            },
            self.user,
        )
        training.refresh_from_db()
        visit.refresh_from_db()
        self.assertEqual(training.salesforce_activity_id, "TS-PAIR-1001")
        self.assertEqual(visit.salesforce_activity_id, "SVE-PAIR-1001")
        self.assertEqual(training.status, "submitted_to_pl")
        self.assertEqual(visit.status, "submitted_to_pl")
        self.assertEqual(
            ActivitySalesforceReference.objects.filter(
                activity_id__in=[training.id, visit.id]
            ).count(),
            2,
        )
        feedback = SchoolVisitFeedback.objects.get(activity=visit)
        self.assertEqual(
            feedback.finding, "School leaders applied the training content."
        )
        self.assertEqual(len(feedback.improvements), 2)

    def test_invalid_visit_id_rolls_back_training_salesforce_reservation(self):
        training, visit = self._scheduled_pair_ready_for_completion()
        with self.assertRaisesMessage(BadRequest, "SVE-"):
            complete_in_school_training_pair(
                training.id,
                {
                    "trainingSalesforceId": "TS-PAIR-ROLLBACK",
                    "visitSalesforceId": "SV-WRONG-PREFIX",
                    "teachersAttended": 2,
                    "leadersAttended": 0,
                },
                self.user,
            )
        training.refresh_from_db()
        visit.refresh_from_db()
        self.assertEqual(training.status, "completion_started")
        self.assertEqual(visit.status, "completion_started")
        self.assertIsNone(training.salesforce_activity_id)
        self.assertIsNone(visit.salesforce_activity_id)
        self.assertFalse(
            ActivitySalesforceReference.objects.filter(
                activity_id__in=[training.id, visit.id]
            ).exists()
        )

    def test_completion_drawer_exposes_both_ids_and_both_evidence_inputs(self):
        result = schedule_in_school_training_pair(self.payload(), self.user)
        client = Client()
        client.force_login(self.user)
        response = client.get(
            f"/my-plan/{result['id']}/complete-drawer",
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertIn('name="salesforce_id"', html)
        self.assertIn('name="visit_salesforce_id"', html)
        self.assertIn('name="training_evidence_file"', html)
        self.assertIn('name="visit_evidence_file"', html)
        self.assertIn('name="feedback_finding"', html)
        self.assertIn('name="school_improvements"', html)
        self.assertIn("Submit training + visit", html)
