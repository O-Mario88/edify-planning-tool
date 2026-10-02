"""The planning rulebook (owner, 2026-10-01) as tests.

Who plans (a role held, not the role in use), which visits count and for
whom, what each school type needs, when Partner work is planned, what a
duplicate is — and the figures check the Country Director reads them in.
"""

from __future__ import annotations

import io

from django.test import Client

from apps.activities.models import Activity
from apps.core.enums import ActivityType, ExecutorType
from apps.core.rbac import EdifyRole
from apps.planning.country_oversight import figures_check, people, rules
from apps.planning.test_country_planning_oversight import FY, World, day

PURPOSES = (
    None,
    "ssa_support",
    "in_school_training",
    "in_school_coaching",
    "training_follow_up",
    "donor_visit",
    "story_gathering",
    "school_invitation",
    "social_visit",
    rules.COMPANION_PURPOSE,
)


class RulebookDefinitionsTest(World):
    def test_rulebook_names_only_real_types(self):
        rules.check()

    def test_database_and_python_agree_on_which_visits_count(self):
        school = self.school("client", self.cceo)
        Activity.objects.bulk_create(
            Activity(
                activity_type=activity_type,
                purpose_type=purpose,
                school=school,
                fy=FY,
                quarter="Q1",
                planned_date=day(30),
                status="scheduled",
            )
            for activity_type in ActivityType.values
            for purpose in PURPOSES
        )
        rows = Activity.objects.filter(school=school)
        counted = dict(
            rows.filter(rules.counted_visit_q())
            .annotate(kind=rules.visit_kind_case())
            .values_list("id", "kind")
        )
        outreach = set(
            rows.filter(rules.outreach_visit_q()).values_list("id", flat=True)
        )
        for pk, activity_type, purpose in rows.values_list(
            "id", "activity_type", "purpose_type"
        ):
            with self.subTest(activity_type=activity_type, purpose=purpose):
                self.assertEqual(
                    counted.get(pk), rules.visit_kind(activity_type, purpose)
                )
                # A visit is counted or it is outreach, never both.
                self.assertFalse(pk in counted and pk in outreach)

    def test_the_three_kinds_and_nothing_else(self):
        self.assertEqual(
            rules.visit_kind("school_visit_ssa_collection", "ssa_support"),
            rules.KIND_SSA,
        )
        self.assertEqual(
            rules.visit_kind("in_school_training", "in_school_training"),
            rules.KIND_IN_SCHOOL,
        )
        self.assertEqual(
            rules.visit_kind("training_follow_up_visit", "training_follow_up"),
            rules.KIND_FOLLOW_UP,
        )
        # A Core visit is what its purpose says it is.
        self.assertEqual(rules.visit_kind("core_visit", "ssa_support"), rules.KIND_SSA)
        for activity_type, purpose in (
            ("donor_visit", "donor_visit"),
            ("story_gathering_visit", "story_gathering"),
            ("school_invitation", "school_invitation"),
            ("social_visit", "social_visit"),
            ("core_visit", "donor_visit"),
            ("school_visit", rules.COMPANION_PURPOSE),
            ("cluster_meeting", None),
            ("training", None),
        ):
            with self.subTest(activity_type=activity_type):
                self.assertIsNone(rules.visit_kind(activity_type, purpose))

    def test_what_each_school_type_needs(self):
        core = rules.requirement_for("core")
        self.assertEqual((core.staff_visits, core.partner_visits), (2, 2))
        self.assertEqual((core.staff_trainings, core.partner_trainings), (2, 2))
        for school_type in ("client", "core_trained"):
            need = rules.requirement_for(school_type)
            self.assertEqual((need.visits, need.trainings), (1, 1))
        graduate = rules.requirement_for("core_graduate")
        self.assertEqual((graduate.visits, graduate.trainings), (1, 0))
        champion = rules.requirement_for("champion")
        self.assertEqual((champion.visits, champion.trainings), (0, 0))
        # 2 visits per Core school and 1 per client-rule school.
        self.assertEqual(
            rules.portfolio_reach({"core": 80, "client": 100, "core_trained": 20}),
            280,
        )

    def test_a_date_is_the_partners_only_when_the_partner_set_it(self):
        self.assertEqual(rules.date_author(self.partner_admin()), "partner")
        # Staff putting the first date on Partner work make it a staff date.
        self.assertEqual(rules.date_author(self.cceo_user), "staff")
        # Staff moving a date the Partner chose leave it the Partner's plan,
        # and an older row nobody recorded stays unrecorded.
        self.assertEqual(rules.date_author(self.cceo_user, "partner"), "partner")
        self.assertEqual(rules.date_author(self.cceo_user, "", already_dated=True), "")

    def partner_admin(self):
        from apps.accounts.models import User

        return User.objects.create(
            email="rule-partner@cpo.test",
            name="Partner Admin",
            roles=[EdifyRole.PARTNER_ADMIN.value],
            active_role=EdifyRole.PARTNER_ADMIN.value,
            is_active=True,
        )


class RosterTest(World):
    def test_every_lead_with_their_cceos_and_targets(self):
        teams = {team.name: team for team in rules.roster("Uganda")}
        self.assertEqual(
            [p.name for p in teams["Lead A"].people],
            ["Lead A", "Officer One", "Officer Two"],
        )
        self.assertEqual(teams["Lead A"].target, 280 + 560 + 560)
        self.assertEqual(teams["Lead B"].target, 280 + 560)

    def test_a_lead_using_another_role_today_is_still_a_lead(self):
        self.pl_user.roles = [
            EdifyRole.COUNTRY_PROGRAM_LEAD.value,
            EdifyRole.ADMIN.value,
        ]
        self.pl_user.active_role = EdifyRole.ADMIN.value
        self.pl_user.save()
        teams = {team.name: team for team in rules.roster("Uganda")}
        lead = teams["Lead A"].people[0]
        self.assertTrue(lead.is_lead)
        self.assertEqual(lead.role_in_use, EdifyRole.ADMIN.value)
        # Their CCEOs stay under them rather than falling to "No Programme Lead".
        self.assertEqual(len(teams["Lead A"].people), 3)
        self.assertNotIn(rules.NO_LEAD_LABEL, teams)

    def test_a_cceo_with_no_lead_is_still_followed(self):
        _user, loose = self.person("cpo-loose", "Officer Loose", EdifyRole.CCEO)
        teams = {team.key: team for team in rules.roster("Uganda")}
        self.assertEqual(
            [p.key for p in teams[rules.NO_LEAD_KEY].people], [str(loose.id)]
        )

    def test_somebody_who_has_left_is_not_on_the_roster(self):
        self.cceo2_user.is_active = False
        self.cceo2_user.save()
        names = [p.name for team in rules.roster("Uganda") for p in team.people]
        self.assertNotIn("Officer Two", names)


class PeoplePlanTest(World):
    def plan(self):
        return people.people_plan(self.cd_user, FY)

    def row(self, plan, profile):
        return next(p for p in plan.people if p.key == str(profile.id))

    def test_expected_is_headcount_times_the_role_target(self):
        plan = self.plan()
        self.assertEqual(plan.headcount(rules.CCEO_ROLE), 3)
        self.assertEqual(plan.headcount(rules.PROGRAM_LEAD_ROLE), 2)
        self.assertEqual(plan.total("target"), 3 * 560 + 2 * 280)

    def test_a_visit_counts_for_whoever_planned_it(self):
        school = self.school("client", self.cceo)
        other = self.school("client", self.cceo)
        self.activity(school, "training_follow_up_visit", owner=self.cceo2)
        self.activity(other, "school_visit_ssa_collection", owner=self.pl)
        plan = self.plan()
        holder = self.row(plan, self.cceo)
        self.assertEqual(holder.visits_planned, 0)
        self.assertEqual(holder.schools, {"client": 2})
        self.assertEqual(self.row(plan, self.cceo2).visits_planned, 1)
        lead = self.row(plan, self.pl)
        self.assertEqual(lead.visits[rules.KIND_SSA], 1)
        self.assertEqual(lead.target, 280)

    def test_an_id_from_either_id_space_is_the_same_person(self):
        school = self.school("client", self.cceo)
        visit = self.activity(school, "school_visit")
        # Some paths write the User id where others write the StaffProfile id.
        Activity.objects.filter(pk=visit.pk).update(
            responsible_staff_id=self.cceo_user.id
        )
        self.assertEqual(self.row(self.plan(), self.cceo).visits_planned, 1)

    def test_outreach_is_shown_and_not_counted(self):
        school = self.school("client", self.cceo)
        self.activity(school, "donor_visit", owner=self.cceo)
        self.activity(school, "story_gathering_visit", owner=self.cceo)
        self.activity(
            school, "school_visit", owner=self.cceo, purpose_type="social_visit"
        )
        officer = self.row(self.plan(), self.cceo)
        self.assertEqual(officer.visits_planned, 0)
        self.assertEqual(officer.outreach, 3)

    def test_an_in_school_training_is_one_visit_and_one_training(self):
        school = self.school("client", self.cceo)
        self.activity(
            school,
            "in_school_training",
            owner=self.cceo,
            purpose_type="in_school_training",
        )
        self.activity(
            school,
            "school_visit",
            owner=self.cceo,
            purpose_type=rules.COMPANION_PURPOSE,
        )
        officer = self.row(self.plan(), self.cceo)
        self.assertEqual(officer.visits, {**officer.visits, rules.KIND_IN_SCHOOL: 1})
        self.assertEqual(officer.visits_planned, 1)
        self.assertEqual(officer.trainings, 1)

    def test_undated_cancelled_and_partner_work_is_not_a_staff_plan(self):
        school = self.school("client", self.cceo)
        self.activity(
            school, "school_visit", owner=self.cceo, on=None, status="planned"
        )
        self.activity(school, "school_visit", owner=self.cceo, status="cancelled")
        self.activity(school, "school_visit", owner=self.cceo, partner=self.partner)
        self.assertEqual(self.row(self.plan(), self.cceo).visits_planned, 0)

    def test_a_person_short_of_schools_keeps_the_whole_target(self):
        for _ in range(3):
            self.school("core", self.cceo)
        self.school("client", self.cceo)
        officer = self.row(self.plan(), self.cceo)
        self.assertEqual(officer.reach, 3 * 2 + 1)
        self.assertEqual(officer.target, 560)
        self.assertEqual(officer.shortfall, 560 - 7)

    def test_someone_else_who_planned_a_visit_shows_without_a_target(self):
        school = self.school("client", self.cceo)
        self.activity(school, "school_visit", owner=self.admin)
        plan = self.plan()
        other = self.row(plan, self.admin)
        self.assertEqual((other.role, other.target, other.visits_planned), ("", 0, 1))
        self.assertEqual(other.team_key, people.OTHER_TEAM_KEY)
        self.assertEqual(plan.total("target"), 3 * 560 + 2 * 280)

    def test_partner_planning_is_nothing_until_a_partner_sets_a_date(self):
        waiting = self.school("client", self.cceo)
        moved = self.school("client", self.cceo)
        agency = self.school("client", self.cceo)
        self.handover(
            waiting,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        self.activity(
            moved,
            "school_visit",
            partner=self.partner,
            on=None,
            status="assigned_to_partner",
            monitored_by_staff_id=self.cceo.id,
        )
        # Staff booked a certified agency onto a day: assigned, not planned.
        self.activity(
            agency,
            "school_visit",
            partner=self.partner2,
            status="partner_scheduled",
            monitored_by_staff_id=self.cceo.id,
            executor_type=ExecutorType.CERTIFIED_PARTNER_AGENCY.value,
            partner_date_set_by="staff",
        )
        officer = self.row(self.plan(), self.cceo)
        self.assertEqual(officer.partner_schools, 3)
        self.assertEqual(officer.partner_assigned, 3)
        self.assertEqual(officer.partner_planned, 0)

        Activity.objects.filter(school=moved).update(
            status="partner_scheduled",
            planned_date=day(50),
            partner_date_set_by="partner",
        )
        officer = self.row(self.plan(), self.cceo)
        self.assertEqual((officer.partner_assigned, officer.partner_planned), (3, 1))

    def test_partner_work_dated_before_the_author_was_kept_is_the_partners(self):
        school = self.school("client", self.cceo)
        self.activity(
            school,
            "school_visit",
            partner=self.partner,
            status="partner_scheduled",
            monitored_by_staff_id=self.cceo.id,
        )
        self.assertEqual(self.row(self.plan(), self.cceo).partner_planned, 1)

    def test_a_dated_hand_over_is_counted_once_as_its_activity(self):
        school = self.school("client", self.cceo)
        visit = self.activity(
            school,
            "school_visit",
            partner=self.partner,
            status="partner_scheduled",
            partner_date_set_by="partner",
        )
        self.handover(
            school,
            self.partner,
            status="partner_scheduled",
            assigning_staff_id=self.cceo.id,
            scheduled_activity=visit,
        )
        officer = self.row(self.plan(), self.cceo)
        self.assertEqual((officer.partner_assigned, officer.partner_planned), (1, 1))
        self.assertEqual(officer.partner_schools, 1)

    def test_trainings_and_cluster_meetings_follow_their_planner(self):
        school = self.school("client", self.cceo, cluster=self.cluster)
        self.session("cluster_meeting", [school])
        self.session("cluster_training", [school])
        officer = self.row(self.plan(), self.cceo)
        self.assertEqual((officer.meetings, officer.cluster_trainings), (1, 1))


class SchoolYearTest(World):
    def year(self):
        return people.school_year(self.cd_user, FY)

    def types(self):
        return {row.school_type: row for row in people.summarise_by_type(self.year())}

    def test_types_are_read_apart_and_champions_need_nothing(self):
        self.school("core", self.cceo)
        self.school("client", self.cceo)
        self.school("core_trained", self.cceo)
        self.school("core_graduate", self.cceo)
        self.school("champion", self.cceo)
        types = self.types()
        self.assertEqual(
            {
                t: (row.schools, row.visit_slots, row.training_slots)
                for t, row in types.items()
            },
            {
                "core": (1, 4, 4),
                "client": (1, 1, 1),
                "core_trained": (1, 1, 1),
                "core_graduate": (1, 1, 0),
                "champion": (1, 0, 0),
            },
        )

    def test_a_slot_is_filled_once_however_many_plans_land_on_it(self):
        core = self.school("core", self.cceo)
        for offset in (10, 20, 30):
            self.activity(core, "core_visit", owner=self.cceo, on=offset)
        for offset in (40, 50, 60):
            self.activity(
                core,
                "core_visit",
                partner=self.partner,
                on=offset,
                status="partner_scheduled",
                partner_date_set_by="partner",
            )
        row = self.types()["core"]
        self.assertEqual((row.staff_slots, row.partner_slots), (2, 2))
        self.assertEqual(row.schools_with_visit, 1)

    def test_a_client_schools_one_slot_is_staffs_when_staff_planned_it(self):
        school = self.school("client", self.cceo)
        self.activity(school, "training_follow_up_visit", owner=self.cceo)
        self.activity(
            school,
            "school_visit",
            partner=self.partner,
            status="partner_scheduled",
            partner_date_set_by="partner",
        )
        row = self.types()["client"]
        self.assertEqual((row.staff_slots, row.partner_slots), (1, 0))
        self.assertEqual(row.with_partner, 1)
        self.assertEqual(row.duplicates, 1)

    def test_a_school_assigned_and_not_dated_is_held_and_not_planned(self):
        school = self.school("client", self.cceo)
        self.handover(school, self.partner, status="pending_scheduling")
        row = self.types()["client"]
        self.assertEqual(row.with_partner, 1)
        self.assertEqual(row.schools_with_visit, 0)
        self.assertEqual(row.partner_slots, 0)

    def test_core_trainings_split_two_staff_and_two_partner(self):
        core = self.school("core", self.cceo)
        for offset in (10, 20, 30):
            self.activity(core, "core_training", owner=self.cceo, on=offset)
        self.activity(
            core,
            "core_training",
            partner=self.partner,
            on=40,
            status="partner_scheduled",
            partner_date_set_by="partner",
        )
        row = self.types()["core"]
        self.assertEqual(row.trainings, 2 + 1)
        self.assertEqual(row.schools_with_training, 1)

    def test_a_cluster_session_trains_only_the_schools_on_its_roster(self):
        named = self.school("client", self.cceo, cluster=self.cluster)
        self.school("client", self.cceo, cluster=self.cluster)
        self.session("cluster_training", [named])
        self.session("cluster_meeting", [named])
        row = self.types()["client"]
        self.assertEqual((row.trainings, row.schools_with_training), (1, 1))
        self.assertEqual((row.clustered, row.in_meeting), (2, 1))

    def reasons(self, school):
        return self.year()[school.id].duplicate_reasons

    def test_support_and_ssa_support_together_are_not_a_duplicate(self):
        school = self.school("client", self.cceo)
        self.activity(school, "training_follow_up_visit", owner=self.cceo, on=10)
        self.activity(school, "school_visit_ssa_collection", owner=self.cceo, on=20)
        self.assertEqual(self.reasons(school), ())

    def test_the_same_kind_twice_is_a_duplicate(self):
        school = self.school("core_trained", self.cceo)
        self.activity(school, "training_follow_up_visit", owner=self.cceo, on=10)
        self.activity(school, "in_school_training", owner=self.cceo2, on=20)
        self.assertEqual(self.reasons(school), (rules.DUPLICATE_STAFF,))

    def test_staff_and_a_partner_at_one_client_school_is_a_duplicate(self):
        school = self.school("core_graduate", self.cceo)
        self.activity(school, "school_visit_ssa_collection", owner=self.cceo)
        self.handover(school, self.partner, status="pending_scheduling")
        self.assertEqual(self.reasons(school), (rules.DUPLICATE_BOTH,))
        self.assertEqual([s.id for s in people.duplicates(self.year())], [school.id])

    def test_a_core_school_is_meant_to_have_both(self):
        core = self.school("core", self.cceo)
        self.activity(core, "core_visit", owner=self.cceo)
        self.handover(core, self.partner, status="pending_scheduling")
        self.assertEqual(self.reasons(core), ())


class FiguresCheckTest(World):
    def workbook(self, user, query=""):
        import openpyxl

        client = Client()
        client.force_login(user)
        response = client.get(
            f"/country-planning-oversight/coverage-export?fy={FY}{query}"
        )
        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(response.content))

    def test_the_country_director_reads_today_beside_the_new_rules(self):
        school = self.school("client", self.cceo)
        self.activity(school, "training_follow_up_visit", owner=self.cceo2)
        self.handover(
            school,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        workbook = self.workbook(self.cd_user)
        titles = [name for name in workbook.sheetnames if name.startswith("Check")]
        self.assertEqual(len(titles), 4)
        cards = {row[0]: row for row in workbook[titles[0]].iter_rows(values_only=True)}
        staff = cards["Staff Visit Planning"]
        # Today: the holder's slot against the holder's capacity. New rules:
        # the planner's visit against every Lead's and CCEO's target.
        self.assertEqual((staff[1], staff[2]), (1, 1))
        self.assertEqual((staff[4], staff[5]), (1, 3 * 560 + 2 * 280))
        partner = cards["Partner Planning"]
        self.assertEqual((partner[4], partner[5]), (0, 1))
        self.assertEqual(cards["Schools planned twice"][4], 1)

        people_rows = list(workbook[titles[1]].iter_rows(values_only=True))
        header = people_rows[0]
        by_name = {row[2]: dict(zip(header, row)) for row in people_rows[1:]}
        self.assertEqual(by_name["Officer Two"]["New rules: visits planned"], 1)
        self.assertEqual(by_name["Officer One"]["New rules: visits planned"], 0)
        self.assertEqual(by_name["Officer One"]["Today: staff planned"], 1)
        self.assertEqual(by_name["Officer One"]["Schools assigned to Partners"], 1)
        self.assertEqual(by_name["All teams"]["Visit target"], 3 * 560 + 2 * 280)

        twice = list(workbook[titles[3]].iter_rows(values_only=True))
        self.assertEqual(twice[1][0], school.school_id)

    def test_only_the_country_director_and_the_admin_get_the_check(self):
        self.assertTrue(figures_check.may_check(self.admin_user))
        workbook = self.workbook(self.ia_user)
        self.assertFalse(
            [name for name in workbook.sheetnames if name.startswith("Check")]
        )

    def test_the_csv_is_still_the_school_sheet(self):
        client = Client()
        client.force_login(self.cd_user)
        response = client.get(
            f"/country-planning-oversight/coverage-export?fy={FY}&format=csv"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
