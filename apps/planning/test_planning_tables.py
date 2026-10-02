"""The consolidated tables behind the Country Director's cards (owner,
2026-10-01): each card links to the table of what it counts, grouped by
Programme Lead, and a table that lists schools carries the School ID.

And the tables Impact Assessment reads the country from (owner, 2026-10-02):
every visit by staff and Partners, the schools assigned to Partners, the
schools in a project, the schools nobody has planned and every activity —
under the Programme Lead, sub-region, district, CCEO or Programme Lead and
cluster of the school, with the activity planned and its SSA intervention,
exported with every planning detail."""

from __future__ import annotations

import io
import re

from apps.activities.models import Activity
from apps.core.enums import ExecutorType
from apps.core.rbac import EdifyRole
from apps.geography.models import SubRegion
from apps.projects.models import Project, ProjectSchoolAssignment
from apps.planning.country_oversight import service as svc
from apps.planning.country_oversight import tables
from apps.planning.test_country_planning_oversight import FY, World, day

PLACE = ("Lead A", "Rule Buganda", "CPO Kampala", "Officer One", "CPO Cluster")


def tables_filter_fields():
    from apps.frontend.views.country_oversight_views import TABLE_FILTER_FIELDS

    return TABLE_FILTER_FIELDS


class TableWorld(World):
    def table(self, key, extra=None, user=None, **filters):
        return tables.build(
            user or self.cd_user, svc.Filters(fy=FY, **filters), key, extra
        )

    def placed(self):
        """The district in a sub-region, so a school has a whole place."""
        sub_region = SubRegion.objects.create(name="Rule Buganda", region=self.region)
        self.district.sub_region = sub_region
        self.district.save()
        return sub_region

    def shown(self, key, **filters):
        """The page's column headings for a table."""
        return [column.label for column in self.table(key, **filters).shown_columns]

    def page(self, key, user=None, query=""):
        return self.as_user(user or self.cd_user).get(
            f"/country-planning-oversight/table/{key}?fy={FY}{query}"
        )

    def country(self, **filters):
        return svc.snapshot_for(
            self.cd_user, svc.Filters(fy=FY, **filters)
        ).tree.country


class TableRulesTest(TableWorld):
    def test_every_card_opens_a_table(self):
        self.assertEqual(tables.METRIC_TABLES, svc.KPI_TABLES)
        self.assertEqual(set(svc.KPI_TABLES), set(svc.KPI_KEYS))
        self.assertEqual(set(tables.TAB_ORDER), set(tables.SPECS))

    def test_a_table_that_lists_schools_carries_the_school_id(self):
        """Owner, 2026-10-01: "every consolidated table with list of school
        should have school ID" — beside the school, on the page and in the
        workbook."""
        for spec in tables.SPECS.values():
            keys = [column.key for column in spec.columns]
            with self.subTest(table=spec.key):
                if "school" in keys:
                    self.assertIn("school_id", keys)
                    self.assertEqual(keys.index("school_id") + 1, keys.index("school"))
                    self.assertNotIn("school_id", spec.export_only)

    def test_every_table_opens_with_where_the_school_sits(self):
        """Owner, 2026-10-02: "a normal table with sub region, cceo, PL,
        cluster all in columns" — on the page, not only in the workbook."""
        place = ["Programme Lead", "Sub-region", "District", "CCEO / PL", "Cluster"]
        for key in tables.SPECS:
            with self.subTest(table=key):
                self.assertEqual(self.shown(key)[:5], place)
        for key in ("visits", "all-visits", "plans", "not-planned", "duplicates"):
            with self.subTest(table=key):
                self.assertEqual(
                    self.shown(key)[5:8], ["School ID", "School Name", "School Type"]
                )
        self.assertEqual(self.shown("partners")[5:7], ["Partner", "School ID"])
        self.assertEqual(self.shown("projects")[5:7], ["Project", "School ID"])
        # The activity planned and its SSA intervention are on the page.
        for key in ("visits", "all-visits", "partners", "projects", "plans"):
            with self.subTest(table=key):
                self.assertIn("SSA Intervention", self.shown(key))
                self.assertIn("Planned Date", self.shown(key))


class PeopleTablesTest(TableWorld):
    def test_the_visit_table_lists_exactly_what_the_card_counts(self):
        here = self.school("client", self.cceo)
        there = self.school("core", self.cceo3)
        self.activity(here, "training_follow_up_visit", owner=self.cceo2, on=30)
        self.activity(here, "school_visit_ssa_collection", owner=self.pl, on=20)
        self.activity(there, "core_visit", owner=self.cceo3, on=10)
        self.activity(here, "donor_visit", owner=self.cceo)
        table = self.table("visits")
        self.assertEqual(len(table.rows), self.country().p_visits)
        # Under whoever holds the school; who planned the visit is a column.
        self.assertEqual(
            [
                (row["lead"], row["holder"], row["staff"], row["activity"])
                for row in table.rows
            ],
            [
                ("Lead A", "Officer One", "Lead A", "SSA Support"),
                ("Lead A", "Officer One", "Officer Two", "Follow up"),
                ("Lead B", "Officer Three", "Officer Three", "Follow up"),
            ],
        )
        self.assertEqual(table.rows[0]["school_id"], here.school_id)
        self.assertEqual(table.summary, "3 visits planned · 3 of 2,240 · 2,237 to plan")
        # The page's Programme Lead choice keeps that Lead's team.
        self.assertEqual(len(self.table("visits", program_lead=self.pl2.id).rows), 1)
        self.assertEqual(len(self.table("visits", school_type="core").rows), 1)

    def test_the_partner_table_shows_assigned_and_what_the_partner_dated(self):
        waiting = self.school("client", self.cceo)
        booked = self.school("client", self.cceo)
        dated = self.school("core_trained", self.cceo3)
        self.handover(
            waiting,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        self.activity(
            booked,
            "school_visit",
            partner=self.partner,
            status="partner_scheduled",
            monitored_by_staff_id=self.cceo.id,
            executor_type=ExecutorType.CERTIFIED_PARTNER_AGENCY.value,
            partner_date_set_by="staff",
        )
        self.activity(
            dated,
            "school_visit_ssa_collection",
            partner=self.partner2,
            status="partner_scheduled",
            monitored_by_staff_id=self.cceo3.id,
            partner_date_set_by="partner",
        )
        table = self.table("partners")
        country = self.country()
        self.assertEqual(len(table.rows), country.pa_work)
        states = {row["school_id"]: row["state"] for row in table.rows}
        self.assertEqual(states[waiting.school_id], "Awaiting partner schedule")
        # A day staff entered is not the Partner's plan (owner, 2026-10-02).
        self.assertEqual(states[booked.school_id], "Awaiting partner schedule")
        self.assertEqual(states[dated.school_id], "Planned by the Partner")
        self.assertEqual(
            sum(1 for state in states.values() if state == "Planned by the Partner"),
            country.pp_work,
        )
        by_school = {row["school_id"]: row for row in table.rows}
        self.assertIsNone(by_school[waiting.school_id]["date"])
        self.assertEqual(by_school[waiting.school_id]["staff"], "Officer One")
        self.assertEqual(by_school[dated.school_id]["partner"], "Partner Beta")
        self.assertEqual(by_school[dated.school_id]["activity"], "SSA Support")
        self.assertIn("1 of 3 planned by the Partner", table.summary)
        # Until the Partner sets it, no date is shown — only that staff
        # entered one, for the workbook.
        self.assertIsNone(by_school[booked.school_id]["date"])
        self.assertIsNotNone(by_school[booked.school_id]["staff_date"])
        self.assertEqual(by_school[booked.school_id]["status"], "Assigned to Partner")

    def test_all_plans_lists_each_activity_once(self):
        school = self.school("client", self.cceo, cluster=self.cluster)
        self.activity(
            school,
            "in_school_training",
            owner=self.cceo,
            purpose_type="in_school_training",
        )
        self.activity(school, "donor_visit", owner=self.cceo)
        self.session("cluster_meeting", [school])
        self.session("cluster_training", [school])
        self.handover(school, self.partner, status="pending_scheduling")
        table = self.table("plans")
        self.assertEqual(len(table.rows), 5)
        counted = {row["activity"]: row["counted"] for row in table.rows}
        self.assertEqual(counted["In-school Training"], "Yes")
        self.assertEqual(counted["Donor Visit"], "No")
        trainings = self.table("trainings")
        # The in-school training is a training too; the cluster training
        # reaches the school on its roster.
        self.assertEqual(len(trainings.rows), 2)
        self.assertEqual(
            sorted(row["invited"] for row in trainings.rows if "invited" in row), [1]
        )

    def test_cluster_tables_say_the_lead_sub_region_district_and_cluster(self):
        sub_region = SubRegion.objects.create(name="Rule Buganda", region=self.region)
        self.district.sub_region = sub_region
        self.district.save()
        self.cluster.responsible_staff_id = self.cceo2.id
        self.cluster.save()
        member = self.school("client", self.cceo, cluster=self.cluster)
        self.school("client", self.cceo)
        self.session("cluster_meeting", [member])
        clusters = self.table("clusters")
        self.assertEqual(len(clusters.rows), self.country().clustered)
        row = clusters.rows[0]
        # The CCEO responsible for the cluster, not whoever holds the school.
        place = ("Lead A", "Rule Buganda", "CPO Kampala", "Officer Two", "CPO Cluster")
        self.assertEqual(clusters.place_of(row), place)
        self.assertEqual(row["school_id"], member.school_id)
        meetings = self.table("meetings")
        self.assertEqual(len(meetings.rows), self.country().p_meetings)
        meeting = meetings.rows[0]
        self.assertEqual(meetings.place_of(meeting), place)
        self.assertEqual((meeting["staff"], meeting["invited"]), ("Officer One", 1))


class SchoolTablesTest(TableWorld):
    def test_planned_twice_lists_every_plan_under_its_school(self):
        """One unique school, however many plans: the duplicate is shown with
        each plan that makes it, never removed."""
        twice = self.school("client", self.cceo)
        self.activity(twice, "training_follow_up_visit", owner=self.cceo2, on=10)
        self.handover(
            twice,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        self.activity(
            twice,
            "school_visit_ssa_collection",
            partner=self.partner2,
            status="partner_scheduled",
            partner_date_set_by="partner",
            on=40,
        )
        self.school("client", self.cceo3)
        country = self.country()
        self.assertEqual(country.duplicates, 1)
        # Coverage counts the school once.
        self.assertEqual((country.any_visit, country.planned), (1, 1))
        table = self.table("duplicates")
        self.assertEqual(len(table.rows), 3)
        self.assertEqual({row["school_id"] for row in table.rows}, {twice.school_id})
        self.assertEqual(
            sorted((row["channel"], row["staff"]) for row in table.rows),
            [
                ("Partner", "Partner Alpha (assigned by Officer One)"),
                ("Partner", "Partner Beta (assigned by Officer One)"),
                ("Staff", "Officer Two"),
            ],
        )
        row = table.rows[0]
        # Under whoever HOLDS the school, whoever planned there.
        self.assertEqual((row["lead"], row["holder"]), ("Lead A", "Officer One"))
        self.assertEqual(row["why"], "Planned by staff and held by a Partner")
        self.assertEqual(table.summary, "1 school planned twice · 3 plans between them")

    def test_a_plan_at_a_school_planned_twice_is_flagged_where_it_is_listed(self):
        twice = self.school("client", self.cceo)
        once = self.school("client", self.cceo)
        self.activity(twice, "training_follow_up_visit", owner=self.cceo, on=10)
        self.activity(twice, "training_follow_up_visit", owner=self.cceo, on=20)
        self.activity(once, "training_follow_up_visit", owner=self.cceo, on=30)
        flags = {(row["school_id"], row["flag"]) for row in self.table("visits").rows}
        self.assertEqual(
            flags, {(twice.school_id, "Planned twice"), (once.school_id, "")}
        )
        # Both instances are listed: the count of plans is not reduced.
        self.assertEqual(len(self.table("visits").rows), 3)
        self.assertEqual(self.country().p_visits, 3)
        self.assertEqual(self.country().any_visit, 2)

    def test_a_core_school_with_more_visits_than_it_takes_is_captured(self):
        core = self.school("core", self.cceo)
        for offset in (10, 20, 30):
            self.activity(core, "core_visit", owner=self.cceo, on=offset)
        country = self.country()
        self.assertEqual((country.duplicates, country.staff), (1, 2))
        table = self.table("duplicates")
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(
            table.rows[0]["why"], "More staff visits than a Core school takes"
        )

    def test_not_yet_planned_no_training_and_unclustered_match_the_page(self):
        open_core = self.school("core", self.cceo3)
        self.school("champion", self.cceo3)
        clustered = self.school("client", self.cceo, cluster=self.cluster)
        self.activity(clustered, "training_follow_up_visit", owner=self.cceo)
        country = self.country()
        not_planned = self.table("not-planned")
        self.assertEqual(len(not_planned.rows), country.unplanned)
        # A Champion school needs no visit: it is not "not yet planned".
        self.assertEqual(
            [row["school_id"] for row in not_planned.rows], [open_core.school_id]
        )
        self.assertEqual(len(self.table("no-training").rows), country.no_training)
        self.assertEqual(len(self.table("unclustered").rows), country.unclustered)
        self.assertEqual(len(self.table("not-planned", school_type="client").rows), 0)


class FlatTablesTest(TableWorld):
    """Owner, 2026-10-02: an ordinary table, as My Plan and Planning Oversight
    draw theirs — every row a record, no heading rows between them."""

    def test_a_page_is_rows_only_and_the_summary_states_the_balance(self):
        school = self.school("client", self.cceo)
        self.activity(school, "training_follow_up_visit", owner=self.cceo)
        self.activity(school, "school_visit_ssa_collection", owner=self.pl, on=40)
        table = self.table("visits")
        page = tables.page_of(table, 1)
        self.assertEqual(len(page["lines"]), 2)
        for line in page["lines"]:
            self.assertEqual(len(line["cells"]), len(page["columns"]))
            self.assertEqual(
                [value for _column, value, _blank in line["cells"][:5]],
                ["Lead A", "", "CPO Kampala", "Officer One", "Unclustered"],
            )
        # Planned AND remaining (owner, 2026-10-01), for the selection.
        self.assertEqual(table.summary, "2 visits planned · 2 of 2,240 · 2,238 to plan")
        self.assertEqual(
            self.table("visits", program_lead=self.pl.id).summary,
            "2 visits planned · 2 of 1,400 · 1,398 to plan",
        )
        body = self.page("visits").content.decode()
        self.assertNotIn("edify-group-head", body)
        # No column stays put while the rest scroll (owner, 2026-10-02).
        self.assertIn('data-columns="unfrozen"', body)
        for label in (
            "Programme Lead",
            "Sub-region",
            "District",
            "CCEO / PL",
            "Cluster",
        ):
            with self.subTest(column=label):
                self.assertIn(f">{label}</th>", body)

    def test_the_partner_table_states_what_the_partner_has_dated(self):
        school = self.school("client", self.cceo)
        self.handover(
            school,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        table = self.table("partners")
        self.assertEqual(
            table.summary,
            "1 assignment at 1 school · 0 of 1 planned by the Partner · "
            "1 awaiting partner schedule",
        )
        self.assertEqual(table.rows[0]["partner"], "Partner Alpha")
        self.assertIn("Partner", self.shown("partners"))


class CardsAreTheirTablesTest(TableWorld):
    """KPI → table → records: the number on a card is the number of records
    its table lists, and the levels add up."""

    def setUp(self):
        core = self.school("core", self.cceo, cluster=self.cluster)
        client = self.school("client", self.cceo, cluster=self.cluster)
        other = self.school("core_trained", self.cceo3)
        self.school("core_graduate", self.cceo3)
        self.school("champion", self.cceo2)
        self.activity(core, "core_visit", owner=self.cceo, on=10)
        self.activity(client, "training_follow_up_visit", owner=self.cceo2, on=20)
        self.activity(other, "school_visit_ssa_collection", owner=self.pl2, on=30)
        self.activity(other, "donor_visit", owner=self.cceo3, on=35)
        self.handover(
            client,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        self.activity(
            core,
            "core_visit",
            partner=self.partner2,
            status="partner_scheduled",
            partner_date_set_by="partner",
            monitored_by_staff_id=self.cceo.id,
            on=50,
        )
        self.session("cluster_meeting", [core])
        self.session("cluster_training", [client])

    def cards(self, **filters):
        snapshot = svc.snapshot_for(self.cd_user, svc.Filters(fy=FY, **filters))
        return {
            card["metric_key"]: (
                int(card["part"].replace(",", "")),
                int(card["whole"].replace(",", "")),
            )
            for card in svc.kpis(snapshot)
        }, snapshot

    def test_each_card_is_the_count_of_its_tables_records(self):
        for filters in ({}, {"program_lead": self.pl.id}, {"school_type": "core"}):
            with self.subTest(filters=filters):
                cards, snapshot = self.cards(**filters)
                country = snapshot.tree.country
                visits = self.table("visits", **filters)
                self.assertEqual(len(visits.rows), cards["cpo_staff_visit_planning"][0])
                partners = self.table("partners", **filters)
                self.assertEqual(len(partners.rows), cards["cpo_partner_planning"][1])
                self.assertEqual(
                    sum(
                        1 for r in partners.rows if r["state"] == tables.PARTNER_PLANNED
                    ),
                    cards["cpo_partner_planning"][0],
                )
                self.assertEqual(
                    len(self.table("clusters", **filters).rows),
                    cards["cpo_cluster_membership"][0]
                    if not filters.get("program_lead")
                    else len(self.table("clusters", **filters).rows),
                )
                self.assertEqual(
                    len(self.table("meetings", **filters).rows), country.p_meetings
                )
                self.assertEqual(
                    len(self.table("not-planned", **filters).rows), country.unplanned
                )
                self.assertEqual(
                    len(self.table("no-training", **filters).rows), country.no_training
                )
                plans = self.table("plans", **filters)
                self.assertEqual(
                    sum(1 for r in plans.rows if r["counted"] == "Yes"),
                    country.p_visits,
                )
                self.assertEqual(
                    sum(1 for r in plans.rows if r["counted"] == "No"),
                    country.p_outreach,
                )

    def test_country_is_its_leads_and_each_lead_its_people(self):
        _cards, snapshot = self.cards()
        tree = snapshot.tree
        for field in (
            "target",
            "p_visits",
            "pa_work",
            "pp_work",
            "p_meetings",
            "visit_slots",
            "staff",
            "partner_scheduled",
            "any_visit",
            "no_visit",
            "unplanned",
            "training",
            "duplicates",
        ):
            with self.subTest(field=field):
                self.assertEqual(
                    getattr(tree.country, field),
                    sum(getattr(lead.tally, field) for lead in tree.leads),
                )
                for lead in tree.leads:
                    self.assertEqual(
                        getattr(lead.tally, field),
                        sum(getattr(owner.tally, field) for owner in lead.owners),
                    )

    def test_the_type_rows_add_up_to_the_cards(self):
        cards, snapshot = self.cards()
        rows = svc.type_rows(snapshot)
        total = rows[-1]
        self.assertTrue(total["is_total"])
        number = lambda text: int(text.replace(",", ""))  # noqa: E731
        types = rows[:-1]
        for key in ("schools", "visit_slots", "staff", "partner_scheduled", "training"):
            with self.subTest(column=key):
                self.assertEqual(
                    sum(number(row[key]) for row in types), number(total[key])
                )
        self.assertEqual(
            (
                number(total["staff"]) + number(total["partner_scheduled"]),
                number(total["visit_slots"]),
            ),
            cards["cpo_total_visit_coverage"],
        )
        self.assertEqual(
            (number(total["training"]), number(total["training_slots"])),
            cards["cpo_training_planning"],
        )
        core = next(row for row in types if row["key"] == "core")
        self.assertEqual(
            (core["visit_slots"], core["staff"], core["partner_scheduled"]),
            ("4", "1", "1"),
        )
        self.assertEqual(core["visit_remaining"], "2")


class ImpactAssessmentTablesTest(TableWorld):
    """Owner, 2026-10-02: the tables Impact Assessment reads the country
    from, each under the Programme Lead, sub-region, district, CCEO or
    Programme Lead and cluster of the school."""

    def cells(self, table, school):
        """One school's row of the workbook, by column heading."""
        sheet = tables.sheet(table)
        index = sheet["headers"].index("School ID")
        row = next(row for row in sheet["rows"] if row[index] == school.school_id)
        return dict(zip(sheet["headers"], row))

    def test_all_visit_plans_holds_staff_and_partner_visits_under_who_holds_the_school(
        self,
    ):
        self.placed()
        school = self.school("client", self.cceo, cluster=self.cluster)
        other = self.school("core_trained", self.cceo3)
        # Planned by another Lead's officer: still under the school's holder.
        self.activity(
            school,
            "training_follow_up_visit",
            owner=self.cceo3,
            on=10,
            focus_intervention="leadership",
        )
        self.activity(school, "donor_visit", owner=self.cceo, on=12)
        self.activity(
            other,
            "school_visit_ssa_collection",
            partner=self.partner2,
            status="partner_scheduled",
            partner_date_set_by="partner",
            monitored_by_staff_id=self.cceo3.id,
            on=20,
        )
        # Neither is a visit plan yet: a hand-over the Partner has not dated,
        # and a day staff entered for the Partner.
        self.handover(
            school,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        self.activity(
            other,
            "school_visit",
            partner=self.partner,
            status="partner_scheduled",
            partner_date_set_by="staff",
            executor_type=ExecutorType.CERTIFIED_PARTNER_AGENCY.value,
            on=25,
        )
        table = self.table("all-visits")
        self.assertEqual(
            [
                (
                    r["school_id"],
                    r["activity"],
                    r["channel"],
                    r["by_name"],
                    r["counted"],
                )
                for r in table.rows
            ],
            [
                (school.school_id, "Follow up", "Staff", "Officer Three", "Yes"),
                (school.school_id, "Donor Visit", "Staff", "Officer One", "No"),
                (other.school_id, "SSA Support", "Partner", "Partner Beta", ""),
            ],
        )
        self.assertEqual(table.place_of(table.rows[0]), PLACE)
        self.assertEqual(
            table.place_of(table.rows[2]),
            ("Lead B", "Rule Buganda", "CPO Kampala", "Officer Three", "Unclustered"),
        )
        self.assertEqual(table.rows[0]["intervention"], "Leadership")
        self.assertEqual(table.summary, "3 visits planned · 2 by staff · 1 by Partners")
        # The Programme Lead choice keeps a person's plan, as the cards do.
        self.assertEqual(
            [
                r["by_name"]
                for r in self.table("all-visits", program_lead=self.pl.id).rows
            ],
            ["Officer One"],
        )

    def test_awaiting_partner_schedule_becomes_the_date_the_partner_sets(self):
        self.placed()
        school = self.school("client", self.cceo, cluster=self.cluster)
        handover = self.handover(
            school,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
            focus_intervention="leadership",
            purpose_of_visit="training_follow_up",
        )
        table = self.table("partners")
        self.assertEqual(table.place_of(table.rows[0]), PLACE)
        cells = self.cells(table, school)
        self.assertEqual(cells["Planned Date"], "Awaiting partner schedule")
        self.assertEqual(cells["Activity"], "Follow up")
        self.assertEqual(cells["SSA Intervention"], "Leadership")
        self.assertEqual(cells["Assigned By"], "Officer One")
        body = self.page("partners").content.decode()
        self.assertIn(">Awaiting partner schedule</span>", body)
        self.assertNotIn("Not yet scheduled", body)

        # The Partner schedules it: the hand-over becomes a dated activity.
        activity = self.activity(
            school,
            "training_follow_up_visit",
            partner=self.partner,
            status="partner_scheduled",
            partner_date_set_by="partner",
            monitored_by_staff_id=self.cceo.id,
            focus_intervention="leadership",
            on=33,
        )
        handover.status = "partner_scheduled"
        handover.scheduled_activity = activity
        handover.save()
        table = self.table("partners")
        self.assertEqual(len(table.rows), 1)
        self.assertEqual(table.rows[0]["state"], tables.PARTNER_PLANNED)
        self.assertEqual(self.cells(table, school)["Planned Date"], day(33))
        self.assertNotIn(
            "Awaiting partner schedule</span>", self.page("partners").content.decode()
        )

    def test_unplanned_schools_leave_out_a_school_in_a_partners_hands(self):
        """Owner, 2026-10-02: a school assigned to a Partner is the Partner's
        to plan, and is followed in the Partner table."""
        self.placed()
        nobody = self.school("client", self.cceo, cluster=self.cluster)
        held = self.school("client", self.cceo)
        partly = self.school("core", self.cceo3)
        self.school("champion", self.cceo3)
        self.handover(
            held,
            self.partner,
            status="pending_scheduling",
            assigning_staff_id=self.cceo.id,
        )
        self.activity(partly, "core_visit", owner=self.cceo3, on=10)
        country = self.country()
        self.assertEqual(
            (country.no_visit, country.unplanned, country.partner_to_plan), (2, 1, 1)
        )
        unplanned = self.table("not-planned")
        self.assertEqual([r["school_id"] for r in unplanned.rows], [nobody.school_id])
        self.assertEqual(unplanned.place_of(unplanned.rows[0]), PLACE)
        cells = self.cells(unplanned, nobody)
        self.assertEqual(
            (
                cells["Visits Needed"],
                cells["Visits Remaining"],
                cells["Trainings Needed"],
            ),
            (1, 1, 1),
        )
        self.assertEqual(
            [r["school_id"] for r in self.table("partners").rows], [held.school_id]
        )
        partly_planned = self.table("partly-planned")
        self.assertEqual(
            [r["school_id"] for r in partly_planned.rows], [partly.school_id]
        )
        cells = self.cells(partly_planned, partly)
        self.assertEqual(
            (
                cells["Visits Needed"],
                cells["Visits Planned"],
                cells["Visits Remaining"],
            ),
            (4, 1, 3),
        )
        # The card and the by-type row say the table's number.
        snapshot = svc.snapshot_for(self.cd_user, svc.Filters(fy=FY))
        card = next(
            card
            for card in svc.kpis(snapshot)
            if card["metric_key"] == "cpo_total_visit_coverage"
        )
        self.assertIn("1 schools not yet planned", card["extras"])
        self.assertIn("1 awaiting a Partner's date", card["extras"])
        client_row = next(r for r in svc.type_rows(snapshot) if r["key"] == "client")
        self.assertEqual((client_row["no_visit"], client_row["unplanned"]), ("2", "1"))

    def project(self):
        self.coord_user, self.coord = self.person(
            "cpo-pc", "Coordinator Cee", EdifyRole.PROJECT_COORDINATOR
        )
        return Project.objects.create(
            name="Rule Project",
            code="SP-RULE",
            category="pilot",
            status="active",
            intervention="leadership",
            target_interventions=["leadership"],
            manager_staff_id=self.coord.id,
        )

    def test_project_schools_are_read_from_the_coordinators_records(self):
        self.placed()
        project = self.project()
        waiting = self.school("client", self.cceo, cluster=self.cluster)
        planned = self.school("core", self.cceo3)
        handed = self.school("client", self.cceo2)
        for school in (waiting, planned, handed):
            ProjectSchoolAssignment.objects.create(
                project=project, school=school, assigned_by=self.cceo_user.id
            )
        Activity.objects.create(
            activity_type="school_visit",
            school=planned,
            fy=FY,
            quarter="Q1",
            planned_date=day(20),
            status="scheduled",
            responsible_staff_id=self.coord.id,
            delivery_type="staff",
            project_id=project.id,
            purpose_type="training_follow_up",
            focus_intervention="leadership",
        )
        self.handover(
            handed,
            self.partner,
            status="pending_scheduling",
            project=project,
            assigning_staff_id=self.coord.id,
        )
        for reader in (self.cd_user, self.ia_user):
            with self.subTest(role=reader.active_role):
                table = self.table("projects", user=reader)
                rows = {row["school_id"]: row for row in table.rows}
                self.assertEqual(
                    set(rows), {s.school_id for s in (waiting, planned, handed)}
                )
                row = rows[waiting.school_id]
                self.assertEqual(table.place_of(row), PLACE)
                self.assertEqual(row["project"], "Rule Project")
                self.assertEqual(
                    (row["date"], row["date_note"], row["status"]),
                    (
                        None,
                        "Awaiting Project Coordinator",
                        "Awaiting Project Coordinator Action",
                    ),
                )
                self.assertEqual(
                    (row["coordinator"], row["added_by"]),
                    ("Coordinator Cee", "Officer One"),
                )
                row = rows[planned.school_id]
                self.assertEqual(
                    (row["lead"], row["holder"], row["channel"], row["by_name"]),
                    ("Lead B", "Officer Three", "Staff", "Coordinator Cee"),
                )
                self.assertEqual(
                    (row["date"], row["intervention"]), (day(20), "Leadership")
                )
                row = rows[handed.school_id]
                self.assertEqual(
                    (row["date"], row["date_note"], row["channel"], row["by_name"]),
                    (None, "Awaiting partner schedule", "Partner", "Partner Alpha"),
                )
                self.assertEqual(
                    table.summary,
                    "3 schools in 1 project · 2 planned · 1 awaiting a plan",
                )
        self.assertEqual(
            self.cells(self.table("projects"), handed)["Planned Date"],
            "Awaiting partner schedule",
        )
        # The coordinator's work says its project wherever it is listed, and
        # a table can be narrowed to one project.
        activities = self.table("plans")
        self.assertEqual(self.cells(activities, planned)["Project"], "Rule Project")
        only = tables.Extra(project=project.id)
        self.assertEqual(
            sorted(r["school_id"] for r in self.table("plans", extra=only).rows),
            sorted([planned.school_id, handed.school_id]),
        )
        self.school("client", self.cceo)
        self.assertEqual(
            [r["school_id"] for r in self.table("not-planned", extra=only).rows],
            [waiting.school_id],
        )

    def test_all_activities_names_every_activity_and_its_ssa_intervention(self):
        self.placed()
        self.cluster.responsible_staff_id = self.cceo.id
        self.cluster.save()
        school = self.school("client", self.cceo, cluster=self.cluster)
        facilitated = self.session("cluster_training", [school], on=41)
        facilitated.facilitating_partner_id = self.partner.id
        facilitated.focus_intervention = "leadership"
        facilitated.save()
        # A cluster training a Partner delivers has no school of its own.
        Activity.objects.create(
            activity_type="cluster_training",
            cluster=self.cluster,
            fy=FY,
            quarter="Q1",
            planned_date=day(45),
            status="partner_scheduled",
            delivery_type="partner",
            assigned_partner_id=self.partner2.id,
            partner_date_set_by="partner",
            monitored_by_staff_id=self.cceo.id,
            focus_intervention="financial_health",
        )
        self.activity(
            school,
            "school_visit_ssa_collection",
            owner=self.cceo,
            on=20,
            salesforce_activity_id="SVE-0001",
        )
        self.activity(school, "field_event", owner=self.cceo, on=22)
        table = self.table("plans")
        self.assertEqual(
            [
                (r["activity"], r["channel"], r["by_name"], r["intervention"])
                for r in table.rows
            ],
            [
                ("SSA Support", "Staff", "Officer One", "Data Gathering"),
                ("Field Event", "Staff", "Officer One", ""),
                ("Cluster Training", "Staff", "Officer One", "Leadership"),
                ("Cluster Training", "Partner", "Partner Beta", "Financial Health"),
            ],
        )
        for row in table.rows:
            self.assertEqual(table.place_of(row), PLACE)
        self.assertEqual(table.rows[2]["facilitator"], "Partner Alpha")
        self.assertEqual(table.summary, "4 activities · 3 by staff · 1 with Partners")
        headers = tables.sheet(table)["headers"]
        for label in (
            "Sub-region",
            "Cluster",
            "Activity",
            "SSA Intervention",
            "Purpose",
            "Project",
            "Facilitated By",
            "Planned Cost (UGX)",
            "Salesforce ID",
            "IA Verification",
        ):
            with self.subTest(column=label):
                self.assertIn(label, headers)
        self.assertEqual(self.cells(table, school)["Salesforce ID"], "SVE-0001")
        # The training table holds the Partner's cluster training too.
        self.assertEqual(
            sorted((r["channel"], r["by_name"]) for r in self.table("trainings").rows),
            [("Partner", "Partner Beta"), ("Staff", "Officer One")],
        )

    def test_a_table_narrows_to_a_sub_region_and_a_cluster(self):
        sub_region = self.placed()
        inside = self.school("client", self.cceo, cluster=self.cluster)
        outside = self.school("client", self.cceo)
        for school in (inside, outside):
            self.activity(school, "training_follow_up_visit", owner=self.cceo)
        here = tables.Extra(cluster=self.cluster.id)
        self.assertEqual(
            [r["school_id"] for r in self.table("all-visits", extra=here).rows],
            [inside.school_id],
        )
        self.assertEqual(
            len(
                self.table(
                    "all-visits", extra=tables.Extra(sub_region=sub_region.id)
                ).rows
            ),
            2,
        )
        self.assertEqual(
            len(
                self.table("all-visits", extra=tables.Extra(sub_region="nowhere")).rows
            ),
            0,
        )
        body = self.page(
            "all-visits", query=f"&cluster={self.cluster.id}"
        ).content.decode()
        self.assertIn(f">{inside.school_id}</td>", body)
        self.assertNotIn(f">{outside.school_id}</td>", body)
        # The choice rides along to the other tables and to the workbook.
        self.assertIn(
            f"/country-planning-oversight/table/partners?fy={FY}&amp;period=fy&amp;cluster={self.cluster.id}",
            body,
        )
        self.assertIn(
            f"/country-planning-oversight/table-export/all-visits?fy={FY}&amp;period=fy&amp;cluster={self.cluster.id}",
            body,
        )


class TablePageTest(TableWorld):
    def test_the_cards_are_links_to_their_tables(self):
        self.school("client", self.cceo)
        body = (
            self.as_user(self.cd_user)
            .get(f"/country-planning-oversight/?fy={FY}")
            .content.decode()
        )
        for key in svc.KPI_TABLES.values():
            with self.subTest(table=key):
                self.assertRegex(
                    body,
                    rf'<a class="cpo-metric"[^>]*href="/country-planning-oversight/table/{key}\?',
                )
        self.assertIn(
            "/country-planning-oversight/table/not-planned?school_type=client", body
        )

    def test_a_table_page_shows_groups_school_ids_and_one_export(self):
        school = self.school("client", self.cceo)
        self.activity(school, "training_follow_up_visit", owner=self.cceo)
        response = self.page("visits")
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn("Staff Visit Plans", body)
        self.assertIn(">Lead A</td>", body)
        self.assertIn(">Officer One</td>", body)
        self.assertIn(">School ID</th>", body)
        self.assertIn(f">{school.school_id}</td>", body)
        self.assertIn(f'href="/schools/{school.school_id}"', body)
        self.assertEqual(
            body.count("/country-planning-oversight/table-export/visits"), 1
        )
        # Every table is a tab away.
        for key in tables.TAB_ORDER:
            self.assertIn(f"/country-planning-oversight/table/{key}?", body)
        # And can be narrowed to a place, a person, a Partner or a project.
        for name in tables_filter_fields():
            with self.subTest(filter=name):
                self.assertIn(f'<select name="{name}"', body)

    def test_impact_assessment_exports_the_tables_in_one_workbook(self):
        import openpyxl

        school = self.school("client", self.cceo, cluster=self.cluster)
        self.activity(school, "training_follow_up_visit", owner=self.cceo)
        client = self.as_user(self.ia_user)
        body = client.get(
            f"/country-planning-oversight/table/plans?fy={FY}"
        ).content.decode()
        self.assertEqual(body.count("/country-planning-oversight/table-export/all?"), 1)
        export = client.get(f"/country-planning-oversight/table-export/all?fy={FY}")
        self.assertEqual(export.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(export.content))
        self.assertEqual(
            workbook.sheetnames, [tables.SPECS[key].title for key in tables.WORKBOOK]
        )
        for key in tables.WORKBOOK:
            sheet = workbook[tables.SPECS[key].title]
            with self.subTest(sheet=key):
                self.assertEqual(
                    [cell.value for cell in sheet[1]][:5],
                    [
                        "Programme Lead",
                        "Sub-region",
                        "District",
                        "CCEO / PL",
                        "Cluster",
                    ],
                )
        self.assertEqual(workbook["All Activities"].max_row, 2)
        # Without the export permission there is no workbook.
        self.assertNotEqual(
            self.as_user(self.pl_user)
            .get(f"/country-planning-oversight/table-export/all?fy={FY}")
            .status_code,
            200,
        )

    def test_every_table_opens_and_exports(self):
        import openpyxl

        school = self.school("client", self.cceo, cluster=self.cluster)
        self.activity(school, "training_follow_up_visit", owner=self.cceo)
        client = self.as_user(self.cd_user)
        for key, spec in tables.SPECS.items():
            with self.subTest(table=key):
                self.assertEqual(self.page(key).status_code, 200)
                export = client.get(
                    f"/country-planning-oversight/table-export/{key}?fy={FY}"
                )
                self.assertEqual(export.status_code, 200)
                sheet = openpyxl.load_workbook(io.BytesIO(export.content)).active
                headers = [cell.value for cell in sheet[1]]
                # The workbook carries the groups as columns.
                self.assertEqual(headers, [column.label for column in spec.columns])
                self.assertEqual(headers[0], "Programme Lead")

    def test_every_table_in_the_platform_holds_fifty_rows_a_page(self):
        """Owner, 2026-10-02: "All table in the platform should hold 50
        records in each page the rest hidden in the paginations"."""
        from apps.core.pagination import TABLE_PAGE_SIZE, paginate_rows

        self.assertEqual(TABLE_PAGE_SIZE, 50)
        self.assertEqual(tables.ROWS_PER_PAGE, TABLE_PAGE_SIZE)
        pager = paginate_rows(list(range(120)), page=2)
        self.assertEqual((len(pager["rows"]), pager["page_count"]), (50, 3))

    def test_a_long_table_is_paged_and_keeps_its_filters(self):
        school = self.school("client", self.cceo)
        for offset in range(tables.ROWS_PER_PAGE + 5):
            self.activity(
                school, "training_follow_up_visit", owner=self.cceo, on=offset
            )
        body = self.page("visits", query=f"&program_lead={self.pl.id}").content.decode()
        self.assertEqual(
            len(re.findall(r">Follow up</td>", body)), tables.ROWS_PER_PAGE
        )
        self.assertIn("page=2", body)
        self.assertIn(f"program_lead={self.pl.id}", body)
        second = self.page("visits", query="&page=2").content.decode()
        self.assertEqual(len(re.findall(r">Follow up</td>", second)), 5)
        # The heading says how many the whole group holds, on every page.
        self.assertIn(f"{tables.ROWS_PER_PAGE + 5} visits", second)

    def test_who_may_open_a_table(self):
        self.school("client", self.cceo)
        for user, allowed in (
            (self.cd_user, True),
            (self.ia_user, True),
            (self.pl_user, False),
            (self.cceo_user, False),
        ):
            with self.subTest(role=user.active_role):
                status = self.page("visits", user).status_code
                self.assertEqual(status == 200, allowed)
        # The RVP reads the country's figures, never its school rows.
        body = self.page("visits", self.rvp_user).content.decode()
        self.assertIn("School-level rows are not part of this role", body)
        self.assertNotIn("<td", body)
        self.assertEqual(self.page("nonsense").status_code, 404)
